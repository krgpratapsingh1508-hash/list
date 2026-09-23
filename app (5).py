import streamlit as st
import pandas as pd
import sqlite3
import json
import io
import os
import hashlib
import xml.etree.ElementTree as ET
from html.parser import HTMLParser
import streamlit.components.v1 as components
import html as _html
import re as _re

st.set_page_config(page_title="NEP Master Data System", page_icon="🎓", layout="wide")

# =========================================================================
# 📋 लिस्ट / टेबल डिज़ाइन: जिन लिस्ट में अपना रंग नहीं है उनमें एक-एक छोड़कर हल्की धारीदार (zebra) रो
# (जिन टेबल में लाल/नीले/हरे रंग पहले से हैं, उन्हें बिल्कुल नहीं छेड़ा जाता)
# =========================================================================
_orig_st_dataframe = st.dataframe

def _zebra_rows(d):
    out = pd.DataFrame("", index=d.index, columns=d.columns)
    out.iloc[1::2, :] = "background-color: #f4f7fd"
    return out

def _designed_dataframe(data=None, *args, **kwargs):
    try:
        if (isinstance(data, pd.DataFrame) and 0 < len(data) <= 3000 and data.shape[1] > 0
                and data.index.is_unique and data.columns.is_unique):
            data = data.style.apply(_zebra_rows, axis=None)
    except Exception:
        pass
    return _orig_st_dataframe(data, *args, **kwargs)

st.dataframe = _designed_dataframe

# =========================================================================
# 🏷️ हेडलाइन डिज़ाइन: st.title / st.header / st.subheader और "#, ##, ###, ####" वाली markdown हेडिंग
# अपने-आप ग्रेडिएंट बैनर बन जाती हैं (कोड में कहीं कुछ बदलने की ज़रूरत नहीं)
# =========================================================================
_orig_st_markdown = st.markdown
_HEAD_RE = _re.compile(r'^\s*(#{1,4})\s+([^\n]+?)\s*$')

def _banner_html(text, level):
    t = _html.escape(str(text)).replace("$", "&#36;")
    t = _re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)
    return f'<div class="hb hb{int(level)}">{t}</div>'

def _designed_markdown(body="", *args, **kwargs):
    if isinstance(body, str) and not args and not kwargs.get("unsafe_allow_html"):
        _m = _HEAD_RE.match(body)
        if _m:
            return _orig_st_markdown(_banner_html(_m.group(2), len(_m.group(1))), unsafe_allow_html=True)
    return _orig_st_markdown(body, *args, **kwargs)

def _make_head_fn(level):
    def _fn(body="", *args, **kwargs):
        return _orig_st_markdown(_banner_html(body, level), unsafe_allow_html=True)
    return _fn

# =========================================================================
# ✍️ लिखा हुआ टेक्स्ट (st.write / st.caption / st.info / st.success / st.warning / st.error) का डिज़ाइन
# =========================================================================
_orig_st_caption = st.caption
_orig_st_write = st.write
_LIST_LIKE_RE = _re.compile(r'(^|\n)\s*([-*+]\s|\d+[.)]\s|\||#|>)')

def _rich_text(t):
    t = _html.escape(str(t)).replace("$", "&#36;")
    t = _re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)
    t = _re.sub(r"`([^`]+?)`", r"<code>\1</code>", t)
    return t.replace("\n", "<br>")

def _plain_text_ok(body):
    return isinstance(body, str) and body.strip() != "" and not _LIST_LIKE_RE.search(body)

_ALERT_SYMBOL = {"success": "✓", "info": "i", "warning": "!", "error": "✕"}

_orig_alerts = {_k: getattr(st, _k) for _k in ("success", "info", "warning", "error")}

def _make_alert_fn(kind):
    def _fn(body="", *args, icon=None, **kwargs):
        if not _plain_text_ok(body):
            # सूची/टेबल जैसा मार्कडाउन या गैर-टेक्स्ट → Streamlit का अपना अलर्ट (कुछ भी छूटेगा नहीं)
            if icon is not None:
                kwargs["icon"] = icon
            return _orig_alerts[kind](body, *args, **kwargs)
        _ic = _html.escape(str(icon)) if icon else _ALERT_SYMBOL[kind]
        return _orig_st_markdown(
            f'<div class="al al-{kind}"><div class="al-ic">{_ic}</div><div class="al-tx">{_rich_text(body)}</div></div>',
            unsafe_allow_html=True
        )
    return _fn

def _designed_caption(body="", *args, **kwargs):
    if _plain_text_ok(body) and not args and not kwargs:
        return _orig_st_markdown(f'<div class="cap">{_rich_text(body)}</div>', unsafe_allow_html=True)
    return _orig_st_caption(body, *args, **kwargs)

def _designed_write(*args, **kwargs):
    if len(args) == 1 and not kwargs and _plain_text_ok(args[0]):
        return _orig_st_markdown(f'<div class="txt">{_rich_text(args[0])}</div>', unsafe_allow_html=True)
    return _orig_st_write(*args, **kwargs)

st.success = _make_alert_fn("success")
st.info = _make_alert_fn("info")
st.warning = _make_alert_fn("warning")
st.error = _make_alert_fn("error")
st.caption = _designed_caption
st.write = _designed_write

st.markdown = _designed_markdown
st.title = _make_head_fn(1)
st.header = _make_head_fn(2)
st.subheader = _make_head_fn(3)

# =========================================================================
# 🎨 प्रोफेशनल UI स्टाइलिंग (पूरे ऐप में कस्टम CSS)
# =========================================================================
st.markdown("""
<style>
    /* मुख्य कंटेनर पैडिंग */
    .block-container {
        padding-top: 1.5rem;
        padding-bottom: 3rem;
    }
    /* हैडिंग्स */
    h1, h2, h3 {
        font-family: 'Segoe UI', 'Trebuchet MS', sans-serif;
        letter-spacing: -0.3px;
    }
    h1 { color: #1a3c6e; }
    /* बटन */
    div.stButton > button {
        border-radius: 8px;
        font-weight: 600;
        border: 1px solid #d0d7e2;
        transition: all 0.15s ease-in-out;
    }
    div.stButton > button:hover {
        border-color: #1a73e8;
        color: #1a73e8;
    }
    div.stButton > button[kind="primary"] {
        background-color: #1a73e8;
    }
    /* डाउनलोड बटन */
    div.stDownloadButton > button {
        border-radius: 8px;
        font-weight: 600;
        background-color: #0f9d58;
        color: white;
        border: none;
    }
    div.stDownloadButton > button:hover {
        background-color: #0c8043;
        color: white;
    }
    /* साइडबार */
    section[data-testid="stSidebar"] {
        background-color: #f6f8fb;
        border-right: 1px solid #e3e8ef;
    }
    /* डेटाफ़्रेम / टेबल कार्ड जैसा दिखे */
    div[data-testid="stDataFrame"] {
        border: 1px solid #e3e8ef;
        border-radius: 10px;
        overflow: hidden;
    }
    /* मेट्रिक कार्ड्स */
    div[data-testid="stMetric"] {
        background-color: #f8fafc;
        border: 1px solid #e3e8ef;
        border-radius: 10px;
        padding: 12px 16px;
    }
    /* एक्सपैंडर */
    div[data-testid="stExpander"] {
        border: 1px solid #e3e8ef;
        border-radius: 10px;
    }
    /* टैब्स */
    button[data-baseweb="tab"] {
        font-weight: 600;
    }
    /* फुटर क्रेडिट */
    .app-footer {
        text-align: center;
        color: #8a94a6;
        font-size: 12.5px;
        padding: 18px 0 4px 0;
        border-top: 1px solid #e3e8ef;
        margin-top: 30px;
    }

    /* ============== 🔐 लॉगिन स्क्रीन स्टाइलिंग ============== */
    @keyframes floatIn {
        0% { opacity: 0; transform: translateY(14px); }
        100% { opacity: 1; transform: translateY(0); }
    }
    @keyframes gradientShift {
        0% { background-position: 0% 50%; }
        50% { background-position: 100% 50%; }
        100% { background-position: 0% 50%; }
    }
    .login-hero {
        text-align: center;
        animation: floatIn 0.6s ease-out;
        margin-bottom: 6px;
        margin-top: 38px;
        padding-top: 10px;
    }
    .login-hero .emoji-badge {
        font-size: 46px;
        display: inline-block;
        animation: floatIn 0.5s ease-out;
    }
    .login-hero .login-logo-img {
        border-radius: 14px;
        box-shadow: 0 6px 20px rgba(26, 60, 110, 0.18);
        animation: floatIn 0.5s ease-out;
        object-fit: contain;
        background: #ffffff;
        padding: 6px;
    }
    .login-hero h1 {
        font-size: 34px;
        font-weight: 800;
        margin: 6px 0 2px 0;
        background: linear-gradient(270deg, #1a73e8, #6a4cff, #0f9d58, #1a73e8);
        background-size: 600% 600%;
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
        background-clip: text;
        animation: gradientShift 6s ease infinite;
    }
    .login-hero p {
        color: #6b7688;
        font-size: 15px;
        margin-top: 0;
    }
    div[data-testid="stVerticalBlockBorderWrapper"] {
        animation: floatIn 0.7s ease-out;
    }
    /* लॉगिन कार्ड (st.container(border=True)) को थोड़ा प्रीमियम लुक */
    div[data-testid="stForm"], div[data-testid="stVerticalBlockBorderWrapper"] > div {
        border-radius: 16px !important;
    }
    .login-badge-row {
        display: flex;
        justify-content: center;
        gap: 8px;
        flex-wrap: wrap;
        margin: 10px 0 18px 0;
    }
    .login-badge {
        background: linear-gradient(135deg, #eef3ff, #f3eefc);
        border: 1px solid #dde4f5;
        color: #4a5b8c;
        font-size: 12px;
        font-weight: 600;
        padding: 5px 12px;
        border-radius: 20px;
    }
    .farewell-box {
        text-align: center;
        animation: floatIn 0.5s ease-out;
        background: linear-gradient(135deg, #eafff0, #eef8ff);
        border: 1px solid #cdeedd;
        border-radius: 14px;
        padding: 14px;
        margin-bottom: 18px;
        color: #1a6e3c;
        font-weight: 600;
    }

    /* ============== 🎛️ Admin सेक्शन Hide / Show टॉगल बटन ============== */
    div[class*="st-key-admin_sec_btn_"] button {
        border-radius: 999px !important;
        font-weight: 700 !important;
        letter-spacing: 0.2px;
        padding: 6px 14px !important;
        border: none !important;
        box-shadow: 0 3px 10px rgba(0, 0, 0, 0.15);
        transition: transform .12s ease, box-shadow .12s ease, filter .12s ease;
    }
    div[class*="st-key-admin_sec_btn_"] button,
    div[class*="st-key-admin_sec_btn_"] button p {
        color: #ffffff !important;
    }
    div[class*="st-key-admin_sec_btn_"] button:hover {
        transform: translateY(-1px);
        box-shadow: 0 6px 16px rgba(0, 0, 0, 0.22);
        filter: brightness(1.07);
    }
    div[class*="st-key-admin_sec_btn_"] button:active {
        transform: translateY(0) scale(0.98);
    }
    div[class*="st-key-admin_sec_btn_"][class*="__hidebtn"] button {
        background: linear-gradient(135deg, #ff7a59, #e8433f) !important;
    }
    div[class*="st-key-admin_sec_btn_"][class*="__showbtn"] button {
        background: linear-gradient(135deg, #2bb673, #0f9d58) !important;
    }
    .admin-sec-hidden-note {
        background: #f6f8fb;
        border: 1px dashed #c7d1e0;
        color: #6b7688;
        font-size: 13px;
        border-radius: 10px;
        padding: 8px 14px;
        margin: 2px 0 6px 0;
    }

    /* ============== ✨ प्रीमियम बटन डिज़ाइन (सभी बटन) ============== */
    div.stButton > button,
    div.stDownloadButton > button,
    div[data-testid="stFormSubmitButton"] > button {
        border-radius: 12px;
        font-weight: 700;
        font-size: 14.5px;
        letter-spacing: .2px;
        min-height: 44px;
        padding: .5rem 1.1rem;
        transition: transform .15s ease, box-shadow .15s ease, background .15s ease, border-color .15s ease, filter .15s ease;
    }
    /* सामान्य बटन: सफ़ेद कार्ड जैसा, हल्का शैडो */
    div.stButton > button:not([kind="primary"]) {
        background: linear-gradient(180deg, #ffffff, #f3f6fc);
        color: #1a3c6e;
        border: 1.5px solid #cfd9ea;
        box-shadow: 0 1px 2px rgba(26, 60, 110, .08);
    }
    div.stButton > button:not([kind="primary"]):hover {
        transform: translateY(-2px);
        border-color: #1a73e8;
        color: #1a73e8;
        background: linear-gradient(180deg, #ffffff, #e9f1ff);
        box-shadow: 0 8px 18px rgba(26, 115, 232, .18);
    }
    /* प्राइमरी बटन: नीला-जामुनी ग्रेडिएंट */
    div.stButton > button[kind="primary"],
    div[data-testid="stFormSubmitButton"] > button[kind="primary"],
    div[data-testid="stFormSubmitButton"] > button {
        background: linear-gradient(135deg, #1a73e8, #5b4bdb);
        color: #ffffff;
        border: none;
        box-shadow: 0 5px 16px rgba(91, 75, 219, .35);
    }
    div.stButton > button[kind="primary"] p,
    div[data-testid="stFormSubmitButton"] > button p { color: #ffffff; }
    div.stButton > button[kind="primary"]:hover,
    div[data-testid="stFormSubmitButton"] > button:hover {
        transform: translateY(-2px);
        filter: brightness(1.08);
        color: #ffffff;
        box-shadow: 0 10px 22px rgba(91, 75, 219, .42);
    }
    div.stButton > button:active,
    div.stDownloadButton > button:active,
    div[data-testid="stFormSubmitButton"] > button:active {
        transform: translateY(0) scale(.98);
        box-shadow: 0 1px 4px rgba(0, 0, 0, .18);
    }
    div.stButton > button:focus-visible,
    div.stDownloadButton > button:focus-visible {
        outline: 3px solid rgba(26, 115, 232, .35);
        outline-offset: 2px;
    }
    /* डाउनलोड बटन: हरा ग्रेडिएंट */
    div.stDownloadButton > button {
        background: linear-gradient(135deg, #22b573, #0b8f52);
        color: #ffffff;
        border: none;
        box-shadow: 0 5px 16px rgba(15, 157, 88, .32);
    }
    div.stDownloadButton > button p { color: #ffffff; }
    div.stDownloadButton > button:hover {
        transform: translateY(-2px);
        filter: brightness(1.07);
        color: #ffffff;
        background: linear-gradient(135deg, #22b573, #0b8f52);
        box-shadow: 0 10px 22px rgba(15, 157, 88, .42);
    }
    /* समरी का CSV बटन: आउटलाइन स्टाइल (XLSX हरे भरे बटन से अलग पहचान) */
    div[class*="st-key-db_dl_"]:not([class*="st-key-db_dl_xlsx_"]) button {
        background: #ffffff !important;
        border: 1.8px solid #0f9d58 !important;
        box-shadow: none !important;
    }
    div[class*="st-key-db_dl_"]:not([class*="st-key-db_dl_xlsx_"]) button,
    div[class*="st-key-db_dl_"]:not([class*="st-key-db_dl_xlsx_"]) button p {
        color: #0b7a4b !important;
    }
    div[class*="st-key-db_dl_"]:not([class*="st-key-db_dl_xlsx_"]) button:hover {
        background: #eafaf2 !important;
        box-shadow: 0 8px 18px rgba(15, 157, 88, .22) !important;
    }

    /* ============== 🧩 फ़िल्टर / इनपुट / टैब / मेट्रिक पॉलिश ============== */
    div[data-baseweb="select"] > div {
        border-radius: 10px;
        border-color: #cfd9ea;
        transition: border-color .15s ease, box-shadow .15s ease;
    }
    div[data-baseweb="select"] > div:hover { border-color: #1a73e8; }
    div[data-baseweb="select"] > div:focus-within {
        border-color: #1a73e8;
        box-shadow: 0 0 0 3px rgba(26, 115, 232, .18);
    }
    div[data-testid="stTextInput"] input {
        border-radius: 10px;
    }
    div[data-testid="stTextInput"] div[data-baseweb="input"]:focus-within {
        box-shadow: 0 0 0 3px rgba(26, 115, 232, .18);
        border-radius: 10px;
    }
    button[data-baseweb="tab"] {
        border-radius: 10px 10px 0 0;
        padding: 10px 16px;
        transition: background .15s ease, color .15s ease;
    }
    button[data-baseweb="tab"]:hover { background: #eef4ff; }
    button[data-baseweb="tab"][aria-selected="true"] {
        background: linear-gradient(180deg, #eaf2ff, #ffffff);
        color: #1a73e8;
    }
    div[data-testid="stMetric"] {
        border-left: 5px solid #1a73e8;
        box-shadow: 0 2px 8px rgba(26, 60, 110, .07);
        transition: transform .15s ease, box-shadow .15s ease;
    }
    div[data-testid="stMetric"]:hover {
        transform: translateY(-2px);
        box-shadow: 0 8px 18px rgba(26, 60, 110, .13);
    }

    /* ============== 🎓 डिग्री+ब्रांच समरी: बैनर + डाउनलोड पैनल ============== */
    .sum-banner {
        background: linear-gradient(135deg, #1a3c6e, #2f5fb3 60%, #5b4bdb);
        color: #ffffff;
        border-radius: 14px;
        padding: 14px 20px;
        margin: 4px 0 10px 0;
        box-shadow: 0 8px 22px rgba(26, 60, 110, .25);
        animation: floatIn .5s ease-out;
    }
    .sum-banner .t { font-size: 20px; font-weight: 800; letter-spacing: -.2px; }
    .dl-panel-head {
        display: flex; align-items: center; gap: 10px; flex-wrap: wrap;
        margin: 18px 0 6px 0;
    }
    .dl-panel-head .h { font-size: 17px; font-weight: 800; color: #1a3c6e; }
    .dl-chip {
        font-size: 12.5px; font-weight: 700; padding: 4px 12px; border-radius: 999px;
        border: 1px solid #dde4f5; background: #f3f6fc; color: #3b4c73;
    }
    .dl-chip.red { background: #fdecee; border-color: #f5c2c7; color: #a61e2b; }
    /* ============== 🏷️ हेडलाइन बैनर (st.title / subheader / ### सब जगह) ============== */
    .hb {
        color: #ffffff;
        border-radius: 14px;
        margin: 8px 0 14px 0;
        font-family: 'Segoe UI', 'Trebuchet MS', sans-serif;
        font-weight: 800;
        letter-spacing: -.2px;
        line-height: 1.35;
        box-shadow: 0 8px 22px rgba(26, 60, 110, .25);
        animation: floatIn .5s ease-out;
    }
    .hb1 { font-size: 25px; padding: 16px 22px; background: linear-gradient(135deg, #1a3c6e, #2f5fb3 60%, #5b4bdb); }
    .hb2 { font-size: 22px; padding: 13px 20px; background: linear-gradient(135deg, #1f4d8f, #3468c4 60%, #6a5be0); }
    .hb3 { font-size: 19px; padding: 11px 18px; background: linear-gradient(135deg, #2b5aa8, #4a6fd6 60%, #6a5be0); box-shadow: 0 6px 16px rgba(43, 90, 168, .22); }
    .hb4 {
        font-size: 16px; padding: 8px 14px; color: #1a3c6e;
        background: linear-gradient(90deg, #eef3ff, #ffffff);
        border-left: 6px solid #5b4bdb; border-radius: 10px; box-shadow: none;
    }
    .mini-head {
        border-left: 5px solid #1a73e8; background: linear-gradient(90deg, #eef4ff, #ffffff);
        border-radius: 10px; padding: 7px 14px; margin: 4px 0 8px 0;
        color: #1a3c6e; font-weight: 800; font-size: 15px; line-height: 1.4;
    }
    .mini-head span { color: #6b7688; font-weight: 500; font-size: 12.5px; }

    /* ============== ✍️ लिखा हुआ टेक्स्ट: पैराग्राफ, कैप्शन, अलर्ट कार्ड ============== */
    .txt {
        font-size: 15px; line-height: 1.75; color: #2f3b57;
        background: linear-gradient(90deg, #f3f6fd, #ffffff 70%);
        border: 1px solid #e3e8ef; border-left: 5px solid #5b4bdb; border-radius: 12px;
        padding: 10px 16px; margin: 4px 0 10px 0;
        box-shadow: 0 2px 8px rgba(26, 60, 110, .05);
    }
    .txt b, .cap b, .al-tx b { color: #1a3c6e; font-weight: 800; }
    .al .al-tx b { color: inherit; }
    .cap {
        font-size: 13.5px; line-height: 1.7; color: #55607a;
        background: #f8fafc; border-left: 4px solid #b9c9ec; border-radius: 8px;
        padding: 7px 12px; margin: 2px 0 8px 0;
    }
    .txt code, .cap code, .al code {
        background: rgba(26, 60, 110, .09); border-radius: 6px; padding: 1px 6px; font-size: 90%;
    }
    .al {
        display: flex; align-items: center; gap: 12px;
        border: 1px solid transparent; border-left-width: 6px; border-radius: 12px;
        padding: 11px 16px; margin: 8px 0;
        box-shadow: 0 3px 10px rgba(0, 0, 0, .06);
        animation: floatIn .4s ease-out;
    }
    .al .al-ic {
        flex: 0 0 auto; width: 28px; height: 28px; border-radius: 50%;
        display: flex; align-items: center; justify-content: center;
        font-weight: 800; font-size: 15px; color: #ffffff;
    }
    .al .al-tx { font-size: 14.5px; line-height: 1.6; font-weight: 600; }
    .al-success { background: linear-gradient(90deg, #e8f9ef, #ffffff 85%); border-color: #bfe8d0; border-left-color: #0f9d58; color: #0b6b3a; }
    .al-success .al-ic { background: linear-gradient(135deg, #2bc07a, #0b8f52); }
    .al-info { background: linear-gradient(90deg, #e8f1ff, #ffffff 85%); border-color: #c3d8fa; border-left-color: #1a73e8; color: #1a4fa0; }
    .al-info .al-ic { background: linear-gradient(135deg, #3b8cf0, #1a5fd0); }
    .al-warning { background: linear-gradient(90deg, #fff4de, #ffffff 85%); border-color: #f6dca4; border-left-color: #f59e0b; color: #8a5a00; }
    .al-warning .al-ic { background: linear-gradient(135deg, #fbbf24, #e08a00); }
    .al-error { background: linear-gradient(90deg, #fdeaec, #ffffff 85%); border-color: #f5c2c7; border-left-color: #d92d3a; color: #a61e2b; }
    .al-error .al-ic { background: linear-gradient(135deg, #ff6a5c, #d92d3a); }

    /* ============== 🧭 विजेट लेबल / कैप्शन / डिवाइडर / टैब बार ============== */
    div[data-testid="stWidgetLabel"] p, div[data-testid="stWidgetLabel"] label {
        color: #1a3c6e;
        font-weight: 700;
    }
    div[data-testid="stCaptionContainer"] {
        border-left: 3px solid #c7d6f5;
        padding-left: 10px;
        color: #55607a;
    }
    .block-container hr {
        border: none !important;
        height: 3px !important;
        border-radius: 3px;
        background: linear-gradient(90deg, #1a73e8, #5b4bdb 45%, rgba(91, 75, 219, 0)) !important;
        opacity: .55;
        margin: 1.4rem 0 !important;
    }
    div[data-baseweb="tab-list"] {
        gap: 6px;
        background: #f3f6fc;
        padding: 6px 8px 0 8px;
        border-radius: 14px 14px 0 0;
        border-bottom: 1px solid #dde4f5;
    }
    div[data-baseweb="tab-highlight"] {
        height: 4px !important;
        border-radius: 4px;
        background: linear-gradient(90deg, #1a73e8, #5b4bdb) !important;
    }
    button[data-baseweb="tab"] { font-size: 15px; }

    /* ============== 🏷️ (पुराना तरीका, बचा हुआ) हेडलाइन CSS ============== */
    .block-container div[data-testid="stHeading"] h1,
    .block-container div[data-testid="stMarkdownContainer"] > h1,
    .block-container div[data-testid="stHeading"] h2,
    .block-container div[data-testid="stMarkdownContainer"] > h2,
    .block-container div[data-testid="stHeading"] h3,
    .block-container div[data-testid="stMarkdownContainer"] > h3 {
        background: linear-gradient(135deg, #1a3c6e, #2f5fb3 60%, #5b4bdb);
        color: #ffffff !important;
        -webkit-text-fill-color: #ffffff;
        border-radius: 14px;
        padding: 13px 20px !important;
        margin: 6px 0 12px 0 !important;
        box-shadow: 0 8px 22px rgba(26, 60, 110, .25);
        font-family: 'Segoe UI', 'Trebuchet MS', sans-serif;
        font-weight: 800 !important;
        letter-spacing: -.2px;
        line-height: 1.35 !important;
        animation: floatIn .5s ease-out;
    }
    /* बैनर के अंदर की हर चीज़ (लिंक-आइकन समेत) सफ़ेद */
    .block-container div[data-testid="stHeading"] h1 *,
    .block-container div[data-testid="stMarkdownContainer"] > h1 *,
    .block-container div[data-testid="stHeading"] h2 *,
    .block-container div[data-testid="stMarkdownContainer"] > h2 *,
    .block-container div[data-testid="stHeading"] h3 *,
    .block-container div[data-testid="stMarkdownContainer"] > h3 * {
        color: #ffffff !important;
        -webkit-text-fill-color: #ffffff;
        fill: #ffffff;
    }
    /* बड़ा पैनल-टाइटल (st.title) */
    .block-container div[data-testid="stHeading"] h1,
    .block-container div[data-testid="stMarkdownContainer"] > h1 {
        font-size: 25px !important;
        padding: 16px 22px !important;
    }
    /* डैशबोर्ड बोर्ड आदि (##) : थोड़ा अलग नीला शेड */
    .block-container div[data-testid="stHeading"] h2,
    .block-container div[data-testid="stMarkdownContainer"] > h2 {
        font-size: 22px !important;
        background: linear-gradient(135deg, #1f4d8f, #3468c4 60%, #6a5be0);
    }
    /* सेक्शन हेडिंग (st.subheader / ###) */
    .block-container div[data-testid="stHeading"] h3,
    .block-container div[data-testid="stMarkdownContainer"] > h3 {
        font-size: 19px !important;
        padding: 11px 18px !important;
        background: linear-gradient(135deg, #2b5aa8, #4a6fd6 60%, #6a5be0);
        box-shadow: 0 6px 16px rgba(43, 90, 168, .22);
    }
    /* छोटी हेडिंग (####) : हल्का कार्ड + बाईं पट्टी */
    .block-container div[data-testid="stMarkdownContainer"] > h4 {
        background: linear-gradient(90deg, #eef3ff, #ffffff);
        border-left: 6px solid #5b4bdb;
        border-radius: 10px;
        padding: 8px 14px !important;
        margin: 8px 0 8px 0 !important;
        color: #1a3c6e;
        font-weight: 800;
    }

    /* ============== 📋 लिस्ट / टेबल कार्ड डिज़ाइन ============== */
    div[data-testid="stDataFrame"] {
        border: 1px solid #d7e0f0;
        border-radius: 14px;
        box-shadow: 0 4px 14px rgba(26, 60, 110, .08);
        transition: box-shadow .2s ease, border-color .2s ease;
    }
    div[data-testid="stDataFrame"]:hover {
        border-color: #9db7e8;
        box-shadow: 0 10px 26px rgba(26, 60, 110, .15);
    }
    .list-head {
        display: flex; align-items: center; gap: 10px; flex-wrap: wrap;
        margin: 10px 0 6px 0;
    }
    .list-head .lh-title { font-size: 16px; font-weight: 800; color: #1a3c6e; }
    .list-head .lh-count {
        font-size: 12.5px; font-weight: 700; color: #ffffff; padding: 3px 12px; border-radius: 999px;
        background: linear-gradient(135deg, #1a73e8, #5b4bdb);
    }
    /* 📝 कारण की लिस्ट: हर ब्रांच का अलग कार्ड */
    .reason-card {
        border: 1px solid #e3e8ef; border-left: 6px solid #6b7688; border-radius: 12px;
        background: #ffffff; padding: 10px 14px; margin: 8px 0;
        box-shadow: 0 2px 8px rgba(26, 60, 110, .06);
    }
    .reason-card.red { border-left-color: #d92d3a; background: linear-gradient(90deg, #fff5f6, #ffffff 40%); }
    .reason-card.blue { border-left-color: #17a2b8; background: linear-gradient(90deg, #f0fbfd, #ffffff 40%); }
    .reason-card .rc-title { font-weight: 800; color: #1a3c6e; font-size: 14.5px; margin-bottom: 6px; }
    .reason-card .rs-line {
        font-size: 13px; line-height: 1.6; padding: 4px 10px; border-radius: 8px; margin: 3px 0;
    }
    .reason-card .rs-line.red { background: #fdecee; color: #a61e2b; }
    .reason-card .rs-line.blue { background: #e3f6fa; color: #0c5460; }
    .dl-hint {
        background: #f8fafc; border: 1px solid #e3e8ef; border-left: 4px solid #0f9d58;
        border-radius: 10px; padding: 8px 14px; margin-bottom: 10px;
        color: #55607a; font-size: 13px; line-height: 1.6;
    }

    /* ============== 🎯 अर्थ के हिसाब से बटन रंग (पूरे कोड में) ============== */
    /* 🔴 डिलीट / रीसेट / Revoke = लाल */
    div[class*="st-key-danger_"] button,
    div[class*="st-key-revoke_btn_"] button {
        background: linear-gradient(135deg, #ff6a5c, #d92d3a) !important;
        color: #ffffff !important;
        border: none !important;
        box-shadow: 0 5px 16px rgba(217, 45, 58, .32) !important;
    }
    div[class*="st-key-danger_"] button p,
    div[class*="st-key-revoke_btn_"] button p { color: #ffffff !important; }
    div[class*="st-key-danger_"] button:hover,
    div[class*="st-key-revoke_btn_"] button:hover {
        box-shadow: 0 10px 22px rgba(217, 45, 58, .42) !important;
        filter: brightness(1.07);
    }
    /* ✅ Approve = हरा */
    div[class*="st-key-approve_"] button {
        background: linear-gradient(135deg, #2bc07a, #0b8f52) !important;
        color: #ffffff !important;
        border: none !important;
        box-shadow: 0 5px 16px rgba(15, 157, 88, .32) !important;
    }
    div[class*="st-key-approve_"] button p { color: #ffffff !important; }
    div[class*="st-key-approve_"] button:hover {
        box-shadow: 0 10px 22px rgba(15, 157, 88, .42) !important;
        filter: brightness(1.07);
    }
    /* 💾 Save = आसमानी-नीला */
    div[class*="st-key-save_"] button,
    div[class*="st-key-blank_exempt_save_"] button {
        background: linear-gradient(135deg, #0ea5e9, #2563eb) !important;
        color: #ffffff !important;
        border: none !important;
        box-shadow: 0 5px 16px rgba(37, 99, 235, .32) !important;
    }
    div[class*="st-key-save_"] button p,
    div[class*="st-key-blank_exempt_save_"] button p { color: #ffffff !important; }
    div[class*="st-key-save_"] button:hover,
    div[class*="st-key-blank_exempt_save_"] button:hover {
        box-shadow: 0 10px 22px rgba(37, 99, 235, .42) !important;
        filter: brightness(1.07);
    }
    /* 🔒 Lock = सुनहरा */
    div[class*="st-key-lock_"] button {
        background: linear-gradient(135deg, #f7b731, #e08a00) !important;
        color: #ffffff !important;
        border: none !important;
        box-shadow: 0 5px 16px rgba(224, 138, 0, .32) !important;
    }
    div[class*="st-key-lock_"] button p { color: #ffffff !important; }
    div[class*="st-key-lock_"] button:hover {
        box-shadow: 0 10px 22px rgba(224, 138, 0, .42) !important;
        filter: brightness(1.07);
    }
    /* 🔓 Logout = लाल आउटलाइन */
    div[class*="st-key-logout_btn"] button {
        background: #ffffff !important;
        border: 1.8px solid #e57373 !important;
        box-shadow: none !important;
    }
    div[class*="st-key-logout_btn"] button,
    div[class*="st-key-logout_btn"] button p { color: #c62828 !important; }
    div[class*="st-key-logout_btn"] button:hover {
        background: #fdecee !important;
        box-shadow: 0 8px 18px rgba(198, 40, 40, .2) !important;
    }
    /* बंद (disabled) बटन */
    div.stButton > button:disabled,
    div.stDownloadButton > button:disabled {
        opacity: .5 !important;
        cursor: not-allowed !important;
        transform: none !important;
        box-shadow: none !important;
        filter: grayscale(.4);
    }

    /* ============== 🔘 रेडियो = पिल (गोली) स्विच ============== */
    div[data-testid="stRadio"] div[role="radiogroup"] { gap: 8px; flex-wrap: wrap; }
    div[data-testid="stRadio"] label[data-baseweb="radio"] {
        background: #f3f6fc;
        border: 1.5px solid #cfd9ea;
        border-radius: 999px;
        padding: 8px 18px;
        margin: 0 !important;
        cursor: pointer;
        display: flex !important;
        align-items: center;
        transition: all .15s ease;
    }
    div[data-testid="stRadio"] label[data-baseweb="radio"] > div:first-of-type:not(:has([data-testid="stMarkdownContainer"])) {
        display: none !important;
    }
    div[data-testid="stRadio"] label[data-baseweb="radio"]:hover {
        border-color: #1a73e8;
        background: #e9f1ff;
        transform: translateY(-1px);
    }
    div[data-testid="stRadio"] label[data-baseweb="radio"]:has(input:checked) {
        background: linear-gradient(135deg, #1a73e8, #5b4bdb);
        border-color: transparent;
        box-shadow: 0 5px 14px rgba(91, 75, 219, .32);
    }
    div[data-testid="stRadio"] label[data-baseweb="radio"]:has(input:checked),
    div[data-testid="stRadio"] label[data-baseweb="radio"]:has(input:checked) * {
        color: #ffffff !important;
        font-weight: 700;
    }

    /* ============== ☑️ चेकबॉक्स / एक्सपैंडर / टैग / अपलोडर ============== */
    div[data-testid="stCheckbox"] label {
        padding: 6px 12px;
        border-radius: 10px;
        transition: background .15s ease;
    }
    div[data-testid="stCheckbox"] label:hover { background: #eef4ff; }
    div[data-testid="stExpander"] summary {
        font-weight: 700;
        border-radius: 10px;
        transition: background .15s ease;
    }
    div[data-testid="stExpander"] summary:hover { background: #f3f6fc; }
    span[data-baseweb="tag"] {
        background: linear-gradient(135deg, #1a73e8, #5b4bdb) !important;
        border-radius: 999px !important;
        color: #ffffff !important;
        font-weight: 600;
    }
    span[data-baseweb="tag"] span, span[data-baseweb="tag"] svg { color: #ffffff !important; fill: #ffffff !important; }
    section[data-testid="stFileUploaderDropzone"] {
        border: 2px dashed #9db7e8;
        border-radius: 14px;
        background: linear-gradient(180deg, #f8fbff, #eef4ff);
        transition: all .15s ease;
    }
    section[data-testid="stFileUploaderDropzone"]:hover {
        border-color: #1a73e8;
        background: #e9f1ff;
        box-shadow: 0 6px 18px rgba(26, 115, 232, .15);
    }
    section[data-testid="stFileUploaderDropzone"] button {
        border-radius: 10px;
        font-weight: 700;
        background: linear-gradient(135deg, #1a73e8, #5b4bdb);
        color: #ffffff;
        border: none;
        box-shadow: 0 4px 12px rgba(91, 75, 219, .3);
        transition: transform .15s ease, filter .15s ease;
    }
    section[data-testid="stFileUploaderDropzone"] button:hover {
        transform: translateY(-1px);
        filter: brightness(1.08);
        color: #ffffff;
    }
</style>
""", unsafe_allow_html=True)

def render_footer():
    st.markdown(
        "<div class='app-footer'>🛠️ Professionally Developed &amp; Maintained &nbsp;|&nbsp; "
        "NEP Master Data System © 2026</div>",
        unsafe_allow_html=True
    )

def inline_section_toggle(key, header_html, default_open=True):
    """किसी भी टेबल/सेक्शन के हेडर के साथ डिज़ाइनर Hide / Show बटन लगाता है।
    True = खुला (टेबल दिखाओ), False = छिपा।"""
    state_key = f"inline_sec_open_{key}"
    if state_key not in st.session_state:
        st.session_state[state_key] = default_open
    is_open = st.session_state[state_key]

    head_col, btn_col = st.columns([5, 1.6])
    with head_col:
        st.markdown(header_html, unsafe_allow_html=True)
    with btn_col:
        if is_open:
            btn_label, btn_key = "🙈 Hide करें", f"admin_sec_btn_{key}__hidebtn"
        else:
            btn_label, btn_key = "👁️ Show करें", f"admin_sec_btn_{key}__showbtn"
        if st.button(btn_label, key=btn_key, use_container_width=True):
            st.session_state[state_key] = not is_open
            st.rerun()

    if not is_open:
        st.markdown(
            "<div class='admin-sec-hidden-note'>🙈 यह समरी अभी छिपी हुई है — दिखाने के लिए दाईं ओर 'Show करें' दबाएँ।</div>",
            unsafe_allow_html=True
        )
    return is_open

def admin_section_toggle(key, title, caption=None, default_open=True):
    """Admin पैनल के हर सेक्शन के हेडर के साथ डिज़ाइनर Hide / Show बटन लगाता है।
    True लौटाए तो सेक्शन खुला है (कंटेंट दिखाओ), False हो तो छिपा है।"""
    state_key = f"admin_sec_open_{key}"
    if state_key not in st.session_state:
        st.session_state[state_key] = default_open
    is_open = st.session_state[state_key]

    head_col, btn_col = st.columns([5, 1.6])
    with head_col:
        st.subheader(title)
        if caption:
            st.caption(caption)
    with btn_col:
        st.markdown("<div style='height:10px'></div>", unsafe_allow_html=True)
        if is_open:
            btn_label, btn_key = "🙈 Hide करें", f"admin_sec_btn_{key}__hidebtn"
        else:
            btn_label, btn_key = "👁️ Show करें", f"admin_sec_btn_{key}__showbtn"
        if st.button(btn_label, key=btn_key, use_container_width=True):
            st.session_state[state_key] = not is_open
            st.rerun()

    if not is_open:
        st.markdown(
            "<div class='admin-sec-hidden-note'>🙈 यह सेक्शन अभी छिपा हुआ है — दिखाने के लिए दाईं ओर 'Show करें' दबाएँ।</div>",
            unsafe_allow_html=True
        )
    return is_open

# =========================================================================
# डेटाबेस सेटअप - टेबल्स संरचना (Raw, Permanent और Rules Lock)
# =========================================================================
conn = sqlite3.connect("nep_master_perma_db.db", check_same_thread=False)
cursor = conn.cursor()

# 1. अस्थायी स्टेजिंग स्टोरेज (Panel 1 से Upload होकर यहाँ आएगा)
cursor.execute("""
    CREATE TABLE IF NOT EXISTS raw_store (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        data_json TEXT
    )
""")

# 2. परमानेंट स्टोरेज (Panel 2 से Approve होकर UG/PG यहाँ आएगा)
cursor.execute("""
    CREATE TABLE IF NOT EXISTS perma_store (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        data_json TEXT,
        course_type TEXT
    )
""")

# 3. परमानेंट नियम लॉकिंग स्टोरेज (पुरानी खराब टेबल को डिलीट करके नया बनाने का ऑटो-सिस्टम)
try:
    # चेक करना कि क्या टेबल सही है
    cursor.execute("SELECT panel_prefix FROM locked_rules LIMIT 1")
except sqlite3.OperationalError:
    # अगर कोई भी गड़बड़ (जैसे कॉलम गायब होना) मिले, तो पुरानी टेबल हटा दें
    cursor.execute("DROP TABLE IF EXISTS locked_rules")
    conn.commit()

# अब बिल्कुल सही और नए स्ट्रक्चर के साथ टेबल बनाएं
cursor.execute("""
    CREATE TABLE IF NOT EXISTS locked_rules (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        panel_prefix TEXT UNIQUE,
        rules_json TEXT
    )
""")
conn.commit()

# सुनिश्चित करें कि टेबल बनने के बाद डेटाबेस में बदलाव सुरक्षित (Commit) हो जाएं
conn.commit()

# 4. पैनल-वाइज पासवर्ड + हाइड/अनहाइड स्टोरेज (6 पैनल्स के लिए)
cursor.execute("""
    CREATE TABLE IF NOT EXISTS panel_auth (
        panel_key TEXT PRIMARY KEY,
        password TEXT,
        hidden INTEGER DEFAULT 0
    )
""")
conn.commit()

# डिफ़ॉल्ट पासवर्ड (सिर्फ पहली बार, जब टेबल में डेटा न हो, तभी डाले जाएंगे)
_default_panel_passwords = {
    "p1": "op",       # पहले Operator का पासवर्ड
    "p2": "p2pass",
    "p3": "ug",       # पहले Teacher_UG का पासवर्ड
    "p4": "pg",       # पहले Teacher_PG का पासवर्ड
    "p5": "p5pass",
    "p6": "psv123",   # पहले Admin का पासवर्ड
}
for _pk, _pw in _default_panel_passwords.items():
    cursor.execute("INSERT OR IGNORE INTO panel_auth (panel_key, password, hidden) VALUES (?, ?, 0)", (_pk, _pw))
conn.commit()

# पैनल-की और उसके डिस्प्ले नाम की मैपिंग
PANEL_KEY_TO_NAME = {
    "p1": "📥 1. Entry / Upload Panel",
    "p2": "💻 2. Work / Approve Panel",
    "p3": "🎓 3. UG Panel",
    "p4": "📜 4. PG Panel",
    "p5": "📊 5. Dashboard / Counter Panel",
    "p6": "⚙️ 6. Admin Panel",
}
PANEL_NAME_TO_KEY = {v: k for k, v in PANEL_KEY_TO_NAME.items()}

def get_panel_password(panel_key):
    cursor.execute("SELECT password FROM panel_auth WHERE panel_key = ?", (panel_key,))
    row = cursor.fetchone()
    return row[0] if row else None

def is_panel_hidden(panel_key):
    cursor.execute("SELECT hidden FROM panel_auth WHERE panel_key = ?", (panel_key,))
    row = cursor.fetchone()
    return bool(row[0]) if row else False

def set_panel_password(panel_key, new_password):
    cursor.execute("UPDATE panel_auth SET password = ? WHERE panel_key = ?", (new_password, panel_key))
    conn.commit()

def set_panel_hidden(panel_key, hidden_flag):
    cursor.execute("UPDATE panel_auth SET hidden = ? WHERE panel_key = ?", (1 if hidden_flag else 0, panel_key))
    conn.commit()

# 5b. ✅ सब्जेक्ट अप्रूवल स्टोरेज (Panel 3/4 में "गलत विषय" को मैन्युअली Approve करने के लिए)
cursor.execute("""
    CREATE TABLE IF NOT EXISTS subject_approvals (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        student_key TEXT,
        panel_prefix TEXT,
        approved_by TEXT,
        approved_at TEXT,
        UNIQUE(student_key, panel_prefix)
    )
""")
conn.commit()

# 5. ऐप सेटिंग्स स्टोरेज (Login टाइटल + लोगो — सिर्फ Admin बदल सकता है)
cursor.execute("""
    CREATE TABLE IF NOT EXISTS app_settings (
        setting_key TEXT PRIMARY KEY,
        setting_value TEXT
    )
""")
conn.commit()

def get_app_setting(key, default=None):
    cursor.execute("SELECT setting_value FROM app_settings WHERE setting_key = ?", (key,))
    row = cursor.fetchone()
    return row[0] if row and row[0] is not None else default

def set_app_setting(key, value):
    cursor.execute("""
        INSERT INTO app_settings (setting_key, setting_value) VALUES (?, ?)
        ON CONFLICT(setting_key) DO UPDATE SET setting_value = excluded.setting_value
    """, (key, value))
    conn.commit()

def delete_app_setting(key):
    cursor.execute("DELETE FROM app_settings WHERE setting_key = ?", (key,))
    conn.commit()

# =========================================================================
# परमानेंट डेटा लोड करने का फंक्शन (perma_store से UG/PG डेटा पढ़ने के लिए)
# =========================================================================
def load_permanent_data(course_type):
    cursor.execute("SELECT data_json FROM perma_store WHERE course_type = ?", (course_type,))
    rows = cursor.fetchall()
    if not rows:
        return pd.DataFrame()
    all_records = []
    for (data_json,) in rows:
        try:
            records = json.loads(data_json)
            if isinstance(records, list):
                all_records.extend(records)
            else:
                all_records.append(records)
        except (json.JSONDecodeError, TypeError):
            continue
    if not all_records:
        return pd.DataFrame()
    return pd.DataFrame(all_records)

def load_raw_data():
    cursor.execute("SELECT data_json FROM raw_store ORDER BY id DESC LIMIT 1")
    row = cursor.fetchone()
    if not row or not row[0]:
        return pd.DataFrame()
    try:
        records = json.loads(row[0])
    except (json.JSONDecodeError, TypeError):
        return pd.DataFrame()
    if isinstance(records, dict):
        records = [records]
    if not records:
        return pd.DataFrame()
    return pd.DataFrame(records)

import datetime

# =========================================================================
# ✅ "गलत विषय" Approve सिस्टम — हेल्पर फंक्शन्स
# =========================================================================
def find_student_key_col(df):
    """छात्र की पहचान के लिए सबसे उपयुक्त कॉलम ढूंढना (Roll/Enrollment/Name)"""
    priority_keywords = ['roll', 'enrollment', 'enroll', 'admission', 'regn', 'registration', 'uid', 'scholar']
    for kw in priority_keywords:
        col = next((c for c in df.columns if kw in c.lower()), None)
        if col:
            return col
    return next((c for c in df.columns if 'name' in c.lower()), None)

def get_student_key(row, position, key_col):
    """हर छात्र के लिए एक यूनीक 'key' बनाना (approvals स्टोर करने के लिए)"""
    if key_col and key_col in row.index and not pd.isna(row[key_col]) and str(row[key_col]).strip():
        return str(row[key_col]).strip()
    return f"row-{position}"

def get_approval(student_key, prefix):
    cursor.execute(
        "SELECT approved_by, approved_at FROM subject_approvals WHERE student_key = ? AND panel_prefix = ?",
        (student_key, prefix)
    )
    return cursor.fetchone()

def add_approval(student_key, prefix, approved_by):
    cursor.execute("""
        INSERT INTO subject_approvals (student_key, panel_prefix, approved_by, approved_at)
        VALUES (?, ?, ?, ?)
        ON CONFLICT(student_key, panel_prefix) DO UPDATE SET
            approved_by = excluded.approved_by, approved_at = excluded.approved_at
    """, (student_key, prefix, approved_by, datetime.datetime.now().strftime("%d-%m-%Y %H:%M")))
    conn.commit()

def remove_approval(student_key, prefix):
    cursor.execute("DELETE FROM subject_approvals WHERE student_key = ? AND panel_prefix = ?", (student_key, prefix))
    conn.commit()

def get_all_approvals(prefix):
    cursor.execute(
        "SELECT student_key, approved_by, approved_at FROM subject_approvals WHERE panel_prefix = ? ORDER BY approved_at DESC",
        (prefix,)
    )
    return cursor.fetchall()

def render_print_button(df, title, button_label="🖨️ इस लिस्ट को A4 पर प्रिंट करें", key_suffix="", orientation="portrait"):
    """दिए गए DataFrame को A4-साइज़ प्रिंट-फ्रेंडली फॉर्मेट में एक नई विंडो में खोलकर सीधे प्रिंट डायलॉग खोलता है।
    orientation: 'portrait' या 'landscape' — यह तय करता है कि प्रिंट पेज सीधा (खड़ा) रहेगा या आड़ा।"""
    orientation = "landscape" if str(orientation).lower().startswith("land") else "portrait"
    _pdf = df.copy()
    if "क्र.सं." not in _pdf.columns and "क्र." not in _pdf.columns:
        _pdf.insert(0, "क्र.", range(1, len(_pdf) + 1))
    table_html = _pdf.to_html(index=False, escape=True, border=0)
    _orient_txt = "लैंडस्केप" if orientation == "landscape" else "पोर्ट्रेट"
    full_html = f"""
    <html>
    <head>
    <meta charset="utf-8">
    <title>{title}</title>
    <style>
        @page {{ size: A4 {orientation}; margin: 14mm; }}
        * {{ -webkit-print-color-adjust: exact; print-color-adjust: exact; }}
        body {{ font-family: 'Segoe UI', Arial, sans-serif; color: #1a1a1a; }}
        .hdr {{ text-align: center; border-bottom: 3px solid #1a3c6e; padding-bottom: 8px; margin-bottom: 12px; }}
        h2 {{ color: #1a3c6e; margin: 0 0 4px 0; font-size: 20px; letter-spacing: -.2px; }}
        p.meta {{ color: #55607a; font-size: 12px; margin: 0; }}
        p.meta b {{ color: #1a3c6e; }}
        table {{ width: 100%; border-collapse: collapse; border: 1px solid #9db0d3; }}
        thead {{ display: table-header-group; }}
        tr {{ page-break-inside: avoid; }}
        th, td {{ border: 1px solid #c3cee6; padding: 6px 8px; font-size: 12px; text-align: left; word-break: break-word; }}
        th {{ background-color: #1a3c6e; color: #ffffff; font-weight: 700; border-color: #1a3c6e; }}
        tbody tr:nth-child(even) {{ background-color: #f1f5fc; }}
        td:first-child, th:first-child {{ text-align: center; width: 4%; }}
        .foot {{ margin-top: 10px; font-size: 11.5px; color: #55607a; text-align: right; }}
    </style>
    </head>
    <body>
        <div class="hdr">
            <h2>{title}</h2>
            <p class="meta">तारीख़: {datetime.datetime.now().strftime('%d-%m-%Y %H:%M')} &nbsp;|&nbsp; पेज: A4 ({_orient_txt}) &nbsp;|&nbsp; <b>कुल रिकॉर्ड: {len(_pdf)}</b></p>
        </div>
        {table_html}
        <div class="foot">— लिस्ट समाप्त • कुल {len(_pdf)} रिकॉर्ड —</div>
    </body>
    </html>
    """
    escaped_json = json.dumps(full_html)
    btn_html = f"""
    <style>
        body {{ margin: 0; padding: 4px 2px; background: transparent; }}
        .pbtn {{
            display: inline-flex; align-items: center; justify-content: center; gap: 8px;
            background: linear-gradient(135deg, #1a73e8, #5b4bdb);
            color: #ffffff; border: none; border-radius: 12px;
            padding: 11px 22px; min-height: 44px; box-sizing: border-box;
            font-family: 'Segoe UI', 'Trebuchet MS', sans-serif;
            font-weight: 700; font-size: 14.5px; letter-spacing: .2px;
            cursor: pointer;
            box-shadow: 0 5px 16px rgba(91, 75, 219, .35);
            transition: transform .15s ease, box-shadow .15s ease, filter .15s ease;
        }}
        .pbtn:hover {{ transform: translateY(-2px); filter: brightness(1.08); box-shadow: 0 10px 22px rgba(91, 75, 219, .42); }}
        .pbtn:active {{ transform: translateY(0) scale(.98); box-shadow: 0 1px 4px rgba(0,0,0,.2); }}
        .pbtn:focus-visible {{ outline: 3px solid rgba(26, 115, 232, .35); outline-offset: 2px; }}
    </style>
    <button id="printBtn_{key_suffix}" class="pbtn">
        {button_label}
    </button>
    <script>
        document.getElementById("printBtn_{key_suffix}").onclick = function() {{
            var content = {escaped_json};
            var w = window.open('', '_blank');
            w.document.write(content);
            w.document.close();
            w.focus();
            setTimeout(function() {{ w.print(); }}, 300);
        }};
    </script>
    """
    components.html(btn_html, height=64)

# Session States Management
if "ok" not in st.session_state: st.session_state["ok"] = False
if "deleted_cols" not in st.session_state: st.session_state["deleted_cols"] = []

# --- LOGIN SYSTEM (पैनल-वाइज: हर पैनल का अपना पासवर्ड) ---
# 🔧 फिक्स: पुराने सेशन (जिसमें "ok"=True था लेकिन "panel" key नहीं थी) की वजह से
# KeyError न आए, इसलिए दोनों चीज़ें एक साथ चेक कर रहे हैं
if not st.session_state["ok"] or "panel" not in st.session_state:
    st.session_state["ok"] = False

    left, mid, right = st.columns([1, 1.3, 1])
    with mid:
        # 👋 अगर अभी-अभी Logout किया है, तो एक प्यारा-सा फेयरवेल मैसेज दिखाएं
        if st.session_state.pop("show_farewell", False):
            st.markdown(
                "<div class='farewell-box'>👋 सफलतापूर्वक Logout हो गए! फिर मिलते हैं 😊</div>",
                unsafe_allow_html=True
            )

        # 🎨 Admin द्वारा सेट किया गया टाइटल और लोगो (डिफ़ॉल्ट: इमोजी + "NEP Master Data System")
        _login_title = get_app_setting("login_title", "NEP Master Data System")
        _login_subtitle = get_app_setting("login_subtitle", "अपना पैनल चुनें और आगे बढ़ने के लिए पासवर्ड डालें")
        _logo_b64 = get_app_setting("login_logo_b64")
        _logo_mime = get_app_setting("login_logo_mime", "image/png")
        _logo_width = int(get_app_setting("login_logo_width", "140"))
        _logo_height = int(get_app_setting("login_logo_height", "140"))
        _logo_fit = get_app_setting("login_logo_fit", "contain")  # "contain" = पूरी image दिखेगी (कटेगी नहीं), "cover" = बॉक्स भरेगी (क्रॉप हो सकती है)

        if _logo_b64:
            badge_html = (
                f'<img src="data:{_logo_mime};base64,{_logo_b64}" class="login-logo-img" '
                f'style="width:{_logo_width}px; height:{_logo_height}px; object-fit:{_logo_fit};" />'
            )
        else:
            badge_html = '<div class="emoji-badge">🎓🔒</div>'

        st.markdown(
            f"""
            <div class="login-hero">
                {badge_html}
                <h1>{_login_title}</h1>
                <p>{_login_subtitle}</p>
            </div>
            """,
            unsafe_allow_html=True
        )
        st.markdown(
            """
            <div class="login-badge-row">
                <span class="login-badge">📥 Entry</span>
                <span class="login-badge">💻 Approve</span>
                <span class="login-badge">🎓 UG</span>
                <span class="login-badge">📜 PG</span>
                <span class="login-badge">📊 Dashboard</span>
                <span class="login-badge">⚙️ Admin</span>
            </div>
            """,
            unsafe_allow_html=True
        )

        with st.container(border=True):
            panel_choice_label = st.selectbox("🗂️ पैनल चुनें:", ["-- चुनें --"] + list(PANEL_KEY_TO_NAME.values()))
            pas = st.text_input("🔑 Password:", type="password", placeholder="अपना पासवर्ड यहाँ डालें")
            login_clicked = st.button("🚀 Login करें", use_container_width=True, type="primary")

            if login_clicked:
                if panel_choice_label == "-- चुनें --":
                    st.error("⚠️ कृपया पहले एक पैनल चुनें।")
                else:
                    selected_key = PANEL_NAME_TO_KEY[panel_choice_label]
                    correct_pw = get_panel_password(selected_key)
                    if correct_pw is not None and pas == correct_pw:
                        st.session_state["ok"] = True
                        st.session_state["panel_key"] = selected_key
                        st.session_state["panel"] = panel_choice_label
                        st.success(f"🎉 स्वागत है! {panel_choice_label} में लॉगिन हो रहे हैं...")
                        st.balloons()
                        st.rerun()
                    else:
                        st.error("❌ गलत पासवर्ड! कृपया सही पासवर्ड डालें।")

        st.markdown(
            "<p style='text-align:center; color:#a3adbd; font-size:12px; margin-top:10px;'>"
            "🔐 आपका डेटा सुरक्षित है — हर पैनल का अपना अलग पासवर्ड है</p>",
            unsafe_allow_html=True
        )
    st.stop()

# =========================================================================
# 🔄 लॉगिन किए गए पैनल को लोड करना
# =========================================================================
panel = st.session_state["panel"]
panel_key = st.session_state.get("panel_key")

st.sidebar.markdown(
    f"""
    <div style="background: linear-gradient(135deg, #eef6ff, #f3fff5); border: 1px solid #d7e6f9;
                border-radius: 12px; padding: 12px 14px; margin-bottom: 10px;">
        <div style="font-size:11px; color:#7a869f; font-weight:600; letter-spacing:0.5px;">लॉगिन पैनल</div>
        <div style="font-size:15px; font-weight:700; color:#1a3c6e; margin-top:2px;">{panel}</div>
    </div>
    """,
    unsafe_allow_html=True
)
if st.sidebar.button("🔓 Logout करें", use_container_width=True, key="logout_btn"):
    st.session_state["ok"] = False
    st.session_state.pop("panel", None)
    st.session_state.pop("panel_key", None)
    st.session_state["show_farewell"] = True
    st.rerun()

# 👁️ सिर्फ Admin (P6 लॉगिन) के लिए: बाकी सभी पैनल्स (P1-P5) को भी देखने का विकल्प
is_admin_session = (panel_key == "p6")
if is_admin_session:
    st.sidebar.divider()
    admin_view_choice = st.sidebar.selectbox(
        "👁️ पैनल देखें (Admin View):",
        ["⚙️ 6. Admin Panel"] + [PANEL_KEY_TO_NAME[k] for k in ["p1", "p2", "p3", "p4", "p5"]],
        key="admin_view_selector"
    )
    active_panel = admin_view_choice
else:
    active_panel = panel

import io
from openpyxl.styles import PatternFill, Border, Side

# =========================================================================
# 🔀 UG / PG रूटिंग: कौन सी डिग्री किस डेटाबेस (UG या PG) में जाएगी
# =========================================================================
# P2 का ऑटो-विभाजन और P4 (PG Panel) का फ़िल्टर — दोनों यही एक सूची इस्तेमाल करते हैं, ताकि कोई डिग्री
# P2 में PG में तो जाए पर P4 में दिखे ही नहीं (या उल्टा) — ऐसा न हो। (LL.M. और M.Tech पहले इस सूची में नहीं थे।)
PG_ROUTE_KEYWORDS = ["ma", "msc", "mcom", "mba", "mca", "post grad", "pg", "grad", "llm", "mtech"]
# P3 (UG Panel) में जो डिग्री दिखाई जाती हैं
UG_ALLOWED_KEYWORDS = ["ba", "bsc", "bcom", "bhsc", "bba", "bca", "computer"]

def is_pg_route_value(v):
    val = str(v).lower().replace(".", "").replace(" ", "").strip()
    return any(k in val for k in PG_ROUTE_KEYWORDS)

# =========================================================================
# ⚪ खाली-छूट (Blank Exemption): जिन डिग्री+ब्रांच में Minor/MDC/Voc/PW हमेशा खाली रहते हैं
# =========================================================================
def _norm_blank_label(v):
    s_ = "" if pd.isna(v) else str(v).strip()
    return s_ if (s_ and s_.lower() != "nan") else "(खाली/Blank)"

def detect_deg_branch_cols(df):
    """डिग्री और ब्रांच कॉलम पहचानना (Minor/MDC/Voc/PW वाले कॉलम को ब्रांच न मानना)।"""
    skip_kw = ['minor', 'mdc', 'voc', 'skill', 'pw', 'project']
    cols = list(df.columns)
    deg = next((c for c in cols if any(k in str(c).lower() for k in ['deg', 'course', 'class'])), None)
    br = None
    for kws in (['branch', 'stream'], ['subject']):
        br = next((c for c in cols if c != deg and any(k in str(c).lower() for k in kws)
                   and not any(k in str(c).lower() for k in skip_kw)), None)
        if br is not None:
            break
    return deg, br

def get_blank_exempt_pairs(prefix=None):
    """सेव की हुई (डिग्री, ब्रांच) जोड़ियाँ। prefix=None हो तो UG + PG दोनों की मिलाकर।"""
    pairs = set()
    for pf in ([prefix] if prefix else ["ug", "pg"]):
        raw = get_app_setting(f"blank_exempt_{pf}", "[]")
        try:
            for item in json.loads(raw):
                if isinstance(item, (list, tuple)) and len(item) == 2:
                    pairs.add((str(item[0]), str(item[1])))
        except Exception:
            pass
    return pairs

def set_blank_exempt_pairs(prefix, pairs):
    set_app_setting(f"blank_exempt_{prefix}", json.dumps(sorted([list(p) for p in pairs]), ensure_ascii=False))

def blank_exempt_mask(df, prefix=None):
    """हर रो के लिए True/False: क्या इस रो की डिग्री+ब्रांच 'खाली-छूट' में है (तो Minor..PW का खाली नीला नहीं होगा)।"""
    n = len(df)
    pairs = get_blank_exempt_pairs(prefix)
    if n == 0 or not pairs:
        return pd.Series([False] * n, index=df.index)
    dc, bc = detect_deg_branch_cols(df)
    deg_lbl = df[dc].map(_norm_blank_label) if dc else pd.Series(["—"] * n, index=df.index)
    br_lbl = df[bc].map(_norm_blank_label) if bc else pd.Series(["—"] * n, index=df.index)
    return pd.Series([(d, b) in pairs for d, b in zip(deg_lbl, br_lbl)], index=df.index)

def generate_colored_excel_bytes(df_filtered, deg_col, br_col, minor_col_found, mdc_col_found, voc_col_found, pw_col_found, master_rules=None, sheet_name="Verified_Data", approved_keys=None, key_col=None, prefix=None):
    """
    🔧 रीयूज़ेबल फ़ंक्शन: किसी भी DataFrame को रंगीन (🔴 गलत / 🔵 खाली / 🟢 Approved) Excel bytes में बदलता है।
    Panel 3/4 और Admin Panel — दोनों जगह इसी फ़ंक्शन का इस्तेमाल होता है, ताकि रंग-कोडिंग हमेशा एक जैसी रहे।
    approved_keys: उन छात्रों की student_key की list/set, जिन्हें गलत विषय होने के बावजूद Approve किया जा चुका है (उन्हें लाल नहीं, हरा दिखाया जाएगा)।
    """
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df_filtered.to_excel(writer, index=False, sheet_name=sheet_name)
        workbook = writer.book
        worksheet = writer.sheets[sheet_name]

        blue_fill = PatternFill(start_color="D1ECF1", end_color="D1ECF1", fill_type="solid")   # ब्लैंक = नीला
        red_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")     # गलत = लाल
        green_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")   # Approved = हरा
        thin_border = Border(left=Side(style='thin', color='CCCCCC'), right=Side(style='thin', color='CCCCCC'),
                             top=Side(style='thin', color='CCCCCC'), bottom=Side(style='thin', color='CCCCCC'))

        targets_xl = {minor_col_found: 'minor', mdc_col_found: 'mdc', voc_col_found: 'voc', pw_col_found: 'pw'}
        approved_keys = approved_keys or set()
        _exempt_list = blank_exempt_mask(df_filtered, prefix).tolist()   # ⚪ खाली-छूट वाली डिग्री+ब्रांच

        for idx, (_, row) in enumerate(df_filtered.iterrows()):
            row_num = idx + 2  # एक्सेल डेटा रो

            deg_part = str(row[deg_col]) if deg_col and deg_col in df_filtered.columns else ""
            br_part = str(row[br_col]) if br_col and br_col in df_filtered.columns else ""
            student_deg = (deg_part + " " + br_part).lower().replace(".", "").replace(" ", "").strip()

            is_row_approved = get_student_key(row, idx, key_col) in approved_keys

            matched_key = "Default"
            if master_rules:
                sorted_keys = sorted(master_rules.keys(), key=len, reverse=True)
                for rule_key in sorted_keys:
                    rule_words = [w.lower().replace(".", "").strip() for w in rule_key.split() if w.strip()]
                    if rule_words and all(w in student_deg for w in rule_words):
                        matched_key = rule_key
                        break
            c_rule = master_rules.get(matched_key, {"minor": [], "mdc": [], "voc": [], "pw": []}) if master_rules else {"minor": [], "mdc": [], "voc": [], "pw": []}

            for col_idx, col_name in enumerate(df_filtered.columns, start=1):
                cell = worksheet.cell(row=row_num, column=col_idx)
                val = row[col_name]

                if col_name == br_col:
                    if pd.isna(val) or str(val).strip() == "":
                        cell.fill = blue_fill
                        cell.border = thin_border
                    continue

                if col_name in targets_xl:
                    rule_key = targets_xl[col_name]

                    if pd.isna(val) or str(val).strip() == "":
                        if not _exempt_list[idx]:
                            cell.fill = blue_fill
                            cell.border = thin_border
                    else:
                        val_clean = str(val).strip().lower().replace(".", "").replace(" ", "")
                        valid_list = c_rule.get(rule_key, [])
                        valid_set = {str(x).strip().lower().replace(".", "").replace(" ", "") for x in valid_list}

                        if valid_set and (val_clean not in valid_set):
                            cell.fill = green_fill if is_row_approved else red_fill
                            cell.border = thin_border

    return output.getvalue()

def generate_master_excel_bytes(sheets_data):
    """
    🆕 P5 Dashboard के 'मास्टर एक्सेल डाउनलोड' बटन के लिए रीयूज़ेबल फ़ंक्शन।
    एक ही Excel फ़ाइल में कई (UG/PG) कलर-कोडेड शीट्स एक साथ बनाता है, ताकि P5 पर बनी
    हर लिस्ट (चाहे कोई भी डिग्री/टैब हो) एक ही मास्टर फ़ाइल में मिल जाए।
    sheets_data: dict की list, हर dict में ये keys हो सकती हैं:
      df, sheet_name, deg_col, br_col, minor_col, mdc_col, voc_col, pw_col,
      master_rules, approved_keys, key_col, prefix
    """
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        for sheet in sheets_data:
            df_filtered = sheet.get("df")
            sheet_name = sheet.get("sheet_name", "Sheet1")
            if df_filtered is None or df_filtered.empty:
                continue

            deg_col = sheet.get("deg_col")
            br_col = sheet.get("br_col")
            minor_col_found = sheet.get("minor_col")
            mdc_col_found = sheet.get("mdc_col")
            voc_col_found = sheet.get("voc_col")
            pw_col_found = sheet.get("pw_col")
            master_rules = sheet.get("master_rules")
            approved_keys = sheet.get("approved_keys") or set()
            key_col = sheet.get("key_col")
            prefix = sheet.get("prefix")

            df_filtered.to_excel(writer, index=False, sheet_name=sheet_name)
            worksheet = writer.sheets[sheet_name]

            blue_fill = PatternFill(start_color="D1ECF1", end_color="D1ECF1", fill_type="solid")
            red_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
            green_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
            thin_border = Border(left=Side(style='thin', color='CCCCCC'), right=Side(style='thin', color='CCCCCC'),
                                 top=Side(style='thin', color='CCCCCC'), bottom=Side(style='thin', color='CCCCCC'))

            targets_xl = {minor_col_found: 'minor', mdc_col_found: 'mdc', voc_col_found: 'voc', pw_col_found: 'pw'}
            _exempt_list = blank_exempt_mask(df_filtered, prefix).tolist()

            for idx, (_, row) in enumerate(df_filtered.iterrows()):
                row_num = idx + 2

                deg_part = str(row[deg_col]) if deg_col and deg_col in df_filtered.columns else ""
                br_part = str(row[br_col]) if br_col and br_col in df_filtered.columns else ""
                student_deg = (deg_part + " " + br_part).lower().replace(".", "").replace(" ", "").strip()

                is_row_approved = get_student_key(row, idx, key_col) in approved_keys

                matched_key = "Default"
                if master_rules:
                    sorted_keys = sorted(master_rules.keys(), key=len, reverse=True)
                    for rule_key in sorted_keys:
                        rule_words = [w.lower().replace(".", "").strip() for w in rule_key.split() if w.strip()]
                        if rule_words and all(w in student_deg for w in rule_words):
                            matched_key = rule_key
                            break
                c_rule = master_rules.get(matched_key, {"minor": [], "mdc": [], "voc": [], "pw": []}) if master_rules else {"minor": [], "mdc": [], "voc": [], "pw": []}

                for col_idx, col_name in enumerate(df_filtered.columns, start=1):
                    cell = worksheet.cell(row=row_num, column=col_idx)
                    val = row[col_name]

                    if col_name == br_col:
                        if pd.isna(val) or str(val).strip() == "":
                            cell.fill = blue_fill
                            cell.border = thin_border
                        continue

                    if col_name in targets_xl:
                        rule_key = targets_xl[col_name]

                        if pd.isna(val) or str(val).strip() == "":
                            if not _exempt_list[idx]:
                                cell.fill = blue_fill
                                cell.border = thin_border
                        else:
                            val_clean = str(val).strip().lower().replace(".", "").replace(" ", "")
                            valid_list = c_rule.get(rule_key, [])
                            valid_set = {str(x).strip().lower().replace(".", "").replace(" ", "") for x in valid_list}

                            if valid_set and (val_clean not in valid_set):
                                cell.fill = green_fill if is_row_approved else red_fill
                                cell.border = thin_border

    return output.getvalue()

def generate_p5_master_full_excel(raw_sheets, cat_labels, ug_blocks, pg_blocks,
                                   ug_summary, pg_summary,
                                   ug_students, ug_flags, ug_approved,
                                   pg_students, pg_flags, pg_approved):
    """
    🆕 P5 Dashboard के 'मास्टर एक्सेल डाउनलोड' बटन के लिए पूरी 6-शीट Excel फ़ाइल बनाता है:
      Sheet 1: UG_Master_List        — पूरा UG रॉ डेटा (रंगीन)
      Sheet 2: PG_Master_List        — पूरा PG रॉ डेटा (रंगीन)
      Sheet 3: Branch_Subject_Sheet  — ब्रांच-वाइज विषय शीट (Total Admission + Minor/MDC/Voc/PW नाम व संख्या, merged cells)
      Sheet 4: Degree_Branch_Summary — डिग्री+ब्रांच-वाइज समरी (Minor+MDC+Voc+PW सभी एक साथ, UG फिर PG)
      Sheet 5: Reason_Students       — जिन छात्रों का कोई विषय गलत/खाली है, उनकी पूरी लिस्ट + कारण
      Sheet 6: Approved_Students     — शीट 5 में से जो पहले ही Approve किए जा चुके हैं, उनकी लिस्ट
    raw_sheets: sheet 1/2 के लिए dict की list (generate_master_excel_bytes जैसा फ़ॉर्मेट)।
    """
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:

        # ================== Sheet 1 & 2: रॉ डेटा (रंगीन) ==================
        for sheet in raw_sheets:
            df_filtered = sheet.get("df")
            sheet_name = sheet.get("sheet_name", "Sheet1")
            if df_filtered is None or df_filtered.empty:
                continue

            s_deg_col = sheet.get("deg_col")
            s_br_col = sheet.get("br_col")
            minor_col_found = sheet.get("minor_col")
            mdc_col_found = sheet.get("mdc_col")
            voc_col_found = sheet.get("voc_col")
            pw_col_found = sheet.get("pw_col")
            master_rules = sheet.get("master_rules")
            approved_keys = sheet.get("approved_keys") or set()
            key_col = sheet.get("key_col")
            prefix = sheet.get("prefix")

            df_filtered.to_excel(writer, index=False, sheet_name=sheet_name)
            worksheet = writer.sheets[sheet_name]

            blue_fill = PatternFill(start_color="D1ECF1", end_color="D1ECF1", fill_type="solid")
            red_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
            green_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
            thin_border = Border(left=Side(style='thin', color='CCCCCC'), right=Side(style='thin', color='CCCCCC'),
                                 top=Side(style='thin', color='CCCCCC'), bottom=Side(style='thin', color='CCCCCC'))

            targets_xl = {minor_col_found: 'minor', mdc_col_found: 'mdc', voc_col_found: 'voc', pw_col_found: 'pw'}
            _exempt_list = blank_exempt_mask(df_filtered, prefix).tolist()

            for idx, (_, row) in enumerate(df_filtered.iterrows()):
                row_num = idx + 2

                deg_part = str(row[s_deg_col]) if s_deg_col and s_deg_col in df_filtered.columns else ""
                br_part = str(row[s_br_col]) if s_br_col and s_br_col in df_filtered.columns else ""
                student_deg = (deg_part + " " + br_part).lower().replace(".", "").replace(" ", "").strip()

                is_row_approved = get_student_key(row, idx, key_col) in approved_keys

                matched_key = "Default"
                if master_rules:
                    sorted_keys = sorted(master_rules.keys(), key=len, reverse=True)
                    for rule_key in sorted_keys:
                        rule_words = [w.lower().replace(".", "").strip() for w in rule_key.split() if w.strip()]
                        if rule_words and all(w in student_deg for w in rule_words):
                            matched_key = rule_key
                            break
                c_rule = master_rules.get(matched_key, {"minor": [], "mdc": [], "voc": [], "pw": []}) if master_rules else {"minor": [], "mdc": [], "voc": [], "pw": []}

                for col_idx, col_name in enumerate(df_filtered.columns, start=1):
                    cell = worksheet.cell(row=row_num, column=col_idx)
                    val = row[col_name]

                    if col_name == s_br_col:
                        if pd.isna(val) or str(val).strip() == "":
                            cell.fill = blue_fill
                            cell.border = thin_border
                        continue

                    if col_name in targets_xl:
                        rule_key = targets_xl[col_name]

                        if pd.isna(val) or str(val).strip() == "":
                            if not _exempt_list[idx]:
                                cell.fill = blue_fill
                                cell.border = thin_border
                        else:
                            val_clean = str(val).strip().lower().replace(".", "").replace(" ", "")
                            valid_list = c_rule.get(rule_key, [])
                            valid_set = {str(x).strip().lower().replace(".", "").replace(" ", "") for x in valid_list}

                            if valid_set and (val_clean not in valid_set):
                                cell.fill = green_fill if is_row_approved else red_fill
                                cell.border = thin_border

        # ================== कॉमन स्टाइल ==================
        thin = Side(style="thin", color="BFBFBF")
        border3 = Border(left=thin, right=thin, top=thin, bottom=thin)
        head_fill3 = PatternFill("solid", start_color="1A3C6E", end_color="1A3C6E")
        band_a3 = PatternFill("solid", start_color="EAF1FB", end_color="EAF1FB")
        band_b3 = PatternFill("solid", start_color="FFFFFF", end_color="FFFFFF")
        red_fill3 = PatternFill("solid", start_color="F8D7DA", end_color="F8D7DA")
        blue_fill3 = PatternFill("solid", start_color="D1ECF1", end_color="D1ECF1")
        zebra_fill3 = PatternFill("solid", start_color="F4F7FD", end_color="F4F7FD")
        green_head_fill3 = PatternFill("solid", start_color="0F9D58", end_color="0F9D58")
        green_fill3 = PatternFill("solid", start_color="D4EDDA", end_color="D4EDDA")

        # ================== Sheet 3: ब्रांच-वाइज विषय शीट (UG + PG मर्ज) ==================
        ws3 = writer.book.create_sheet("Branch_Subject_Sheet")
        headers3 = ["Type", "Degree (डिग्री)", "Branch (ब्रांच)", "Total Admission"]
        for l_ in cat_labels:
            headers3 += [f"{l_} (विषय)", f"{l_} Count"]
        for c_i, t_ in enumerate(headers3, start=1):
            cell = ws3.cell(row=1, column=c_i, value=t_)
            cell.fill = head_fill3
            cell.font = Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border3

        r3 = 2
        all_blocks3 = [("UG", b_) for b_ in ug_blocks] + [("PG", b_) for b_ in pg_blocks]
        for bi, (typ_, b_) in enumerate(all_blocks3):
            start, end = r3, r3 + b_["n"] - 1
            fill = band_a3 if bi % 2 == 0 else band_b3
            for rr in range(start, end + 1):
                for cc in range(1, len(headers3) + 1):
                    cell = ws3.cell(row=rr, column=cc)
                    cell.fill = fill
                    cell.border = border3
            for cc, val in ((1, typ_), (2, b_["degree"]), (3, b_["branch"]), (4, b_["total"])):
                cell = ws3.cell(row=start, column=cc, value=val)
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                if end > start:
                    ws3.merge_cells(start_row=start, start_column=cc, end_row=end, end_column=cc)
            for k, items in enumerate(b_["cats"]):
                for i, (name_, cnt_) in enumerate(items):
                    ws3.cell(row=start + i, column=5 + 2 * k, value=name_).alignment = Alignment(vertical="center", wrap_text=True)
                    ws3.cell(row=start + i, column=6 + 2 * k, value=cnt_).alignment = Alignment(horizontal="center", vertical="center")
            r3 = end + 1

        widths3 = [10, 22, 28, 16] + [34, 12] * len(cat_labels)
        for c_i, w_ in enumerate(widths3, start=1):
            ws3.column_dimensions[get_column_letter(c_i)].width = w_
        ws3.freeze_panes = "A2"
        if r3 == 2:
            ws3.cell(row=2, column=1, value="कोई डेटा उपलब्ध नहीं है।")

        # ================== Sheet 4: डिग्री + ब्रांच-वाइज समरी (UG फिर PG) ==================
        def _write_summary_section(ws, start_row, type_label, summary_df):
            if summary_df is None or summary_df.empty:
                return start_row
            cols4 = list(summary_df.columns)
            red_idx4 = [i for i, c in enumerate(cols4) if "🔴" in str(c)]
            blue_idx4 = [i for i, c in enumerate(cols4) if "🔵" in str(c)]
            reason_idx4 = next((i for i, c in enumerate(cols4) if "Reason" in str(c)), None)

            headers4 = ["Type"] + cols4
            for c_i, t_ in enumerate(headers4, start=1):
                cell = ws.cell(row=start_row, column=c_i, value=t_)
                cell.fill = head_fill3
                cell.font = Font(bold=True, color="FFFFFF")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = border3

            r_ = start_row + 1
            for _, row in summary_df.iterrows():
                is_total = str(row.iloc[0]).startswith("कुल योग")
                cell0 = ws.cell(row=r_, column=1, value=type_label)
                cell0.border = border3
                if is_total:
                    cell0.fill = head_fill3
                    cell0.font = Font(bold=True, color="FFFFFF")
                any_w = (not is_total) and any((row.iloc[i] or 0) > 0 for i in red_idx4)
                any_b = (not is_total) and any((row.iloc[i] or 0) > 0 for i in blue_idx4)
                for c_i0, cname in enumerate(cols4):
                    val = row.iloc[c_i0]
                    if c_i0 == reason_idx4 and isinstance(val, str):
                        val = val.replace("  ||  ", "\n")
                    cell = ws.cell(row=r_, column=c_i0 + 2, value=val)
                    cell.border = border3
                    cell.alignment = Alignment(vertical="top", wrap_text=(c_i0 == reason_idx4))
                    if is_total:
                        cell.fill = head_fill3
                        cell.font = Font(bold=True, color="FFFFFF")
                    elif c_i0 in red_idx4 and isinstance(val, (int, float)) and val > 0:
                        cell.fill = red_fill3
                        cell.font = Font(bold=True, color="721C24")
                    elif c_i0 in blue_idx4 and isinstance(val, (int, float)) and val > 0:
                        cell.fill = blue_fill3
                        cell.font = Font(bold=True, color="0C5460")
                    elif c_i0 == reason_idx4 and str(val).strip():
                        if any_w:
                            cell.fill = red_fill3
                            cell.font = Font(bold=True, color="721C24")
                        elif any_b:
                            cell.fill = blue_fill3
                            cell.font = Font(bold=True, color="0C5460")
                r_ += 1
            return r_ + 1   # अगली टेबल से पहले एक खाली रो

        ws4 = writer.book.create_sheet("Degree_Branch_Summary")
        _r4 = _write_summary_section(ws4, 1, "UG", ug_summary)
        _r4 = _write_summary_section(ws4, _r4, "PG", pg_summary)
        if _r4 == 1:
            ws4.cell(row=1, column=1, value="कोई डेटा उपलब्ध नहीं है।")
        ws4.column_dimensions['A'].width = 8
        ws4.column_dimensions['B'].width = 22
        ws4.column_dimensions['C'].width = 26
        for _cl in ['D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M', 'N']:
            ws4.column_dimensions[_cl].width = 16
        ws4.freeze_panes = "C2"

        # ================== Sheet 5: Reason वाले छात्रों की पूरी लिस्ट ==================
        def _write_students_section(ws, start_row, type_label, students_df, flags):
            if students_df is None or students_df.empty:
                return start_row
            cols5 = list(students_df.columns)
            reason_idx5 = next((i for i, c in enumerate(cols5) if "Reason" in str(c)), None)
            headers5 = ["Type"] + cols5
            for c_i, t_ in enumerate(headers5, start=1):
                cell = ws.cell(row=start_row, column=c_i, value=t_)
                cell.fill = head_fill3
                cell.font = Font(bold=True, color="FFFFFF")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = border3

            r_ = start_row + 1
            flags = flags or {}
            for r_i in range(len(students_df)):
                row = students_df.iloc[r_i]
                c0 = ws.cell(row=r_, column=1, value=type_label)
                c0.border = border3
                row_has_w = row_has_b = False
                for c_i0, cname in enumerate(cols5):
                    val = row.iloc[c_i0]
                    if c_i0 == reason_idx5 and isinstance(val, str):
                        val = val.replace("\n", "\n")
                    cell = ws.cell(row=r_, column=c_i0 + 2, value=val)
                    cell.border = border3
                    cell.alignment = Alignment(vertical="top", wrap_text=(c_i0 == reason_idx5))
                    f_list = flags.get(cname)
                    f = f_list[r_i] if f_list is not None else ""
                    if f == "w":
                        cell.fill = red_fill3
                        cell.font = Font(bold=True, color="721C24")
                        row_has_w = True
                    elif f == "b":
                        cell.fill = blue_fill3
                        cell.font = Font(bold=True, color="0C5460")
                        row_has_b = True
                    elif r_i % 2 == 1:
                        cell.fill = zebra_fill3
                if reason_idx5 is not None:
                    rc = ws.cell(row=r_, column=reason_idx5 + 2)
                    if row_has_w:
                        rc.fill = red_fill3
                        rc.font = Font(bold=True, color="721C24")
                    elif row_has_b:
                        rc.fill = blue_fill3
                        rc.font = Font(bold=True, color="0C5460")
                r_ += 1
            return r_ + 1

        ws5 = writer.book.create_sheet("Reason_Students")
        _r5 = _write_students_section(ws5, 1, "UG", ug_students, ug_flags)
        _r5 = _write_students_section(ws5, _r5, "PG", pg_students, pg_flags)
        if _r5 == 1:
            ws5.cell(row=1, column=1, value="कोई भी गलत (🔴) या खाली (🔵) एंट्री वाला छात्र नहीं मिला।")
            ws5.column_dimensions['A'].width = 70
        else:
            ws5.column_dimensions['A'].width = 8
            ws5.freeze_panes = "B2"

        # ================== Sheet 6: Approve किए गए छात्रों की लिस्ट (Sheet 5 में से) ==================
        def _write_approved_section(ws, start_row, type_label, approved_df):
            if approved_df is None or approved_df.empty:
                return start_row
            cols6 = list(approved_df.columns)
            headers6 = ["Type"] + cols6
            for c_i, t_ in enumerate(headers6, start=1):
                cell = ws.cell(row=start_row, column=c_i, value=t_)
                cell.fill = green_head_fill3
                cell.font = Font(bold=True, color="FFFFFF")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = border3

            r_ = start_row + 1
            reason_idx6 = next((i for i, c in enumerate(cols6) if "Reason" in str(c) or "कारण" in str(c)), None)
            for _, row in approved_df.iterrows():
                c0 = ws.cell(row=r_, column=1, value=type_label)
                c0.border = border3
                c0.fill = green_fill3
                for c_i0, cname in enumerate(cols6):
                    cell = ws.cell(row=r_, column=c_i0 + 2, value=row.iloc[c_i0])
                    cell.border = border3
                    cell.fill = green_fill3
                    cell.font = Font(color="155724")
                    cell.alignment = Alignment(vertical="top", wrap_text=(c_i0 == reason_idx6))
                r_ += 1
            return r_ + 1

        ws6 = writer.book.create_sheet("Approved_Students")
        _r6 = _write_approved_section(ws6, 1, "UG", ug_approved)
        _r6 = _write_approved_section(ws6, _r6, "PG", pg_approved)
        if _r6 == 1:
            ws6.cell(row=1, column=1, value="अभी तक इनमें से किसी भी 'गलत/खाली' छात्र को Approve नहीं किया गया है।")
            ws6.column_dimensions['A'].width = 70
        else:
            ws6.column_dimensions['A'].width = 8
            ws6.freeze_panes = "B2"

    return output.getvalue()

def process_panel_validation(df_panel, prefix, allowed_degrees, master_rules=None):
    deg_col = next((c for c in df_panel.columns if any(k in c.lower() for k in ['deg', 'course', 'class'])), df_panel.columns[0])
    br_col = next((c for c in df_panel.columns if any(k in c.lower() for k in ['branch', 'stream', 'subject'])), df_panel.columns[0])
    
    def check_degree(val):
        v = str(val).lower().replace(".", "").replace(" ", "").strip()
        return any(d in v for d in allowed_degrees)
        
    df_filtered = df_panel[df_panel[deg_col].apply(check_degree)].reset_index(drop=True)
    
    if df_filtered.empty:
        st.warning(f"⚠️ {prefix.upper()} पैनल के लिए कोई उपयुक्त डेटा (मैचिंग डिग्री) नहीं मिला।")
        return

    minor_col_found = next((c for c in df_filtered.columns if 'minor' in c.lower()), None)
    mdc_col_found = next((c for c in df_filtered.columns if 'mdc' in c.lower()), None)
    voc_col_found = next((c for c in df_filtered.columns if 'voc' in c.lower() or 'skill' in c.lower()), None)
    pw_col_found = next((c for c in df_filtered.columns if any(k in c.lower() for k in ['pw', 'project', 'ce'])), None)

    # छात्र की पहचान के लिए कॉलम (Roll/Enrollment/Name) — Approve सिस्टम के लिए ज़रूरी
    key_col = find_student_key_col(df_filtered)
    targets = {minor_col_found: 'minor', mdc_col_found: 'mdc', voc_col_found: 'voc', pw_col_found: 'pw'}

    # --- 🔎 हर रो के लिए मिसमैच वाले कॉलम निकालने का साझा फ़ंक्शन ---
    # (लाइव टेबल कलरिंग और नीचे की "Approve" लिस्ट — दोनों जगह इसी लॉजिक का इस्तेमाल होता है)
    def compute_row_mismatches(row):
        # 🔧 फिक्स: Degree column + Branch column दोनों को मिलाकर चेक करना
        # (Biotechnology / Commerce Computer जैसी ब्रांच अक्सर अलग Branch column में होती है, Degree column में नहीं)
        deg_part = str(row[deg_col]) if deg_col else ""
        br_part = str(row[br_col]) if br_col else ""
        student_deg = (deg_part + " " + br_part).lower().replace(".", "").replace(" ", "").strip()

        matched_key = "Default"
        if master_rules:
            sorted_keys = sorted(master_rules.keys(), key=len, reverse=True)
            for rule_key in sorted_keys:
                # 🔧 फिक्स: पूरा नाम एक साथ ढूंढने के बजाय हर word अलग-अलग ढूंढना
                # (जैसे "B.Com. Computer" -> "bcom" और "computer" दोनों कहीं भी मिलने चाहिए)
                rule_words = [w.lower().replace(".", "").strip() for w in rule_key.split() if w.strip()]
                if rule_words and all(w in student_deg for w in rule_words):
                    matched_key = rule_key
                    break

        c_rule = master_rules.get(matched_key, {"minor": [], "mdc": [], "voc": [], "pw": []}) if master_rules else {"minor": [], "mdc": [], "voc": [], "pw": []}

        mismatched_cols = []
        for col_name, rule_key in targets.items():
            if col_name and col_name in df_filtered.columns:
                val = row[col_name]
                if not (pd.isna(val) or str(val).strip() == ""):
                    val_clean = str(val).strip().lower().replace(".", "").replace(" ", "")
                    valid_list = c_rule.get(rule_key, [])
                    valid_set = {str(x).strip().lower().replace(".", "").replace(" ", "") for x in valid_list}
                    if valid_set and (val_clean not in valid_set):
                        mismatched_cols.append(col_name)
        return mismatched_cols

    # ⚪ खाली-छूट वाली डिग्री+ब्रांच की रो (इनमें Minor/MDC/Voc/PW खाली होने पर नीला नहीं होगा)
    _exempt_arr = blank_exempt_mask(df_filtered, prefix).tolist()

    # --- 🖥️ लाइव वैरिफिकेशन स्टाइलर फ़ंक्शन (स्क्रीन ग्रिड के लिए फिक्स) ---
    def cell_styler(dataframe):
        s_df = pd.DataFrame('', index=dataframe.index, columns=dataframe.columns)

        for position, (index, row) in enumerate(dataframe.iterrows()):
            # ब्रांच की खाली चेकिंग
            if br_col and br_col in dataframe.columns:
                b_val = row[br_col]
                if pd.isna(b_val) or str(b_val).strip() == "":
                    s_df.at[index, br_col] = 'background-color: #d1ecf1; color: #0c5460; font-weight: bold; border: 1px solid #17a2b8;'

            student_key = get_student_key(row, position, key_col)
            approval = get_approval(student_key, prefix)
            mismatched_cols = compute_row_mismatches(row)

            # स्क्रीन पर नियमों के अनुसार सटीक लाइव कलर कोडिंग
            for col_name in targets:
                if col_name and col_name in dataframe.columns:
                    val = row[col_name]
                    # 🔵 स्थिति 1: अगर पूरी तरह से ब्लैंक है तो नीला करें (खाली-छूट वाली डिग्री+ब्रांच को छोड़कर)
                    if (pd.isna(val) or str(val).strip() == "") and not _exempt_arr[position]:
                        s_df.at[index, col_name] = 'background-color: #d1ecf1; color: #0c5460; font-weight: bold; border: 1px solid #17a2b8;'

            # 🔴 गलत विषय → लाल | ✅ अगर पहले से Approve हो चुका है → हरा (अब गलत नहीं माना जाएगा)
            for col_name in mismatched_cols:
                if approval:
                    s_df.at[index, col_name] = 'background-color: #d4edda; color: #155724; font-weight: bold; border: 2px solid #28a745;'
                else:
                    s_df.at[index, col_name] = 'background-color: #f8d7da; color: #721c24; font-weight: bold; border: 2px solid red;'
        return s_df

    # =====================================================================
    # ⚪ खाली-छूट सेटिंग: यहाँ बताएँ कि किस डिग्री + ब्रांच में Minor / MDC / Voc / PW खाली ही रहते हैं
    # =====================================================================
    _ex_dc, _ex_bc = detect_deg_branch_cols(df_filtered)
    _ex_deg = df_filtered[_ex_dc].map(_norm_blank_label) if _ex_dc else pd.Series(["—"] * len(df_filtered), index=df_filtered.index)
    _ex_br = df_filtered[_ex_bc].map(_norm_blank_label) if _ex_bc else pd.Series(["—"] * len(df_filtered), index=df_filtered.index)
    _saved_pairs = get_blank_exempt_pairs(prefix)
    _all_pairs = sorted(set(zip(_ex_deg, _ex_br)) | _saved_pairs)
    _pair_label = lambda pr: f"{pr[0]} → {pr[1]}"
    _label_to_pair = {_pair_label(pr): pr for pr in _all_pairs}
    _ex_opts = list(_label_to_pair.keys())
    _ex_sig = hashlib.md5("|".join(_ex_opts).encode("utf-8")).hexdigest()[:8]

    with st.expander(f"⚪ खाली-छूट सेटिंग — जिन डिग्री + ब्रांच में Minor/MDC/Voc/PW खाली ही रहते हैं ({len(_saved_pairs)} चुनी हुई)"):
        st.caption("यहाँ चुनी गई डिग्री + ब्रांच में Minor, MDC, Voc और PW खाली होने पर 🔵 नीला रंग नहीं लगेगा, और खाली-गिनती में भी नहीं जुड़ेगा (टेबल, Excel डाउनलोड, खाली-सूची और डैशबोर्ड — सब जगह)। अगर उनमें कोई विषय भरा हो तो वह पहले की तरह जाँचा जाएगा।")
        _ex_sel = st.multiselect(
            "इन डिग्री + ब्रांच में खाली सेल नीला न हो:",
            options=_ex_opts,
            default=[_pair_label(pr) for pr in sorted(_saved_pairs)],
            key=f"blank_exempt_sel_{prefix}_{_ex_sig}"
        )
        if st.button("💾 खाली-छूट सेव करें", key=f"blank_exempt_save_{prefix}"):
            set_blank_exempt_pairs(prefix, {_label_to_pair[l_] for l_ in _ex_sel})
            st.success("🎉 खाली-छूट सेव हो गई! अब चुनी हुई डिग्री+ब्रांच में खाली सेल नीले नहीं दिखेंगे।")
            st.rerun()

    st.subheader(f"📊 लाइव वैरिफाइड {prefix.upper()} डेटा टेबल")
    st.caption("🔵 नीला सेल = डेटा गायब है | 🔴 लाल सेल = गलत विषय (मास्टर गाइडलाइन से मिसमैच) | 🟢 हरा सेल = Approved (मान्य किया गया, अब गलत नहीं गिना जाएगा)")
    
    # स्क्रीन पर सीरियल नंबर 1 से शुरू करना
    df_filtered.index = range(1, len(df_filtered) + 1)
    st.dataframe(df_filtered.style.apply(cell_styler, axis=None), height=500, use_container_width=True)

    # approved_keys सेट पहले से निकाल लेना ताकि एक्सेल डाउनलोड और नीचे की लिस्ट दोनों इस्तेमाल कर सकें
    all_approvals_now = get_all_approvals(prefix)
    approved_keys_set = {a[0] for a in all_approvals_now}

    # --- 🚨 📥 रंगीन एक्सेल डाउनलोड (अब शेयर्ड फ़ंक्शन का इस्तेमाल कर रहा है) 🚨 ---
    processed_data = generate_colored_excel_bytes(
        df_filtered, deg_col, br_col, minor_col_found, mdc_col_found, voc_col_found, pw_col_found,
        master_rules=master_rules, sheet_name="Verified_Data", approved_keys=approved_keys_set, key_col=key_col,
        prefix=prefix
    )
    st.download_button(
        label=f"📥 रंगीन (🔴/🔵/🟢) {prefix.upper()} डेटा एक्सेल डाउनलोड करें",
        data=processed_data,
        file_name=f"Verified_{prefix.upper()}_Colored_Data.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key=f"download_validated_excel_{prefix}"
    )

    # =====================================================================
    # 🔵 खाली (Blank) डेटा — किस डिग्री + ब्रांच में क्या खाली है
    # =====================================================================
    st.divider()
    st.subheader("🔵 खाली (Blank) डेटा — किस डिग्री + ब्रांच में")
    st.caption("यहाँ दिखता है कि किस डिग्री + ब्रांच में कौन सा कॉलम खाली (🔵) है और कितने छात्रों का। जाँचने वाले कॉलम नीचे के बॉक्स से बदल भी सकते हैं (PG में जहाँ Minor/MDC जैसे कॉलम नहीं होते, वहाँ अपनी पसंद के कॉलम चुनें)।")

    _skip_kw = ['minor', 'mdc', 'voc', 'skill', 'pw', 'project']
    _bl_deg = deg_col if (deg_col and deg_col in df_filtered.columns) else None
    _bl_br = None
    for _kws in (['branch', 'stream'], ['subject']):
        _bl_br = next((c for c in df_filtered.columns if c != _bl_deg
                       and any(k in str(c).lower() for k in _kws)
                       and not any(k in str(c).lower() for k in _skip_kw)), None)
        if _bl_br is not None:
            break

    _bl_default = [c for c in (minor_col_found, mdc_col_found, voc_col_found, pw_col_found) if c]
    if _bl_br is not None and _bl_br not in _bl_default:
        _bl_default.append(_bl_br)

    # कॉलम-सूची बदलने पर (नई फ़ाइल) चुनाव अपने-आप नए सिरे से शुरू हो, इसलिए key में कॉलमों की पहचान जोड़ी है
    _bl_sig = hashlib.md5("|".join(map(str, df_filtered.columns)).encode("utf-8")).hexdigest()[:8]
    _bl_key = f"blank_cols_{prefix}_{_bl_sig}"
    if _bl_key not in st.session_state:
        st.session_state[_bl_key] = _bl_default

    st.multiselect(
        "🔎 किन कॉलम में खाली जाँचना है:",
        options=list(df_filtered.columns),
        key=_bl_key
    )
    _bl_cols = [c for c in st.session_state[_bl_key] if c in df_filtered.columns]

    if not _bl_cols:
        st.info("ℹ️ ऊपर के बॉक्स से कम से कम एक कॉलम चुनें, फिर खाली डेटा की सूची यहाँ दिखेगी।")
    else:
        def _bl_clean(series):
            s_ = series.astype(str).str.strip()
            return s_.mask(series.isna() | (s_ == "") | (s_.str.lower() == "nan"), "(खाली/Blank)")

        _work = df_filtered.copy()
        _work["Degree (डिग्री)"] = _bl_clean(_work[_bl_deg]) if _bl_deg else "—"
        _work["Branch (ब्रांच)"] = _bl_clean(_work[_bl_br]) if _bl_br else "—"
        _ex_s = pd.Series(_exempt_arr, index=_work.index)
        _cat_set = {c_ for c_ in (minor_col_found, mdc_col_found, voc_col_found, pw_col_found) if c_}
        for _c in _bl_cols:
            _fl = _work[_c].isna() | (_work[_c].astype(str).str.strip() == "")
            if _c in _cat_set:
                _fl = _fl & ~_ex_s          # ⚪ खाली-छूट वाली डिग्री+ब्रांच में Minor..PW का खाली नहीं गिनना
            _work[f"__b__{_c}"] = _fl
        _flag_cols = [f"__b__{_c}" for _c in _bl_cols]
        _work["__any__"] = _work[_flag_cols].any(axis=1)
        if int(_ex_s.sum()) > 0:
            st.caption(f"⚪ खाली-छूट वाली डिग्री+ब्रांच के {int(_ex_s.sum())} छात्रों के Minor/MDC/Voc/PW यहाँ गिने नहीं गए (सेटिंग ऊपर 'खाली-छूट सेटिंग' में है)।")

        _agg = {"कुल छात्र (Total)": ("__any__", "size"), "🔵 खाली वाले छात्र": ("__any__", "sum")}
        for _c in _bl_cols:
            _agg[f"🔵 {_c}"] = (f"__b__{_c}", "sum")
        _bl_summary = _work.groupby(["Degree (डिग्री)", "Branch (ब्रांच)"]).agg(**_agg).reset_index()
        for _c in ["🔵 खाली वाले छात्र"] + [f"🔵 {_c}" for _c in _bl_cols]:
            _bl_summary[_c] = _bl_summary[_c].astype(int)

        def _bl_reason(r):
            _parts = [f"{_c}: {int(r[f'🔵 {_c}'])} छात्र" for _c in _bl_cols if r[f"🔵 {_c}"] > 0]
            return ("🔵 खाली → " + " ; ".join(_parts)) if _parts else ""
        _bl_summary["📝 कारण (Reason)"] = _bl_summary.apply(_bl_reason, axis=1)

        _total_combos = len(_bl_summary)
        _blank_combos = int((_bl_summary["🔵 खाली वाले छात्र"] > 0).sum())
        _blank_students = int(_work["__any__"].sum())
        _blank_cells = int(sum(_work[f].sum() for f in _flag_cols))

        bm1, bm2, bm3, bm4 = st.columns(4)
        with bm1:
            st.metric("🎓 कुल डिग्री+ब्रांच", _total_combos)
        with bm2:
            st.metric("🔵 खाली वाली डिग्री+ब्रांच", _blank_combos)
        with bm3:
            st.metric("👥 खाली वाले छात्र", _blank_students)
        with bm4:
            st.metric("🔵 कुल खाली सेल", _blank_cells)

        if _blank_cells == 0:
            st.success("✅ चुने गए कॉलम में कोई भी खाली डेटा नहीं मिला।")
        else:
            _show_all_bl = st.checkbox("बिना खाली वाली डिग्री+ब्रांच भी दिखाएँ", value=False, key=f"blank_show_all_{prefix}")
            _bl_show = _bl_summary if _show_all_bl else _bl_summary[_bl_summary["🔵 खाली वाले छात्र"] > 0]
            _bl_show = _bl_show.sort_values(["🔵 खाली वाले छात्र", "Degree (डिग्री)"], ascending=[False, True]).reset_index(drop=True)

            _bl_blue = 'background-color: #d1ecf1; color: #0c5460; font-weight: bold; border: 2px solid #17a2b8;'
            _bl_blue_cols = [c for c in _bl_show.columns if "🔵" in c]

            def _bl_styler(row):
                _any_b = any(row[c] > 0 for c in _bl_blue_cols)
                out = []
                for c in row.index:
                    if c in _bl_blue_cols and row[c] > 0:
                        out.append(_bl_blue)
                    elif c == "📝 कारण (Reason)" and row[c] and _any_b:
                        out.append(_bl_blue)
                    else:
                        out.append("")
                return out

            st.dataframe(
                _bl_show.style.apply(_bl_styler, axis=1),
                hide_index=True,
                use_container_width=True,
                height=min(38 * (len(_bl_show) + 1) + 3, 650),
                column_config={"📝 कारण (Reason)": st.column_config.TextColumn(width=500)}
            )

            _bd1, _bd2 = st.columns(2)
            with _bd1:
                st.download_button(
                    label="📥 खाली-डेटा सूची CSV में डाउनलोड करें",
                    data=_bl_show.to_csv(index=False).encode("utf-8-sig"),
                    file_name=f"{prefix.upper()}_Blank_Degree_Branch_Report.csv",
                    mime="text/csv",
                    key=f"blank_dl_csv_{prefix}",
                    use_container_width=True
                )
            with _bd2:
                st.download_button(
                    label="📊 खाली-डेटा सूची Excel (XLSX) में डाउनलोड करें",
                    data=generate_summary_excel_bytes(_bl_show, sheet_name=f"{prefix.upper()}_Blank"),
                    file_name=f"{prefix.upper()}_Blank_Degree_Branch_Report.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key=f"blank_dl_xlsx_{prefix}",
                    use_container_width=True
                )

            # 👨‍🎓 छात्र-वार सूची (ताकि सीधे पता चले कि किस छात्र का कौन सा कॉलम भरना है)
            _bl_rows = _work[_work["__any__"]]
            with st.expander(f"👨‍🎓 छात्र-वार खाली सूची ({len(_bl_rows)} छात्र)"):
                _stu = pd.DataFrame({"टेबल में रो नं.": _bl_rows.index})
                if key_col and key_col in _bl_rows.columns:
                    _stu["छात्र (Key)"] = _bl_rows[key_col].astype(str).values
                _stu["Degree (डिग्री)"] = _bl_rows["Degree (डिग्री)"].values
                _stu["Branch (ब्रांच)"] = _bl_rows["Branch (ब्रांच)"].values
                _stu["🔵 खाली कॉलम"] = _bl_rows.apply(
                    lambda r: ", ".join(_c for _c in _bl_cols if r[f"__b__{_c}"]), axis=1
                ).values
                st.dataframe(_stu, hide_index=True, use_container_width=True, height=min(38 * (len(_stu) + 1) + 3, 500))


    # =====================================================================
    # ✅ "गलत विषय" Approve करने का सिस्टम
    # =====================================================================
    st.divider()
    st.subheader("🛠️ गलत विषय Approve करें")
    st.caption("अगर किसी छात्र का विषय मास्टर गाइडलाइन से मैच नहीं हो रहा लेकिन वह सही है, तो यहाँ से उसे Approve करें — फिर वह टेबल में लाल नहीं, हरा (✅ Approved) दिखेगा।")

    pending_students = []
    for position, (index, row) in enumerate(df_filtered.iterrows()):
        mismatched_cols = compute_row_mismatches(row)
        if mismatched_cols:
            student_key = get_student_key(row, position, key_col)
            if not get_approval(student_key, prefix):
                pending_students.append((index, row, student_key, mismatched_cols))

    if not pending_students:
        st.info("✅ फिलहाल कोई भी 'गलत विषय' वाला छात्र Approve होने के लिए बाकी नहीं है।")
    else:
        for index, row, student_key, mismatched_cols in pending_students:
            name_display = str(row[key_col]) if key_col and key_col in df_filtered.columns else f"रो #{index}"
            with st.expander(f"⚠️ {name_display} — गलत कॉलम: {', '.join(mismatched_cols)}"):
                st.write({c: row[c] for c in mismatched_cols})
                appr_col1, appr_col2 = st.columns([2, 1])
                with appr_col1:
                    approver_role = st.selectbox(
                        "किसने Approve किया?",
                        ["Nodal", "Student", "Principal"],
                        key=f"approver_role_{prefix}_{student_key}_{index}"
                    )
                with appr_col2:
                    st.write("")
                    if st.button("✅ Approve करें", key=f"approve_btn_{prefix}_{student_key}_{index}", use_container_width=True):
                        add_approval(student_key, prefix, approver_role)
                        st.success(f"🎉 {name_display} को {approver_role} द्वारा Approve कर दिया गया!")
                        st.rerun()

    # =====================================================================
    # 📋 Approved List (जिन छात्रों का गलत विषय Approve किया जा चुका है)
    # =====================================================================
    st.divider()
    st.subheader("📋 Approved List")
    st.caption("यहाँ वो सभी छात्र दिखेंगे जिनका 'गलत विषय' मान्य (Approve) किया जा चुका है — साथ में किसने Approve किया, यह भी दिखेगा।")

    if not all_approvals_now:
        st.info("अभी तक कोई भी छात्र Approve नहीं हुआ है।")
    else:
        approved_rows = [
            {"छात्र (Key)": sk, "Approve किया": ab, "समय": at}
            for sk, ab, at in all_approvals_now
        ]
        approved_display_df = pd.DataFrame(approved_rows)
        st.dataframe(approved_display_df, use_container_width=True, hide_index=True)
        panel_orientation = st.radio(
            "🖨️ प्रिंट ओरिएंटेशन चुनें:",
            options=["📄 Portrait (सीधा)", "📃 Landscape (आड़ा)"],
            horizontal=True,
            key=f"orientation_panel_{prefix}"
        )
        render_print_button(
            approved_display_df,
            title=f"Approved List - {prefix.upper()}",
            key_suffix=f"panel_{prefix}",
            orientation="landscape" if "Landscape" in panel_orientation else "portrait"
        )

        revoke_choice = st.selectbox(
            "❌ किसी Approval को हटाना है? (Revoke करें):",
            ["-- चुनें --"] + [a[0] for a in all_approvals_now],
            key=f"revoke_select_{prefix}"
        )
        if revoke_choice != "-- चुनें --":
            if st.button("🗑️ चुनी गई Approval हटाएं (Revoke)", key=f"revoke_btn_{prefix}"):
                remove_approval(revoke_choice, prefix)
                st.success("🎉 Approval हटा दी गई — यह छात्र अब फिर से 'गलत विषय' में गिना जाएगा।")
                st.rerun()

# =========================================================================
# 🔄 Excel (.xls / .xlsx) → CSV कन्वर्टर (Panel 1 के लिए)
# =========================================================================
_OLE_MAGIC = b"\xd0\xcf\x11\xe0"   # असली .xls (पुराना Excel)
_ZIP_MAGIC = b"PK"                # असली .xlsx / .ods

class _HTMLTableParser(HTMLParser):
    """बिना किसी एक्स्ट्रा लाइब्रेरी (lxml/bs4) के HTML <table> पढ़ने वाला पार्सर।"""
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tables = []
        self._stack = []
        self._row = None
        self._cell = None

    def _close_cell(self):
        if self._cell is not None and self._row is not None:
            self._row.append("".join(self._cell).strip())
        self._cell = None

    def _close_row(self):
        self._close_cell()
        if self._row is not None and self._stack and self._row:
            self._stack[-1].append(self._row)
        self._row = None

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self._stack.append([])
        elif tag == "tr" and self._stack:
            self._close_row()
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._close_cell()
            self._cell = []
        elif tag == "br" and self._cell is not None:
            self._cell.append(" ")

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)

    def handle_endtag(self, tag):
        if tag in ("td", "th"):
            self._close_cell()
        elif tag == "tr":
            self._close_row()
        elif tag == "table" and self._stack:
            self._close_row()
            self.tables.append(self._stack.pop())

    def finish(self):
        self._close_row()
        while self._stack:
            self.tables.append(self._stack.pop())

def _html_tables_to_df(text):
    parser = _HTMLTableParser()
    parser.feed(text)
    parser.close()
    parser.finish()
    tables = [t for t in parser.tables if t]
    if not tables:
        raise ValueError("HTML में कोई टेबल नहीं मिली")
    rows = max(tables, key=len)                      # सबसे बड़ी टेबल = असली डेटा
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    header = [h if h else f"Unnamed: {i}" for i, h in enumerate(rows[0])]
    return pd.DataFrame(rows[1:], columns=header)

def _spreadsheetml_to_df(raw):
    """Excel 2003 'XML Spreadsheet' फ़ाइल (पहली शीट)।"""
    ns = "{urn:schemas-microsoft-com:office:spreadsheet}"
    root = ET.fromstring(raw)
    ws = root.find(f"{ns}Worksheet")
    if ws is None:
        raise ValueError("XML में कोई वर्कशीट नहीं मिली")
    rows = []
    for row in ws.iter(f"{ns}Row"):
        vals = []
        for cell in row.findall(f"{ns}Cell"):
            idx = cell.get(f"{ns}Index")
            if idx:
                while len(vals) < int(idx) - 1:
                    vals.append("")
            data = cell.find(f"{ns}Data")
            vals.append("".join(data.itertext()).strip() if data is not None else "")
        rows.append(vals)
    if not rows:
        raise ValueError("XML शीट खाली है")
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    header = [h if h else f"Unnamed: {i}" for i, h in enumerate(rows[0])]
    return pd.DataFrame(rows[1:], columns=header)

def _decode_text_bytes(raw):
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return raw.decode("utf-16", errors="replace")
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace")

def _delimited_text_to_df(text):
    if "\x00" in text[:20000]:
        raise ValueError("यह टेक्स्ट नहीं, बाइनरी फ़ाइल लगती है")
    first_line = text.lstrip("\ufeff").split("\n", 1)[0]
    counts = {d: first_line.count(d) for d in ("\t", ",", ";", "|")}
    sep = max(counts, key=counts.get)
    if counts[sep] == 0:
        raise ValueError("कोई कॉलम-सेपरेटर (Tab/Comma) नहीं मिला")
    return pd.read_csv(io.StringIO(text), sep=sep)

@st.cache_data(show_spinner=False)
def get_excel_sheet_names(raw):
    """असली Excel फ़ाइल की शीट्स के नाम (नकली .xls के लिए खाली लिस्ट)।"""
    if raw[:4] == _OLE_MAGIC or raw[:2] == _ZIP_MAGIC:
        try:
            return pd.ExcelFile(io.BytesIO(raw)).sheet_names
        except Exception:
            return []
    return []

def excel_bytes_to_dataframe(raw, sheet_name=0):
    """.xls / .xlsx को DataFrame में पढ़ता है। कई पोर्टल .xls नाम से असल में
    HTML / XML / टेक्स्ट फ़ाइल देते हैं — उन सबको भी संभालता है।"""
    if not raw:
        raise ValueError("फ़ाइल खाली है")
    # 1) असली Excel (xls / xlsx / ods)
    if raw[:4] == _OLE_MAGIC or raw[:2] == _ZIP_MAGIC:
        return pd.read_excel(io.BytesIO(raw), sheet_name=sheet_name)
    # 2) नकली .xls: अंदर से देखकर टाइप पहचानना
    text = _decode_text_bytes(raw)
    sample = text[:500000].lower()
    try:
        if "urn:schemas-microsoft-com:office:spreadsheet" in sample:
            return _spreadsheetml_to_df(raw)
        if "<table" in sample:
            return _html_tables_to_df(text)
        return _delimited_text_to_df(text)
    except Exception as e:
        raise ValueError(
            f"यह फ़ाइल असली Excel नहीं है और पढ़ी भी नहीं जा सकी ({e}). "
            f"फ़ाइल की शुरुआत: {raw[:150]!r}"
        ) from e

@st.cache_data(show_spinner="🔄 Excel फ़ाइल को CSV में बदला जा रहा है...")
def excel_bytes_to_csv_bytes(raw, sheet_name=0):
    """Excel bytes → (CSV bytes, रोज़ की संख्या)। utf-8-sig ताकि हिंदी Excel में भी सही खुले।"""
    df_in = excel_bytes_to_dataframe(raw, sheet_name)
    return df_in.to_csv(index=False).encode("utf-8-sig"), len(df_in)

# =========================================================================
# 📑 ब्रांच-वाइज विषय शीट: Degree/Branch/Total Admission मर्ज + Minor/MDC/Voc/PW नाम+काउंट
# =========================================================================
def build_branch_subject_blocks(df, deg_c, br_c, cats):
    """cats = [(label, column_name), ...] -> हर Degree+Branch के लिए एक ब्लॉक"""
    blocks = []
    for (d_, b_), g_ in df.groupby([deg_c, br_c], sort=False):
        cat_lists = []
        for _lbl, _cn in cats:
            s_ = g_[_cn].dropna().astype(str).str.strip()
            s_ = s_[(s_ != "") & (s_.str.lower() != "nan")]
            items_ = sorted(s_.value_counts().items(), key=lambda kv: (-kv[1], kv[0]))
            cat_lists.append([(str(n_), int(c_)) for n_, c_ in items_])
        n_rows = max([1] + [len(x_) for x_ in cat_lists])
        blocks.append({"degree": d_, "branch": b_, "total": int(len(g_)), "cats": cat_lists, "n": n_rows})
    blocks.sort(key=lambda x_: (str(x_["degree"]), -x_["total"], str(x_["branch"])))
    return blocks


def branch_subject_sheet_html(blocks, cat_labels):
    """स्क्रीन पर दिखाने के लिए rowspan (मर्ज) वाली HTML तालिका"""
    def _e(x):
        return _html.escape(str(x)).replace("$", "&#36;")
    th = "position:sticky;top:0;background:#1a3c6e;color:#fff;padding:8px 10px;border:1px solid #BFBFBF;text-align:center;font-size:13px;white-space:nowrap;z-index:2;"
    h = ['<div style="max-height:650px;overflow:auto;border:1px solid #BFBFBF;border-radius:6px;">',
         '<table style="border-collapse:collapse;width:100%;font-size:13.5px;color:#111;">', '<thead><tr>']
    heads = ["Degree (डिग्री)", "Branch (ब्रांच)", "Total Admission"]
    for l_ in cat_labels:
        heads += [f"{l_} (विषय)", f"{l_} Count"]
    for t_ in heads:
        h.append(f'<th style="{th}">{_e(t_)}</th>')
    h.append("</tr></thead><tbody>")
    for bi, b_ in enumerate(blocks):
        bg = "#EAF1FB" if bi % 2 == 0 else "#FFFFFF"
        td = f"padding:6px 10px;border:1px solid #BFBFBF;background:{bg};"
        for i in range(b_["n"]):
            h.append("<tr>")
            if i == 0:
                rs = b_["n"]
                h.append(f'<td rowspan="{rs}" style="{td}vertical-align:middle;text-align:center;font-weight:700;">{_e(b_["degree"])}</td>')
                h.append(f'<td rowspan="{rs}" style="{td}vertical-align:middle;text-align:center;font-weight:700;">{_e(b_["branch"])}</td>')
                h.append(f'<td rowspan="{rs}" style="{td}vertical-align:middle;text-align:center;font-weight:800;color:#1a3c6e;">{b_["total"]}</td>')
            for items in b_["cats"]:
                if i < len(items):
                    h.append(f'<td style="{td}">{_e(items[i][0])}</td><td style="{td}text-align:center;">{items[i][1]}</td>')
                else:
                    h.append(f'<td style="{td}"></td><td style="{td}"></td>')
            h.append("</tr>")
    h.append("</tbody></table></div>")
    return "".join(h)


def generate_branch_subject_sheet_excel_bytes(blocks, cat_labels, sheet_name="Branch_Subject_Sheet"):
    """असली merged cells वाली Excel शीट"""
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Border, Side, Font, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "".join(ch for ch in str(sheet_name) if ch not in '[]:*?/\\')[:31] or "Sheet1"

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", start_color="1A3C6E", end_color="1A3C6E")
    band_a = PatternFill("solid", start_color="EAF1FB", end_color="EAF1FB")
    band_b = PatternFill("solid", start_color="FFFFFF", end_color="FFFFFF")

    headers = ["Degree (डिग्री)", "Branch (ब्रांच)", "Total Admission"]
    for l_ in cat_labels:
        headers += [f"{l_} (विषय)", f"{l_} Count"]
    ncols = len(headers)

    for c_i, t_ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c_i, value=t_)
        cell.fill = head_fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    r = 2
    for bi, b_ in enumerate(blocks):
        start, end = r, r + b_["n"] - 1
        fill = band_a if bi % 2 == 0 else band_b
        for rr in range(start, end + 1):
            for cc in range(1, ncols + 1):
                cell = ws.cell(row=rr, column=cc)
                cell.fill = fill
                cell.border = border
        for cc, val in ((1, b_["degree"]), (2, b_["branch"]), (3, b_["total"])):
            cell = ws.cell(row=start, column=cc, value=val)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if end > start:
                ws.merge_cells(start_row=start, start_column=cc, end_row=end, end_column=cc)
        for k, items in enumerate(b_["cats"]):
            for i, (name_, cnt_) in enumerate(items):
                ws.cell(row=start + i, column=4 + 2 * k, value=name_).alignment = Alignment(vertical="center", wrap_text=True)
                ws.cell(row=start + i, column=5 + 2 * k, value=cnt_).alignment = Alignment(horizontal="center", vertical="center")
        r = end + 1

    widths = [22, 28, 16] + [34, 12] * len(cat_labels)
    for c_i, w_ in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c_i)].width = w_
    ws.freeze_panes = "A2"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# =========================================================================
# 📊 डिग्री+ब्रांच समरी → रंगीन Excel (लाल/नीली रो + कारण कॉलम)
# =========================================================================
def generate_summary_excel_bytes(summary_df, sheet_name="Summary", students_df=None, student_flags=None, students_sheet_name="Reason_Students"):
    from openpyxl.styles import PatternFill, Border, Side, Font, Alignment
    from openpyxl.utils import get_column_letter

    safe_sheet = "".join(ch for ch in str(sheet_name) if ch not in '[]:*?/\\')[:31] or "Summary"
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        summary_df.to_excel(writer, index=False, sheet_name=safe_sheet)
        ws = writer.sheets[safe_sheet]

        thin = Side(style="thin", color="BFBFBF")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        red_fill = PatternFill("solid", start_color="F8D7DA", end_color="F8D7DA")
        blue_fill = PatternFill("solid", start_color="D1ECF1", end_color="D1ECF1")
        head_fill = PatternFill("solid", start_color="1A3C6E", end_color="1A3C6E")

        cols = list(summary_df.columns)
        red_idx = [i for i, c in enumerate(cols) if "🔴" in str(c)]      # "गलत" वाले कॉलम
        blue_idx = [i for i, c in enumerate(cols) if "🔵" in str(c)]     # "खाली" वाले कॉलम
        reason_idx = next((i for i, c in enumerate(cols) if "Reason" in str(c)), None)

        # हेडर
        for c_i in range(1, len(cols) + 1):
            cell = ws.cell(row=1, column=c_i)
            cell.fill = head_fill
            cell.font = Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border

        # डेटा रो: गलत > 0 = लाल सेल, खाली > 0 = नीला सेल, कारण सेल = लाल (गलत हो तो) वरना नीला
        for r_i, (_, row) in enumerate(summary_df.iterrows(), start=2):
            if str(row.iloc[0]).startswith("कुल योग"):
                for c_i0 in range(len(cols)):
                    cell = ws.cell(row=r_i, column=c_i0 + 1)
                    cell.border = border
                    cell.fill = head_fill
                    cell.font = Font(bold=True, color="FFFFFF")
                continue
            any_w = any(row.iloc[i] > 0 for i in red_idx)
            any_b = any(row.iloc[i] > 0 for i in blue_idx)
            for c_i0 in range(len(cols)):
                cell = ws.cell(row=r_i, column=c_i0 + 1)
                cell.border = border
                is_reason = (c_i0 == reason_idx)
                cell.alignment = Alignment(vertical="top", wrap_text=is_reason)
                fill, fcolor = None, "000000"
                if c_i0 in red_idx and row.iloc[c_i0] > 0:
                    fill, fcolor = red_fill, "721C24"
                elif c_i0 in blue_idx and row.iloc[c_i0] > 0:
                    fill, fcolor = blue_fill, "0C5460"
                elif is_reason and str(row.iloc[c_i0]).strip():
                    fill, fcolor = (red_fill, "721C24") if any_w else ((blue_fill, "0C5460") if any_b else (None, "000000"))
                if is_reason:
                    cell.value = str(row.iloc[c_i0]).replace("  ||  ", "\n")   # हर श्रेणी नई लाइन में
                if fill is not None:
                    cell.fill = fill
                    cell.font = Font(bold=True, color=fcolor)

        # कॉलम चौड़ाई (कारण कॉलम चौड़ा + टेक्स्ट रैप)
        for c_i, name in enumerate(cols):
            if c_i == reason_idx:
                width = 100
            else:
                longest = max([len(str(name))] + [len(str(v)) for v in summary_df.iloc[:, c_i].tolist()])
                width = min(max(longest + 3, 12), 40)
            ws.column_dimensions[get_column_letter(c_i + 1)].width = width

        ws.freeze_panes = "C2"
        _last_r = ws.max_row - 1 if str(summary_df.iloc[-1, 0]).startswith("कुल योग") and ws.max_row > 2 else ws.max_row
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{_last_r}"   # फ़िल्टर/सॉर्ट में Total रो शामिल नहीं

        # 📄 शीट 2: Reason वाले छात्रों की लिस्ट (जिनका Minor/MDC/Voc/PW गलत या खाली है)
        if students_df is not None:
            safe_sheet2 = "".join(ch for ch in str(students_sheet_name) if ch not in '[]:*?/\\')[:31] or "Reason_Students"
            if safe_sheet2 == safe_sheet:
                safe_sheet2 = (safe_sheet2[:28] + "_2")
            if students_df.empty:
                pd.DataFrame({"सूचना": ["कोई भी गलत (🔴) या खाली (🔵) एंट्री वाला छात्र नहीं मिला।"]}).to_excel(
                    writer, index=False, sheet_name=safe_sheet2)
                ws2 = writer.sheets[safe_sheet2]
                ws2.column_dimensions["A"].width = 70
            else:
                students_df.to_excel(writer, index=False, sheet_name=safe_sheet2)
                ws2 = writer.sheets[safe_sheet2]
                cols2 = list(students_df.columns)
                reason_i2 = next((i for i, c in enumerate(cols2) if "Reason" in str(c)), None)
                for c_i in range(1, len(cols2) + 1):
                    cell = ws2.cell(row=1, column=c_i)
                    cell.fill = head_fill
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    cell.border = border
                flags = student_flags or {}
                for r_i in range(len(students_df)):
                    row_has_w = False
                    row_has_b = False
                    for c_i0, cname in enumerate(cols2):
                        cell = ws2.cell(row=r_i + 2, column=c_i0 + 1)
                        cell.border = border
                        cell.alignment = Alignment(vertical="top", wrap_text=(c_i0 == reason_i2))
                        f_list = flags.get(cname)
                        f = f_list[r_i] if f_list is not None else ""
                        if r_i % 2 == 1 and not f:
                            cell.fill = PatternFill("solid", start_color="F4F7FD", end_color="F4F7FD")
                        if f == "w":
                            cell.fill = red_fill
                            cell.font = Font(bold=True, color="721C24")
                            row_has_w = True
                        elif f == "b":
                            cell.fill = blue_fill
                            cell.font = Font(bold=True, color="0C5460")
                            row_has_b = True
                    if reason_i2 is not None:
                        rc = ws2.cell(row=r_i + 2, column=reason_i2 + 1)
                        if row_has_w:
                            rc.fill = red_fill
                            rc.font = Font(bold=True, color="721C24")
                        elif row_has_b:
                            rc.fill = blue_fill
                            rc.font = Font(bold=True, color="0C5460")
                for c_i, name in enumerate(cols2):
                    if c_i == reason_i2:
                        width = 70
                    else:
                        longest = max([len(str(name))] + [len(str(v)) for v in students_df.iloc[:, c_i].tolist()[:500]])
                        width = min(max(longest + 3, 10), 32)
                    ws2.column_dimensions[get_column_letter(c_i + 1)].width = width
                ws2.freeze_panes = "B2"
                ws2.auto_filter.ref = ws2.dimensions
    return output.getvalue()

# =========================================================================
# 📥 PANEL 1: ENTRY / UPLOAD PANEL (डेटा सुरक्षित अपलोड)
# =========================================================================
if active_panel == "📥 1. Entry / Upload Panel":
    st.title("📥 Entry Panel - डेटा सुरक्षित अपलोड")
    if is_panel_hidden("p1") and not is_admin_session:
        st.warning("🔒 यह पैनल फिलहाल Admin द्वारा Hide किया गया है। डेटा उपलब्ध नहीं है।")
        st.dataframe(pd.DataFrame(), use_container_width=True)
        st.stop()
    st.write("यहाँ अपनी मुख्य एक्सेल/CSV फ़ाइल अपलोड करें। यह डेटा सीधे समीक्षा और क्लीनिंग के लिए **Work / Approve Panel (P2)** में ट्रांसफर हो जाएगा।")
    
    # ✅ ट्रांसफर के बाद वाला सफलता संदेश (अपलोडर खाली होने के बाद दिखता है)
    if st.session_state.pop("p1_transfer_done", False):
        st.success("🎉 डेटा सफलतापूर्वक **Work / Approve Panel (P2)** में ट्रांसफर हो गया है! अब यहाँ से फ़ाइल हट गई है — नई फ़ाइल अपलोड की जा सकती है।")
        st.balloons()

    # एक्सेल या सीएसवी फ़ाइल अपलोड करने का विकल्प
    # (key बदलते ही अपलोडर खाली हो जाता है — इसी से ट्रांसफर के बाद फ़ाइल हटाई जाती है)
    if "p1_uploader_gen" not in st.session_state:
        st.session_state["p1_uploader_gen"] = 0
    f = st.file_uploader(
        "अपनी फ़ाइल अपलोड करें (CSV / XLS / XLSX)",
        type=["csv", "xls", "xlsx"],
        key=f"p1_uploader_{st.session_state['p1_uploader_gen']}"
    )
    if f:
        try:
            if f.name.lower().endswith((".xls", ".xlsx")):
                # 🔄 Excel फ़ाइल → पहले CSV में कन्वर्ट, फिर बिल्कुल CSV की तरह ही आगे प्रोसेस
                raw_bytes = f.getvalue()
                sheet_names = get_excel_sheet_names(raw_bytes)
                sheet_choice = 0
                if len(sheet_names) > 1:
                    sheet_choice = st.selectbox("📑 कौन सी शीट लोड करनी है?", sheet_names, key="p1_sheet_select")
                csv_bytes, n_rows = excel_bytes_to_csv_bytes(raw_bytes, sheet_choice)
                st.info(f"🔄 Excel फ़ाइल '{f.name}' अपने-आप CSV में बदल दी गई ({n_rows} रोज़)।")
                st.download_button(
                    "📥 बदली हुई CSV फ़ाइल डाउनलोड करें (ज़रूरत हो तो)",
                    data=csv_bytes,
                    file_name=os.path.splitext(f.name)[0] + ".csv",
                    mime="text/csv",
                    key="p1_converted_csv_dl"
                )
                # CSV से वापस पढ़ना, ताकि डेटा बिल्कुल CSV अपलोड जैसा ही बर्ताव करे
                df = pd.read_csv(io.BytesIO(csv_bytes), encoding="utf-8-sig")
            else:
                df = pd.read_csv(f)
            st.success(f"🎉 फ़ाइल सफलतापूर्वक लोड हो गई ({len(df)} रोज़)!")
            
            # डेटा को P2 में ट्रांसफर करने का बटन
            if st.button("📤 वर्क/अप्रूवल पैनल (P2) में ट्रांसफर करें", key="transfer_p2_btn", type="primary"):
                # पुराने किसी भी रॉ (Temporary) डेटा को साफ़ करना
                cursor.execute("DELETE FROM raw_store")
                
                # डेटाफ़्रेम को JSON में बदलकर सुरक्षित रूप से अस्थायी डेटाबेस में स्टोर करना
                cursor.execute("INSERT INTO raw_store (data_json) VALUES (?)", (json.dumps(df.to_dict(orient='records')),))
                conn.commit()
                
                # पुराने फ़ाइल के डिलीट किए गए कॉलम्स की सेटिंग्स को रीसेट करना
                st.session_state["deleted_cols"] = [] 
                
                # 🧹 ट्रांसफर के बाद Entry Panel से फ़ाइल हटाना (अपलोडर रीसेट)
                st.session_state["p1_uploader_gen"] += 1
                st.session_state.pop("p1_sheet_select", None)
                st.session_state["p1_transfer_done"] = True
                st.rerun()
        except Exception as e:
            st.error(f"त्रुटि: {e}")

# =========================================================================
# 💻 PANEL 2: WORK / APPROVE PANEL (कॉलम मूव + लाइव स्प्लिट + डेटाबेस रूटिंग)
# =========================================================================
elif active_panel == "💻 2. Work / Approve Panel":
    st.title("💻 Work / Approve Panel - डेटा प्रोसेसिंग एवं अप्रूवल")
    if is_panel_hidden("p2") and not is_admin_session:
        st.warning("🔒 यह पैनल फिलहाल Admin द्वारा Hide किया गया है। डेटा उपलब्ध नहीं है।")
        st.dataframe(pd.DataFrame(), use_container_width=True)
        st.stop()
    
    # Panel 1 से ट्रांसफर होकर आया हुआ Staging (Raw) डेटा लोड करना
    raw_df = load_raw_data()
    
    if raw_df is None or raw_df.empty:
        st.info("📥 वर्तमान में कोई नई अपलोड की गई फ़ाइल पेंडिंग नहीं है। कृपया पहले 'Entry / Upload Panel (P1)' से फ़ाइल अपलोड करें।")
    else:
        # --- कार्य 1: बेकार कॉलम हटाना ---
        st.subheader("🗑️ बेकार कॉलम हटाएं (Remove Unwanted Columns)")
        active_cols = [c for c in raw_df.columns if c not in st.session_state["deleted_cols"]]
        
        cols_to_delete = st.multiselect("हटाने के लिए अनुपयोगी कॉलम चुनें:", options=active_cols)
        if cols_to_delete:
            if st.button("🔴 चुने गए कॉलम हटाएं", key="danger_delete_cols_btn"):
                st.session_state["deleted_cols"].extend(cols_to_delete)
                st.success("चयनित कॉलम स्क्रीन से हटा दिए गए!")
                st.rerun()
        
        final_raw_df = raw_df[active_cols]

        # --- कार्य 2: 🔄 कॉलमों का क्रम बदलना (Left/Right Move Feature) ---
        st.divider()
        st.subheader("🔄 कॉलमों का क्रम बदलें (Move Columns Left/Right)")
        st.write("नीचे दिए गए बॉक्स में क्रम बदलकर कॉलम को आगे-पीछे सेट करें। अप्रूवल के बाद इसी क्रम में लिस्ट लॉक होगी:")
        
        reordered_cols = st.multiselect(
            "कॉलमों का नया क्रम तय करें (सभी आवश्यक कॉलम इसी क्रम में चुनें):",
            options=active_cols,
            default=active_cols,
            key="col_reorder_select"
        )
        
        missing_cols = [c for c in active_cols if c not in reordered_cols]
        if missing_cols:
            reordered_cols.extend(missing_cols)
            
        final_raw_df = final_raw_df[reordered_cols]
        
        deg_col = next((c for c in final_raw_df.columns if any(k in c.lower() for k in ['deg', 'course', 'class'])), final_raw_df.columns)
        br_col = next((c for c in final_raw_df.columns if any(k in c.lower() for k in ['branch', 'stream', 'subject'])), final_raw_df.columns if len(final_raw_df.columns) > 1 else final_raw_df.columns)
        
        # --- कार्य 3: विशिष्ट डिग्री / ब्रांच की पूरी रो डिलीट करना ---
        st.divider()
        st.subheader("❌ विशिष्ट डिग्री / ब्रांच की पूरी रो डिलीट करें")
        st.write("यदि आप किसी खास कोर्स या स्ट्रीम का पूरा डेटा हटाना चाहते हैं, तो यहाँ से चुनें:")
        
        c_row1, c_row2 = st.columns(2)
        with c_row1:
            unique_degrees = final_raw_df[deg_col].dropna().unique().tolist()
            selected_degs = st.multiselect("डिलीट करने के लिए डिग्री (Course) चुनें:", options=unique_degrees)
        with c_row2:
            unique_branches = final_raw_df[br_col].dropna().unique().tolist()
            selected_branches = st.multiselect("डिलीट करने के लिए ब्रांच (Stream) चुनें:", options=unique_branches)
            
        if selected_degs or selected_branches:
            if st.button("🗑️ चुनी हुई रोज़ हमेशा के लिए डिलीट करें", key="danger_delete_rows_btn"):
                filtered_rows = []
                for _, row in raw_df.iterrows():
                    match_deg = str(row[deg_col]) in selected_degs if selected_degs else False
                    match_br = str(row[br_col]) in selected_branches if selected_branches else False
                    if not (match_deg or match_br):
                        filtered_rows.append(row.to_dict())
                
                cursor.execute("DELETE FROM raw_store")
                if filtered_rows:
                    cursor.execute("INSERT INTO raw_store (data_json) VALUES (?)", (json.dumps(filtered_rows),))
                conn.commit()
                st.success("🎉  चयनित डिग्री/ब्रांच की सभी रोज़ को सफलतापूर्वक डिलीट कर दिया गया है!")
                st.rerun()

        # फ़िल्टर्ड और रीऑर्डर किए गए डेटा का लाइव प्रीव्यू दिखाना
        st.divider()
        st.subheader("📋 अपलोड किए गए रॉ डेटा का लाइव प्रीव्यू (संशोधित क्रम)")
        st.dataframe(final_raw_df, height=350, use_container_width=True)
        
        el_col = next((c for c in final_raw_df.columns if any(k in c.lower() for k in ['elig', 'qual', 'class', 'course', 'deg'])), final_raw_df.columns)
        st.info(f"🔍 सिस्टम ऑटो-वर्गीकरण के लिए **'{el_col}'** कॉलम का उपयोग कर रहा है।")
        
        st.subheader("👀 लाइव प्री-विभाजन समीक्षा (Live Split Preview)")
        ug_preview_rows = []
        pg_preview_rows = []
        
        for _, row in final_raw_df.iterrows():
            if is_pg_route_value(row[el_col]):
                pg_preview_rows.append(row.to_dict())
            else:
                ug_preview_rows.append(row.to_dict())
                
        df_ug_preview = pd.DataFrame(ug_preview_rows)
        df_pg_preview = pd.DataFrame(pg_preview_rows)
        
        prev_tab1, prev_tab2 = st.tabs([f"🎓 UG में जाने वाला डेटा ({len(df_ug_preview)} रोज़)", f"📜 PG में जाने वाला डेटा ({len(df_pg_preview)} रोज़)"])
        
        with prev_tab1:
            if not df_ug_preview.empty: st.dataframe(df_ug_preview, height=250, use_container_width=True)
            else: st.caption("कोई डेटा UG श्रेणी में नहीं मिला।")
        with prev_tab2:
            if not df_pg_preview.empty: st.dataframe(df_pg_preview, height=250, use_container_width=True)
            else: st.caption("कोई डेटा PG श्रेणी में नहीं मिला।")

        # --- कार्य 4: फाइनल अप्रूवल और रूटिंग एक्शन (नया क्रम डेटाबेस में लॉक होगा) ---
        st.divider()
        st.subheader("🚀 FINAL ACTION")
        st.write("📈 **डेटा ट्रांसफर:** क्लीन और रीऑर्डर किए गए छात्रों के डेटा को आगे UG (P3) और PG (P4) पैनल में भेजने के लिए यह बटन दबाएँ।")
        
        if st.button("✅ डेटा अप्रूव करें और पैनल्स में ट्रांसफर करें", key="approve_transfer_all_btn"):
            if not df_ug_preview.empty:
                ug_json_str = json.dumps(df_ug_preview[reordered_cols].to_dict(orient='records'))
                cursor.execute("INSERT INTO perma_store (data_json, course_type) VALUES (?, ?)", (ug_json_str, "UG"))
            if not df_pg_preview.empty:
                pg_json_str = json.dumps(df_pg_preview[reordered_cols].to_dict(orient='records'))
                cursor.execute("INSERT INTO perma_store (data_json, course_type) VALUES (?, ?)", (pg_json_str, "PG"))
            
            cursor.execute("DELETE FROM raw_store")
            conn.commit()
            st.success("🎉 बधाई हो! डेटा सफलतापूर्वक कस्टमाइज्ड क्रम में ट्रांसफर और लॉक कर दिया गया है।")
            st.balloons()
            st.rerun()

elif active_panel == "🎓 3. UG Panel":
    st.title("🎓 Undergraduate (UG) चेकिंग एवं त्रुटि सुधार पैनल")
    if is_panel_hidden("p3") and not is_admin_session:
        st.warning("🔒 यह पैनल फिलहाल Admin द्वारा Hide किया गया है। डेटा उपलब्ध नहीं है।")
        st.dataframe(pd.DataFrame(), use_container_width=True)
        st.stop()
    df_ug = load_permanent_data("UG")
    
    if df_ug is None or df_ug.empty: 
        st.info("ℹ️ UG डेटाबेस खाली है। कृपया पहले Panel 2 से डेटा अप्रूव करें।")
    else:
        # ℹ️ UG डेटाबेस की वो डिग्री जो इस पैनल की डिग्री-सूची में नहीं आतीं (इसलिए यहाँ नहीं दिखेंगी)
        _ug_dc, _ = detect_deg_branch_cols(df_ug)
        if _ug_dc:
            _hidden_ug = sorted({
                d for d in df_ug[_ug_dc].map(_norm_blank_label).unique()
                if not any(k in d.lower().replace(".", "").replace(" ", "") for k in UG_ALLOWED_KEYWORDS)
            })
            if _hidden_ug:
                _pg_like = [d for d in _hidden_ug if is_pg_route_value(d)]
                _msg = ("ℹ️ UG डेटाबेस में ये डिग्री हैं पर इस पैनल की डिग्री-सूची (BA, B.Sc., B.Com., B.H.Sc., BBA, BCA...) में नहीं आतीं, "
                        "इसलिए इस पैनल की टेबल में नहीं दिखेंगी (Dashboard में दिखेंगी): " + ", ".join(_hidden_ug))
                if _pg_like:
                    _msg += (f"\n\nइनमें {', '.join(_pg_like)} PG की डिग्री लगती है — यह पुराना डेटा है जो पहले UG में चला गया था। "
                             "अब P2 में नई फ़ाइल आने पर ऐसी डिग्री अपने-आप PG में जाएगी।")
                st.info(_msg)

        # ऑटो-कॉलम डिटेक्शन
        minor_col = next((c for c in df_ug.columns if 'minor' in c.lower()), None)
        mdc_col = next((c for c in df_ug.columns if 'mdc' in c.lower()), None)
        voc_col = next((c for c in df_ug.columns if 'voc' in c.lower() or 'skill' in c.lower()), None)
        pw_col = next((c for c in df_ug.columns if any(k in c.lower() for k in ['pw', 'project', 'ce'])), None)
        
        # ड्रॉपडाउन में दिखाने के लिए यूनिक लिस्ट
        opt_minor = df_ug[minor_col].dropna().unique().tolist() if minor_col else []
        opt_mdc = df_ug[mdc_col].dropna().unique().tolist() if mdc_col else []
        opt_voc = df_ug[voc_col].dropna().unique().tolist() if voc_col else []
        opt_pw = df_ug[pw_col].dropna().unique().tolist() if pw_col else []
        
        st.markdown("### 🛠️ स्टेप 1: डिग्री-वाइज मास्टर गाइडलाइन सेट करें")
        st.caption("नीचे दी गई प्रत्येक डिग्री के बॉक्स को खोलकर उसके मान्य विषय चुनें और फिर 'लॉक करें' बटन दबाएं।")
        
        # --- डेटाबेस से पहले से सेव नियमों को सुरक्षित लोड करना ---
        ug_master_rules = {}
        try:
            cursor.execute("SELECT rules_json FROM locked_rules WHERE panel_prefix = 'ug_master'")
            locked_row = cursor.fetchone()
            if locked_row and locked_row[0]:
                ug_master_rules = json.loads(locked_row[0])
        except Exception as e:
            pass

        # आपकी मांगी गई 6 विशिष्ट डिग्रियां
        target_degrees = ["BA", "B.Sc.", "B.Sc. Biotechnology", "B.H.Sc.", "B.Com.", "B.Com. Computer"]
        current_configured_rules = {}
        
        for deg in target_degrees:
            with st.expander(f"📘 {deg} के लिए वैध विषय नियम (Valid Subjects)"):
                c1, c2, c3, c4 = st.columns(4)
                
                # पहले से सेव नियमों को ड्रॉपडाउन में डिफ़ॉल्ट दिखाना
                saved_deg_rule = ug_master_rules.get(deg, {})
                default_min = [x for x in saved_deg_rule.get("minor", []) if x in opt_minor]
                default_mdc = [x for x in saved_deg_rule.get("mdc", []) if x in opt_mdc]
                default_voc = [x for x in saved_deg_rule.get("voc", []) if x in opt_voc]
                default_pw = [x for x in saved_deg_rule.get("pw", []) if x in opt_pw]
                
                with c1:
                    r_minor = st.multiselect(f"Valid Minor", opt_minor, default=default_min, key=f"ug_min_{deg}")
                with c2:
                    r_mdc = st.multiselect(f"Valid MDC", opt_mdc, default=default_mdc, key=f"ug_mdc_{deg}")
                with c3:
                    r_voc = st.multiselect(f"Valid Vocational", opt_voc, default=default_voc, key=f"ug_voc_{deg}")
                with c4:
                    r_pw = st.multiselect(f"Valid PW/Ap/CE", opt_pw, default=default_pw, key=f"ug_pw_{deg}")
                    
                current_configured_rules[deg] = {
                    "minor": r_minor,
                    "mdc": r_mdc,
                    "voc": r_voc,
                    "pw": r_pw
                }
        
        # नियमों को डेटाबेस में लॉक करने का बटन
        if st.button("🔒 UG मास्टर विषय नियमावली लॉक करें", key="lock_master_ug_btn"):
            try:
                cursor.execute("""
                    INSERT INTO locked_rules (panel_prefix, rules_json) 
                    VALUES (?, ?)
                    ON CONFLICT(panel_prefix) DO UPDATE SET rules_json = excluded.rules_json
                """, ("ug_master", json.dumps(current_configured_rules)))
                conn.commit()
                st.success("🎉 सभी डिग्रियों के नियम डेटाबेस में सुरक्षित हो गए हैं और नीचे की लिस्ट रंगीन हो गई है!")
                st.rerun()
            except Exception as e:
                st.error(f"त्रुटि: {e}")
            
        st.divider()
        
        # लाइव वैलिडेशन टेबल रन करना (डेटाबेस से लोड किए गए नियमों को प्राथमिकता दें)
        rules_to_apply = ug_master_rules if ug_master_rules else current_configured_rules
        allowed_ug = list(UG_ALLOWED_KEYWORDS)
        
        process_panel_validation(df_ug, "ug", allowed_ug, master_rules=rules_to_apply)

# =========================================================================
# 📜 PANEL 4: PG PANEL
# =========================================================================
elif active_panel == "📜 4. PG Panel":
    st.title("📜 Postgraduate (PG) चेकिंग एवं त्रुटि सुधार पैनल")
    if is_panel_hidden("p4") and not is_admin_session:
        st.warning("🔒 यह पैनल फिलहाल Admin द्वारा Hide किया गया है। डेटा उपलब्ध नहीं है।")
        st.dataframe(pd.DataFrame(), use_container_width=True)
        st.stop()
    df_pg = load_permanent_data("PG")
    if df_pg is None or df_pg.empty: 
        st.info("ℹ️ PG डेटाबेस खाली है।")
    else: 
        process_panel_validation(df_pg, "pg", list(PG_ROUTE_KEYWORDS))

# =========================================================================
# 📊 PANEL 5: DASHBOARD / COUNTER PANEL (फुल स्क्रीन व्यूअर - भाग 1 और भाग 2 आपस में हाइड/शो)
# =========================================================================
elif active_panel == "📊 5. Dashboard / Counter Panel":
    # शीर्षक का फ़ॉन्ट छोटा किया गया है
    st.markdown("### 📊 Dashboard - डिग्री-वाइज लाइव काउंटर एवं विस्तृत डेटा समीक्षा")
    if is_panel_hidden("p5") and not is_admin_session:
        st.warning("🔒 यह पैनल फिलहाल Admin द्वारा Hide किया गया है। डेटा उपलब्ध नहीं है।")
        st.dataframe(pd.DataFrame(), use_container_width=True)
        st.stop()
    st.write("नीचे दिए गए टैब पर क्लिक करें, फिर बटन चुनकर 'विषय समरी' या 'छात्रों की फुल लिस्ट' को पूरी स्क्रीन पर देखें।")
    
    # 🛠️ फिक्स: UG और PG दोनों का डेटा लोड करना और लिस्ट बनाना
    df_ug_all = load_permanent_data("UG")
    df_pg_all = load_permanent_data("PG")
    
    # वेरिएबल को सही ढंग से इनिशियलाइज़ करना (यह डिलीट होने से एरर आ रहा था)
    all_dfs = []
    if df_ug_all is not None and not df_ug_all.empty: 
        all_dfs.append(df_ug_all)
    if df_pg_all is not None and not df_pg_all.empty: 
        all_dfs.append(df_pg_all)
    
    # अब यह कंडीशन बिना किसी एरर के बिल्कुल सही रन होगी
    if not all_dfs:
        st.info("ℹ️ काउंट प्रदर्शित करने के लिए डेटाबेस में कोई डेटा उपलब्ध नहीं है। कृपया पहले Panel 2 से डेटा अप्रूव करें।")
    else:
        master_df = pd.concat(all_dfs, ignore_index=True)
        
        # ऑटो-कॉलम डिटेक्शन इंजन
        deg_col = next((c for c in master_df.columns if any(k in c.lower() for k in ['deg', 'course', 'class'])), None)
        br_col = next((c for c in master_df.columns if any(k in c.lower() for k in ['branch', 'stream', 'subject'])), None)
        minor_col = next((c for c in master_df.columns if 'minor' in c.lower()), None)
        mdc_col = next((c for c in master_df.columns if 'mdc' in c.lower()), None)
        voc_col = next((c for c in master_df.columns if 'voc' in c.lower() or 'skill' in c.lower()), None)
        pw_col = next((c for c in master_df.columns if any(k in c.lower() for k in ['pw', 'project', 'ce'])), None)
        
        # 🔒 डेटाबेस से UG मास्टर नियमों को सुरक्षित लोड करना
        ug_master_rules = {}
        try:
            cursor.execute("SELECT rules_json FROM locked_rules WHERE panel_prefix = 'ug_master'")
            locked_row = cursor.fetchone()
            if locked_row and locked_row[0]:
                ug_master_rules = json.loads(locked_row[0])
        except:
            pass

        # डिग्रियों की सूची की सटीक मैपिंग (P3 - UG Rules Panel जैसी ही 6 डिग्री/ब्रांच संरचना)
        target_degrees = [
            {"display": "UG", "scope": "UG"},
            {"display": "PG", "scope": "PG"},
            {"display": "BA", "keywords": ["ba"]},
            {"display": "B.Sc.", "keywords": ["bsc"], "exclude": ["biotech"]},
            {"display": "B.Sc. Biotechnology", "keywords": ["bsc", "biotech"]},
            {"display": "B.H.Sc.", "keywords": ["bhsc"]},
            {"display": "B.Com.", "keywords": ["bcom"], "exclude": ["computer"]},
            {"display": "B.Com. Computer", "keywords": ["bcom", "computer"]}
        ]

        # ---- UG / PG टैब के लिए helper: हर रो की डिग्री पहचानकर उसी के मास्टर नियम लगाना ----
        _degree_defs = [d for d in target_degrees if "keywords" in d]

        def _norm_txt(x):
            return str(x).strip().lower().replace(".", "").replace(" ", "")

        def _row_degree_name(row):
            _d = str(row[deg_col]) if deg_col and deg_col in row.index else ""
            _b = str(row[br_col]) if br_col and br_col in row.index else ""
            _v = (_d + " " + _b).lower().replace(".", "").replace(" ", "")
            _cands = []
            for _dd in _degree_defs:
                if all(k in _v for k in _dd["keywords"]) and not any(ex in _v for ex in _dd.get("exclude", [])):
                    _cands.append((-len(_dd["keywords"]), _v.find(_dd["keywords"][0]), _dd["display"]))
            if not _cands:
                return None
            _cands.sort()
            return _cands[0][2]

        # सभी डिग्रियों के लिए इंटरएक्टिव टैब्स
        tab_titles = [deg["display"] for deg in target_degrees]
        tabs = st.tabs(tab_titles)

        for index, deg_info in enumerate(target_degrees):
            with tabs[index]:
                st.markdown(f"## 🎓 {deg_info['display']} डैशबोर्ड बोर्ड")
                
                # छात्र सूची में से इस विशिष्ट डिग्री के छात्रों को फ़िल्टर करना
                _scope = deg_info.get("scope")
                _empty_rules = {"minor": [], "mdc": [], "voc": [], "pw": []}

                def _rules_for_row(row):
                    if _scope == "UG":
                        return ug_master_rules.get(_row_degree_name(row), _empty_rules) or _empty_rules
                    if _scope == "PG":
                        return _empty_rules  # PG के लिए अभी कोई मास्टर नियम सेट नहीं है
                    return ug_master_rules.get(deg_info['display'], _empty_rules) or _empty_rules

                if _scope == "UG":
                    df_deg_filtered = df_ug_all.reset_index(drop=True) if (df_ug_all is not None and not df_ug_all.empty) else pd.DataFrame()
                elif _scope == "PG":
                    df_deg_filtered = df_pg_all.reset_index(drop=True) if (df_pg_all is not None and not df_pg_all.empty) else pd.DataFrame()
                elif deg_col and deg_col in master_df.columns:
                    def match_degree(row):
                        # 🔧 फिक्स: Degree column + Branch column दोनों को मिलाकर चेक करना
                        # (Biotechnology / Commerce Computer जैसी ब्रांच अक्सर अलग Branch column में होती है)
                        deg_part = str(row[deg_col]) if deg_col else ""
                        br_part = str(row[br_col]) if br_col and br_col in master_df.columns else ""
                        v = (deg_part + " " + br_part).lower().replace(".", "").replace(" ", "").strip()
                        match = all(k in v for k in deg_info["keywords"])
                        if "exclude" in deg_info:
                            if any(ex in v for ex in deg_info["exclude"]):
                                match = False
                        return match
                    
                    df_deg_filtered = master_df[master_df.apply(match_degree, axis=1)].reset_index(drop=True)
                else:
                    df_deg_filtered = pd.DataFrame()

                if df_deg_filtered.empty:
                    st.warning(f"⚠️ डेटाबेस में `{deg_info['display']}` का कोई छात्र रिकॉर्ड नहीं मिला।")
                    continue
                
                # मास्टर नियम लोड करना
                deg_rule = ug_master_rules.get(deg_info['display'], {"minor":[], "mdc":[], "voc":[], "pw":[]})

                # ✨ जादुई टॉगल बटन: यह तय करेगा कि भाग 1 देखना है या भाग 2
                view_option = st.radio(
                    "देखने के लिए व्यू चुनें:",
                    options=["📈 भाग 1: विषय काउंटर समरी (Subject Summary Counters)", "📋 भाग 2: छात्रों की विस्तृत लिस्ट (Detailed Student List)"],
                    horizontal=True,
                    key=f"view_toggle_{deg_info['display']}"
                )
                
                st.divider()

                # -------------------------------------------------------------------------
                # 📈 केवल भाग 1 (विषय काउंटर समरी) - लाइव कलर कोडिंग और अमान्य विषय रेड फिक्स
                # -------------------------------------------------------------------------
                if view_option == "📈 भाग 1: विषय काउंटर समरी (Subject Summary Counters)":
                    st.markdown("### 📈 विषयों की लाइव स्थिति (Summary Counters)")
                    
                    # समरी के अंदर माइनर, एमडीसी स्विच करने के लिए हॉरिजॉन्टल बार
                    selected_category = st.radio(
                        "समीक्षा के लिए विषय प्रकार चुनें:",
                        options=["Minor (माइनर)", "MDC (एम.डी.सी.)", "Vocational (व्यवसायिक)", "Project/PW (परियोजना)"],
                        horizontal=True,
                        key=f"cat_selector_{deg_info['display']}"
                    )
                
                    # 🔒 डेटाबेस से P3 (UG Panel) के नियमों को बिल्कुल सही तरीके से लोड करना
                    current_deg_rules = {}
                    try:
                        cursor.execute("SELECT rules_json FROM locked_rules WHERE panel_prefix = 'ug_master'")
                        locked_row = cursor.fetchone()
                        if locked_row and locked_row[0]:
                            all_rules = json.loads(locked_row[0])
                            current_deg_rules = all_rules.get(deg_info['display'], {"minor":[], "mdc":[], "voc":[], "pw":[]})
                    except Exception as e:
                        pass
                
                    cat_mapping = {
                        "Minor (माइनर)": {"col_name": minor_col, "rule_key": "minor", "label": "विषय का नाम (Minor Subject)"},
                        "MDC (एम.डी.सी.)": {"col_name": mdc_col, "rule_key": "mdc", "label": "विषय का नाम (MDC Subject)"},
                        "Vocational (व्यवसायिक)": {"col_name": voc_col, "rule_key": "voc", "label": "विषय का नाम (Vocational Subject)"},
                        "Project/PW (परियोजना)": {"col_name": pw_col, "rule_key": "pw", "label": "प्रोजेक्ट प्रकार (Project Type)"}
                    }
                
                    current_cat = cat_mapping[selected_category]
                
                    if current_cat["col_name"] and current_cat["col_name"] in df_deg_filtered.columns:
                        # काउंट्स (फ्रीक्वेंसी) की लाइव गणना
                        counts = df_deg_filtered[current_cat["col_name"]].dropna().value_counts().reset_index()
                        counts.columns = ["Subject", "Count"]

                        if _scope:
                            # UG/PG टैब: हर छात्र की अपनी डिग्री के नियम से सही/गलत गिनना
                            _c_col = current_cat["col_name"]
                            _c_rk = current_cat["rule_key"]
                            _sub = df_deg_filtered[df_deg_filtered[_c_col].notna()].copy()

                            def _wrong_flag(r):
                                _vs = {_norm_txt(x) for x in _rules_for_row(r).get(_c_rk, [])}
                                return bool(_vs) and (_norm_txt(r[_c_col]) not in _vs)

                            _sub["_wrong"] = _sub.apply(_wrong_flag, axis=1)
                            counts = (
                                _sub.groupby(_c_col)["_wrong"].agg(["size", "sum"]).reset_index()
                                .sort_values("size", ascending=False).reset_index(drop=True)
                            )
                            counts.columns = ["Subject", "Count", "Wrong"]
                            counts["Wrong"] = counts["Wrong"].astype(int)
                        
                        if not counts.empty:
                            # P3 के लॉक नियमों से वैध विषयों का क्लीन सेट बनाना
                            valid_list = current_deg_rules.get(current_cat["rule_key"], [])
                            valid_set = {str(x).strip().lower().replace(".", "").replace(" ", "") for x in valid_list}
                            
                            # ✨ भाग 1 के लिए नया सख्त रो स्टाइलर इंजन (गलत विषय = चमकदार गाढ़ा लाल)
                            def row_styler(row):
                                if "Wrong" in row.index:
                                    if row["Wrong"] > 0:
                                        return ['background-color: #f8d7da; color: #721c24; font-weight: bold; border: 2px solid #dc3545;'] * len(row)
                                    return [''] * len(row)
                                sub_val = str(row["Subject"]).strip().lower().replace(".", "").replace(" ", "")
                                # यदि P3 में विषय चुने गए हैं और छात्र का विषय उसमें नहीं है, तो पूरी रो लाल होगी
                                if valid_set and (sub_val not in valid_set):
                                    return ['background-color: #f8d7da; color: #721c24; font-weight: bold; border: 2px solid #dc3545;'] * len(row)
                                return [''] * len(row)
                            
                            if current_cat["rule_key"] == "pw":
                                _sum_hdr_html = "<div class='mini-head'>Project Type Summary<br><span>(प्रोजेक्ट प्रकार की समरी सूची)</span></div>"
                            else:
                                _sum_hdr_html = "<div class='mini-head'>Subject Distribution Summary<br><span>(विषय आवंटन की समरी सूची)</span></div>"
                            _sum_key = "subsum_" + "".join(ch if ch.isalnum() else "_" for ch in deg_info["display"])
                            _show_summary = inline_section_toggle(_sum_key, _sum_hdr_html)
                            
                            # फुल स्क्रीन चौड़ाई (Width) के साथ काउंटर तालिका रेंडर करना
                            # 🔧 फिक्स: टेबल की ऊंचाई अब रोज़ की संख्या के हिसाब से खुद-ब-खुद सेट होगी,
                            # ताकि पूरी लिस्ट एक बार में दिखे और स्क्रॉल न करना पड़े
                            if _show_summary:
                                dynamic_height = min(38 * (len(counts) + 1) + 3, 2000)
                                st.dataframe(
                                    counts.style.apply(row_styler, axis=1), 
                                    hide_index=True, 
                                    use_container_width=True,
                                    height=dynamic_height,
                                    column_config={
                                        "Subject": st.column_config.TextColumn(label=current_cat["label"], width=600), 
                                        "Count": st.column_config.NumberColumn(label="छात्रों की संख्या (Count)", width=150),
                                        **({"Wrong": st.column_config.NumberColumn(label="🔴 गलत (Wrong)", width=150)} if _scope else {})
                                    }
                                )
                        else:
                            st.caption("इस श्रेणी में कोई डेटा उपलब्ध नहीं है।")
                    else:
                        st.caption("डेटाबेस में संबंधित कॉलम नहीं मिला।")

                    # ---------------------------------------------------------------------
                    # 🎓 केवल UG / PG टैब: डिग्री + ब्रांच-वाइज समरी (Minor + MDC + Voc + PW सभी एक साथ)
                    # ---------------------------------------------------------------------
                    if _scope:
                        st.divider()
                        st.markdown(
                            '<div class="sum-banner"><div class="t">🎓 डिग्री + ब्रांच-वाइज समरी — Minor + MDC + Voc + PW (सभी एक साथ)</div></div>',
                            unsafe_allow_html=True
                        )
                        st.caption("हर डिग्री और ब्रांच के सामने कुल छात्र, और Minor / MDC / Vocational / Project-PW हर श्रेणी में कितने सही (✅), गलत (🔴) और खाली (🔵) हैं। 🔴 लाल सेल = उस डिग्री के मास्टर नियम से मेल न खाने वाला विषय | 🔵 नीला सेल = डेटा खाली है (PG में अभी कोई मास्टर नियम नहीं है, इसलिए वहाँ लाल नहीं दिखेगा)। 📝 कारण कॉलम में हर श्रेणी का पूरा कारण लिखा आता है।")

                        _db = df_deg_filtered.copy()

                        # डिग्री / ब्रांच कॉलम इसी टैब के अपने डेटा से पहचानना (PG के कॉलम नाम अलग हो सकते हैं)
                        _dc = deg_col if (deg_col and deg_col in _db.columns) else next(
                            (c for c in _db.columns if any(k in str(c).lower() for k in ['deg', 'course', 'class'])), None)
                        _skip_kw = ['minor', 'mdc', 'voc', 'skill', 'pw', 'project']
                        _bc = br_col if (br_col and br_col in _db.columns and br_col != _dc) else None
                        if _bc is None:
                            _bc = next((c for c in _db.columns if c != _dc and any(k in str(c).lower() for k in ['branch', 'stream'])), None)
                        if _bc is None:
                            _bc = next((c for c in _db.columns if c != _dc and 'subject' in str(c).lower()
                                        and not any(k in str(c).lower() for k in _skip_kw)), None)

                        def _clean_key(series):
                            s_ = series.astype(str).str.strip()
                            return s_.mask(series.isna() | (s_ == "") | (s_.str.lower() == "nan"), "(खाली/Blank)")

                        _db["Degree (डिग्री)"] = _clean_key(_db[_dc]) if _dc else "—"
                        _db["Branch (ब्रांच)"] = _clean_key(_db[_bc]) if _bc else "—"

                        _deg_options = ["सभी डिग्री"] + sorted(_db["Degree (डिग्री)"].unique().tolist())
                        _deg_pick = st.selectbox(
                            "डिग्री फ़िल्टर:",
                            _deg_options,
                            key=f"db_deg_filter_{deg_info['display']}_all"
                        )
                        if _deg_pick != "सभी डिग्री":
                            _db = _db[_db["Degree (डिग्री)"] == _deg_pick]

                        _red = 'background-color: #f8d7da; color: #721c24; font-weight: bold; border: 2px solid #dc3545;'
                        _blue = 'background-color: #d1ecf1; color: #0c5460; font-weight: bold; border: 2px solid #17a2b8;'

                        # चारों श्रेणियाँ एक साथ: Minor / MDC / Voc / PW
                        _all_cats = [("Minor", "minor", minor_col), ("MDC", "mdc", mdc_col), ("Voc", "voc", voc_col), ("PW", "pw", pw_col)]
                        _present_cats = [(l_, k_, c_) for (l_, k_, c_) in _all_cats if c_ and c_ in _db.columns]
                        _missing_cats = [l_ for (l_, k_, c_) in _all_cats if not (c_ and c_ in _db.columns)]

                        if _present_cats:
                            if _missing_cats:
                                st.caption("ℹ️ इस डेटा में इन श्रेणियों का कॉलम नहीं मिला, इसलिए ये समरी में शामिल नहीं हैं: " + ", ".join(_missing_cats))

                            # हर छात्र की डिग्री के हिसाब से मास्टर-नियम वाली डिग्री (एक बार निकालना)
                            if _scope == "UG" and len(_db):
                                _db["_rd"] = _db.apply(_row_degree_name, axis=1)
                            else:
                                _db["_rd"] = None

                            # ⚪ खाली-छूट वाली डिग्री+ब्रांच (UG/PG की सेटिंग)
                            _ex_pairs = get_blank_exempt_pairs("ug" if _scope == "UG" else "pg")
                            _ex_s = pd.Series([(d_, b_) in _ex_pairs for d_, b_ in zip(_db["Degree (डिग्री)"], _db["Branch (ब्रांच)"])], index=_db.index)

                            # हर श्रेणी के लिए गलत / खाली फ्लैग
                            for _lbl, _rk, _cn in _present_cats:
                                _empty_s = _db[_cn].isna() | (_db[_cn].astype(str).str.strip() == "")
                                _blank_s = _empty_s & ~_ex_s     # ⚪ खाली-छूट वाली डिग्री+ब्रांच में खाली नहीं गिनना
                                _norm_s = _db[_cn].astype(str).map(_norm_txt)
                                _wrong_s = pd.Series(False, index=_db.index)
                                if _scope == "UG":
                                    for _rd_name in _db["_rd"].dropna().unique():
                                        _allowed = {_norm_txt(x) for x in ((ug_master_rules.get(_rd_name) or {}).get(_rk, []))}
                                        if _allowed:
                                            _wrong_s = _wrong_s | ((_db["_rd"] == _rd_name) & ~_empty_s & ~_norm_s.isin(_allowed))
                                _db[f"_w_{_rk}"] = _wrong_s
                                _db[f"_b_{_rk}"] = _blank_s

                            _first_rk = _present_cats[0][1]
                            _agg = {"_total": (f"_w_{_first_rk}", "size")}
                            for _lbl, _rk, _cn in _present_cats:
                                _agg[f"_w_{_rk}"] = (f"_w_{_rk}", "sum")
                                _agg[f"_b_{_rk}"] = (f"_b_{_rk}", "sum")
                            _grp = _db.groupby(["Degree (डिग्री)", "Branch (ब्रांच)"]).agg(**_agg).reset_index()

                            # 📝 कारण (Reason): हर श्रेणी का पूरा कारण (कौन सा विषय गलत, कितने खाली)
                            def _reason_for_group(g):
                                lines = []
                                for _lbl, _rk, _cn in _present_cats:
                                    parts = []
                                    w = g[g[f"_w_{_rk}"]]
                                    if len(w):
                                        for _rd_name, _wg in w.groupby("_rd"):
                                            _vc = _wg[_cn].astype(str).str.strip().value_counts()
                                            _subj = ", ".join(f"{k} ({v})" for k, v in _vc.items())
                                            parts.append(f"🔴 {len(_wg)} छात्रों का विषय {_rd_name} के मास्टर नियम में मान्य नहीं है → {_subj}")
                                    _bn = int(g[f"_b_{_rk}"].sum())
                                    if _bn:
                                        parts.append(f"🔵 {_bn} छात्रों का खाली है (डेटा नहीं भरा गया)")
                                    if parts:
                                        lines.append(f"{_lbl}: " + " ; ".join(parts))
                                return "  ||  ".join(lines)

                            _reasons = {}
                            for (_d_k, _b_k), _g in _db.groupby(["Degree (डिग्री)", "Branch (ब्रांच)"]):
                                _reasons[(_d_k, _b_k)] = _reason_for_group(_g)

                            db_summary = _grp[["Degree (डिग्री)", "Branch (ब्रांच)"]].copy()
                            db_summary["कुल छात्र (Total)"] = _grp["_total"].astype(int)
                            for _lbl, _rk, _cn in _present_cats:
                                _w_col = _grp[f"_w_{_rk}"].astype(int)
                                _b_col = _grp[f"_b_{_rk}"].astype(int)
                                db_summary[f"{_lbl} ✅ सही"] = db_summary["कुल छात्र (Total)"] - _w_col - _b_col
                                db_summary[f"{_lbl} 🔴 गलत"] = _w_col
                                db_summary[f"{_lbl} 🔵 खाली"] = _b_col
                            db_summary["📝 कारण (Reason)"] = [
                                _reasons.get((_d_k, _b_k), "")
                                for _d_k, _b_k in zip(_grp["Degree (डिग्री)"], _grp["Branch (ब्रांच)"])
                            ]
                            db_summary = db_summary.sort_values(
                                ["Degree (डिग्री)", "कुल छात्र (Total)"], ascending=[True, False]
                            ).reset_index(drop=True)
                            _db_core = db_summary.copy()   # मेट्रिक्स के लिए बिना Total रो वाली कॉपी
                            _tot = {c_: "" for c_ in db_summary.columns}
                            _tot["Degree (डिग्री)"] = "कुल योग (GRAND TOTAL)"
                            _tot["Branch (ब्रांच)"] = ""
                            for c_ in db_summary.columns:
                                if c_ not in ("Degree (डिग्री)", "Branch (ब्रांच)", "📝 कारण (Reason)"):
                                    _tot[c_] = int(db_summary[c_].sum())
                            db_summary = pd.concat([db_summary, pd.DataFrame([_tot])], ignore_index=True)

                            _wrong_cols = [c for c in db_summary.columns if "🔴" in c]
                            _blank_cols = [c for c in db_summary.columns if "🔵" in c]

                            def _db_cell_styler(row):
                                if str(row["Degree (डिग्री)"]).startswith("कुल योग"):
                                    return ['background-color: #1a3c6e; color: #ffffff; font-weight: bold;'] * len(row)
                                _any_w = any(row[c] > 0 for c in _wrong_cols)
                                _any_b = any(row[c] > 0 for c in _blank_cols)
                                out = []
                                for c in row.index:
                                    if c in _wrong_cols and row[c] > 0:
                                        out.append(_red)
                                    elif c in _blank_cols and row[c] > 0:
                                        out.append(_blue)
                                    elif c == "📝 कारण (Reason)" and row[c]:
                                        out.append(_red if _any_w else (_blue if _any_b else ""))
                                    else:
                                        out.append("")
                                return out

                            # 📄 शीट 2 के लिए: Reason वाले (गलत/खाली) छात्रों की व्यक्तिगत लिस्ट
                            _any_flag = pd.Series(False, index=_db.index)
                            for _lbl, _rk, _cn in _present_cats:
                                _any_flag = _any_flag | _db[f"_w_{_rk}"] | _db[f"_b_{_rk}"]
                            _stu = _db[_any_flag].sort_values(["Degree (डिग्री)", "Branch (ब्रांच)"], kind="stable")
                            _orig_cols = [c for c in df_deg_filtered.columns if c in _stu.columns]

                            def _stu_reason(r):
                                out_ = []
                                for _lbl, _rk, _cn in _present_cats:
                                    if r[f"_w_{_rk}"]:
                                        _rd_ = r["_rd"] if r["_rd"] else "इस डिग्री"
                                        out_.append(f"🔴 {_lbl}: '{str(r[_cn]).strip()}' — {_rd_} के मास्टर नियम में मान्य नहीं")
                                    elif r[f"_b_{_rk}"]:
                                        out_.append(f"🔵 {_lbl}: खाली (डेटा नहीं भरा गया)")
                                return "\n".join(out_)

                            students_df = _stu[_orig_cols].copy().reset_index(drop=True)
                            students_df.insert(0, "क्र.सं.", range(1, len(students_df) + 1))
                            students_df["📝 कारण (Reason)"] = [_stu_reason(r_) for _, r_ in _stu.iterrows()]
                            student_flags = {}
                            for _lbl, _rk, _cn in _present_cats:
                                student_flags[_cn] = [
                                    "w" if w_ else ("b" if b_ else "")
                                    for w_, b_ in zip(_stu[f"_w_{_rk}"].tolist(), _stu[f"_b_{_rk}"].tolist())
                                ]

                            dm1, dm2, dm3 = st.columns(3)
                            with dm1:
                                st.metric("🎓 कुल डिग्री+ब्रांच कॉम्बिनेशन", len(_db_core))
                            with dm2:
                                st.metric("🔴 कुल गलत एंट्री (सभी श्रेणी)", int(_db_core[_wrong_cols].to_numpy().sum()))
                            with dm3:
                                st.metric("🔵 कुल खाली सेल (सभी श्रेणी)", int(_db_core[_blank_cols].to_numpy().sum()))

                            st.dataframe(
                                db_summary.style.apply(_db_cell_styler, axis=1),
                                hide_index=True,
                                use_container_width=True,
                                height=min(38 * (len(db_summary) + 1) + 3, 650),
                                column_config={"📝 कारण (Reason)": st.column_config.TextColumn(width=900)}
                            )

                            # 📝 लंबा कारण टेबल के सेल में कट सकता है, इसलिए पूरा कारण यहाँ भी दिखाना
                            _reason_rows = db_summary[db_summary["📝 कारण (Reason)"] != ""]
                            if not _reason_rows.empty:
                                with st.expander(f"📝 पूरा कारण देखें (लाल/नीली {len(_reason_rows)} ब्रांच)"):
                                    for _, _rr in _reason_rows.iterrows():
                                        _rtxt = str(_rr["📝 कारण (Reason)"])
                                        _lines_html = ""
                                        for _ln in _rtxt.split("  ||  "):
                                            _lc = "red" if "🔴" in _ln else "blue"
                                            _lines_html += f'<div class="rs-line {_lc}">{_html.escape(_ln).replace("$", "&#36;")}</div>'
                                        _cc = "red" if "🔴" in _rtxt else "blue"
                                        _ttl = _html.escape(f"{_rr['Degree (डिग्री)']} → {_rr['Branch (ब्रांच)']}").replace("$", "&#36;")
                                        st.markdown(
                                            f'<div class="reason-card {_cc}"><div class="rc-title">{_ttl}</div>{_lines_html}</div>',
                                            unsafe_allow_html=True
                                        )
                        else:
                            # इस डेटा में Minor/MDC/Voc/PW में से कोई कॉलम नहीं है — तब भी डिग्री+ब्रांच की गिनती दिखाना
                            students_df, student_flags = None, None
                            st.info("ℹ️ इस डेटा में Minor / MDC / Vocational / Project-PW में से किसी का कॉलम नहीं मिला, इसलिए नीचे सिर्फ डिग्री + ब्रांच के हिसाब से छात्रों की कुल संख्या दिखाई जा रही है।")
                            st.caption("इस डेटा में उपलब्ध कॉलम: " + ", ".join(str(c) for c in df_deg_filtered.columns))
                            db_summary = (
                                _db.groupby(["Degree (डिग्री)", "Branch (ब्रांच)"]).size().reset_index(name="कुल छात्र (Total)")
                                .sort_values(["Degree (डिग्री)", "कुल छात्र (Total)"], ascending=[True, False]).reset_index(drop=True)
                            )
                            dm1, dm2 = st.columns(2)
                            with dm1:
                                st.metric("🎓 कुल डिग्री+ब्रांच कॉम्बिनेशन", len(db_summary))
                            with dm2:
                                st.metric("👥 कुल छात्र", int(db_summary["कुल छात्र (Total)"].sum()))
                            db_summary = pd.concat([db_summary, pd.DataFrame([{
                                "Degree (डिग्री)": "कुल योग (GRAND TOTAL)", "Branch (ब्रांच)": "",
                                "कुल छात्र (Total)": int(db_summary["कुल छात्र (Total)"].sum())}])], ignore_index=True)
                            st.dataframe(
                                db_summary,
                                hide_index=True,
                                use_container_width=True,
                                height=min(38 * (len(db_summary) + 1) + 3, 650)
                            )

                        # 📥 डाउनलोड पैनल (CSV + 2-शीट Excel)
                        _n_reason_stu = 0 if students_df is None else len(students_df)
                        _chips = f'<span class="dl-chip">📑 {len(db_summary) - 1 if len(db_summary) > 1 else len(db_summary)} डिग्री+ब्रांच रो</span>'
                        if students_df is not None:
                            _chips += f'<span class="dl-chip red">👥 Reason वाले छात्र: {_n_reason_stu}</span>'
                        st.markdown(
                            f'<div class="dl-panel-head"><span class="h">📥 समरी डाउनलोड करें</span>{_chips}</div>',
                            unsafe_allow_html=True
                        )
                        st.markdown(
                            '<div class="dl-hint">📄 <b>CSV</b> — सिर्फ़ यह समरी टेबल (Total रो सहित)<br>'
                            '📊 <b>Excel</b> — रंगीन, <b>2 शीट</b>: शीट 1 = समरी, शीट 2 = Reason वाले छात्रों की पूरी लिस्ट</div>',
                            unsafe_allow_html=True
                        )
                        _dl1, _dl2 = st.columns(2)
                        with _dl1:
                            st.download_button(
                                label="📄  CSV डाउनलोड  (सिर्फ़ समरी)",
                                data=db_summary.to_csv(index=False).encode("utf-8-sig"),
                                file_name=f"{deg_info['display']}_Degree_Branch_All_Categories_Summary.csv",
                                mime="text/csv",
                                key=f"db_dl_{deg_info['display']}_all",
                                use_container_width=True
                            )
                        with _dl2:
                            st.download_button(
                                label="📊  Excel डाउनलोड  (2 शीट: समरी + Reason वाले छात्र)",
                                data=generate_summary_excel_bytes(
                                    db_summary,
                                    sheet_name=f"{deg_info['display']}_Summary",
                                    students_df=students_df,
                                    student_flags=student_flags,
                                    students_sheet_name="Reason_Students"
                                ),
                                file_name=f"{deg_info['display']}_Degree_Branch_All_Categories_Summary.xlsx",
                                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                key=f"db_dl_xlsx_{deg_info['display']}_all",
                                use_container_width=True
                            )

                        # ---------------------------------------------------------------------
                        # 📑 ब्रांच-वाइज विषय शीट: Degree/Branch/Total Admission मर्ज,
                        # नीचे हर रो में Minor / MDC / Voc / PW के अलग-अलग नाम + count
                        # ---------------------------------------------------------------------
                        if _present_cats:
                            st.divider()
                            st.markdown(
                                '<div class="sum-banner"><div class="t">📑 ब्रांच-वाइज विषय शीट — Total Admission + Minor / MDC / Voc / PW के नाम व संख्या</div></div>',
                                unsafe_allow_html=True
                            )
                            st.caption("हर Degree/Branch के लिए Degree, Branch और Total Admission एक बार (मर्ज) दिखते हैं, और उसके नीचे की रो में हर श्रेणी के अलग-अलग विषय अपनी संख्या के साथ आते हैं। जिस श्रेणी में विषय कम हैं वहाँ बाकी सेल खाली रहते हैं। (ऊपर चुना गया डिग्री फ़िल्टर यहाँ भी लागू है)")
                            _bs_cats = [(l_, c_) for (l_, k_, c_) in _present_cats]
                            _bs_labels = [l_ for l_, _c in _bs_cats]
                            _bs_blocks = build_branch_subject_blocks(_db, "Degree (डिग्री)", "Branch (ब्रांच)", _bs_cats)
                            st.markdown(branch_subject_sheet_html(_bs_blocks, _bs_labels), unsafe_allow_html=True)
                            st.download_button(
                                label="📊  Excel डाउनलोड  (ब्रांच-वाइज विषय शीट, merged)",
                                data=generate_branch_subject_sheet_excel_bytes(_bs_blocks, _bs_labels, sheet_name=f"{deg_info['display']}_Branch_Subjects"),
                                file_name=f"{deg_info['display']}_Branch_Subject_Sheet.xlsx",
                                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                key=f"db_dl_bs_{deg_info['display']}_all",
                                use_container_width=True
                            )

                # -------------------------------------------------------------------------
                # 📋 केवल भाग 2 (छात्रों की विस्तृत लिस्ट) - 1 से शुरू होने वाला सीरियल नंबर फिक्स
                # -------------------------------------------------------------------------
                elif view_option == "📋 भाग 2: छात्रों की विस्तृत लिस्ट (Detailed Student List)":
                    st.markdown(f"### 📋 {deg_info['display']} के सभी छात्रों का विस्तृत डेटा")
                    
                    df_to_show = df_deg_filtered.copy()
                    
                    # 🔍 लाइव सर्च बार फीचर
                    search_query = st.text_input(
                        f"🔍 {deg_info['display']} डेटा में सर्च करें (नाम, रोल नंबर या विषय डालें):", 
                        key=f"search_{deg_info['display']}"
                    )
                    
                    if search_query:
                        mask = df_to_show.astype(str).apply(lambda row: row.str.contains(search_query, case=False).any(), axis=1)
                        df_to_show = df_to_show[mask]
                
                    # --- 🛠️ सटीक लाइव कॉलम डिटेक्शन ---
                    actual_minor = next((c for c in df_to_show.columns if 'minor' in c.lower()), None)
                    actual_mdc = next((c for c in df_to_show.columns if 'mdc' in c.lower()), None)
                    actual_voc = next((c for c in df_to_show.columns if 'voc' in c.lower() or 'skill' in c.lower()), None)
                    actual_pw = next((c for c in df_to_show.columns if any(k in c.lower() for k in ['pw', 'project', 'ce'])), None)
                    actual_br = next((c for c in df_to_show.columns if any(k in c.lower() for k in ['branch', 'stream', 'subject'])), None)
                
                    # 🔒 डेटाबेस से P3 के नियमों को लोड करना
                    current_deg_rules = {}
                    try:
                        cursor.execute("SELECT rules_json FROM locked_rules WHERE panel_prefix = 'ug_master'")
                        locked_row = cursor.fetchone()
                        if locked_row and locked_row[0]:
                            all_rules = json.loads(locked_row[0])
                            current_deg_rules = all_rules.get(deg_info['display'], {"minor":[], "mdc":[], "voc":[], "pw":[]})
                    except Exception as e:
                        pass
                
                    # --- 📊 लाइव काउंटर मीटर ---
                    blank_count = 0
                    wrong_count = 0
                    
                    targets_for_counting = {
                        actual_minor: 'minor', 
                        actual_mdc: 'mdc', 
                        actual_voc: 'voc', 
                        actual_pw: 'pw'
                    }
                    
                    _ex_prefix = "ug" if _scope == "UG" else ("pg" if _scope == "PG" else None)
                    _ex_list = blank_exempt_mask(df_to_show, _ex_prefix).tolist()   # ⚪ खाली-छूट वाली रो
                    for _pos, (index, row) in enumerate(df_to_show.iterrows()):
                        _row_rules = _rules_for_row(row)
                        for col_name, rule_key in targets_for_counting.items():
                            if col_name and col_name in df_to_show.columns:
                                val = row[col_name]
                                if pd.isna(val) or str(val).strip() == "":
                                    if not _ex_list[_pos]:
                                        blank_count += 1
                                else:
                                    val_clean = str(val).strip().lower().replace(".", "").replace(" ", "")
                                    valid_list = _row_rules.get(rule_key, [])
                                    valid_set = {str(x).strip().lower().replace(".", "").replace(" ", "") for x in valid_list}
                                    if valid_set and (val_clean not in valid_set):
                                        wrong_count += 1
                
                    # स्क्रीन पर लाइव स्टेट्स कार्ड्स दिखाना
                    metric_c1, metric_c2, metric_c3 = st.columns(3)
                    with metric_c1:
                        st.metric(label="👥 कुल छात्र रिकॉर्ड (Total Rows)", value=len(df_to_show))
                    with metric_c2:
                        st.metric(label="🔵 कुल खाली सेल (Missing Data)", value=blank_count)
                    with metric_c3:
                        st.metric(label="🔴 कुल गलत विषय (Rule Mismatch)", value=wrong_count)
                
                    st.caption("🔵 नीला सेल = डेटा गायब है | 🔴 लाल सेल = गलत विषय (Panel 3 में आपके द्वारा चुने गए विषयों के अलावा बाकी सब)")
                
                    # ✨ जादू यहाँ है: टेबल दिखाने से पहले इंडेक्स को 1 से शुरू करने के लिए शिफ्ट करना
                    df_to_show.index = range(1, len(df_to_show) + 1)
                
                    # --- 🎨 लाइव कलर कोडिंग स्टाइलर इंजन ---
                    def full_table_styler(dataframe):
                        s_df = pd.DataFrame('', index=dataframe.index, columns=dataframe.columns)
                        targets = {
                            actual_minor: 'minor', 
                            actual_mdc: 'mdc', 
                            actual_voc: 'voc', 
                            actual_pw: 'pw'
                        }
                        for _pos, (index, row) in enumerate(dataframe.iterrows()):
                            row_has_wrong = False
                            _row_rules = _rules_for_row(row)
                            for col_name, rule_key in targets.items():
                                if col_name and col_name in dataframe.columns:
                                    val = row[col_name]
                                    if pd.isna(val) or str(val).strip() == "":
                                        if not _ex_list[_pos]:
                                            s_df.at[index, col_name] = 'background-color: #d1ecf1; color: #0c5460; font-weight: bold; border: 2px solid #17a2b8;'
                                    else:
                                        val_clean = str(val).strip().lower().replace(".", "").replace(" ", "")
                                        valid_list = _row_rules.get(rule_key, [])
                                        valid_set = {str(x).strip().lower().replace(".", "").replace(" ", "") for x in valid_list}
                                        if valid_set and (val_clean not in valid_set):
                                            s_df.at[index, col_name] = 'background-color: #f8d7da; color: #721c24; font-weight: bold; border: 2px solid #dc3545;'
                                            row_has_wrong = True
                            # 🟡 फिक्स: अगर इस रो में कोई भी सेल (Minor/MDC/Voc/PW) लाल (गलत) है,
                            # तो उसी रो के Branch सेल को पीला (Yellow) कर देना
                            if row_has_wrong and actual_br and actual_br in dataframe.columns:
                                s_df.at[index, actual_br] = 'background-color: #fff3cd; color: #856404; font-weight: bold; border: 2px solid #ffc107;'
                        return s_df
                
                    # full screen view table render
                    st.dataframe(
                        df_to_show.style.apply(full_table_styler, axis=None),
                        height=550,
                        use_container_width=True
                    )

                    # -------------------------------------------------------------------------
                    # 🟡 ब्रांच-वाइज सही/गलत समरी (जिस ब्रांच में कोई गलत छात्र है वो पीली दिखेगी)
                    # -------------------------------------------------------------------------
                    if actual_br and actual_br in df_to_show.columns:
                        st.divider()
                        st.markdown("### 🟡 ब्रांच-वाइज सही / गलत समरी")
                        st.caption("हर ब्रांच के सामने कुल कितने छात्र हैं, उनमें से कितने सही (✅) हैं और कितने गलत (🔴) हैं — जिस ब्रांच में कम-से-कम एक गलत छात्र है, वो रो पीली (Yellow) दिखेगी।")

                        targets_for_branch_summary = {
                            actual_minor: 'minor',
                            actual_mdc: 'mdc',
                            actual_voc: 'voc',
                            actual_pw: 'pw'
                        }

                        def _row_has_wrong_subject(row):
                            _row_rules = _rules_for_row(row)
                            for col_name, rule_key in targets_for_branch_summary.items():
                                if col_name and col_name in df_to_show.columns:
                                    val = row[col_name]
                                    if not (pd.isna(val) or str(val).strip() == ""):
                                        val_clean = str(val).strip().lower().replace(".", "").replace(" ", "")
                                        valid_list = _row_rules.get(rule_key, [])
                                        valid_set = {str(x).strip().lower().replace(".", "").replace(" ", "") for x in valid_list}
                                        if valid_set and (val_clean not in valid_set):
                                            return True
                            return False

                        branch_summary_rows = []
                        for branch_val, group in df_to_show.groupby(df_to_show[actual_br].fillna("(खाली/Blank)").replace("", "(खाली/Blank)")):
                            total_n = len(group)
                            wrong_n = int(group.apply(_row_has_wrong_subject, axis=1).sum())
                            correct_n = total_n - wrong_n
                            branch_summary_rows.append({
                                "Branch (ब्रांच)": branch_val,
                                "कुल छात्र (Total)": total_n,
                                "✅ सही (Correct)": correct_n,
                                "🔴 गलत (Wrong)": wrong_n
                            })

                        branch_summary_df = pd.DataFrame(branch_summary_rows).sort_values(
                            "कुल छात्र (Total)", ascending=False
                        ).reset_index(drop=True)

                        def branch_summary_row_styler(row):
                            if row["🔴 गलत (Wrong)"] > 0:
                                return ['background-color: #fff3cd; color: #856404; font-weight: bold; border: 1px solid #ffc107;'] * len(row)
                            return ['background-color: #d4edda; color: #155724; font-weight: bold; border: 1px solid #28a745;'] * len(row)

                        branch_dyn_height = min(38 * (len(branch_summary_df) + 1) + 3, 1500)
                        st.dataframe(
                            branch_summary_df.style.apply(branch_summary_row_styler, axis=1),
                            hide_index=True,
                            use_container_width=True,
                            height=branch_dyn_height
                        )

                        # -------------------------------------------------------------------------
                        # ✅ Approve किए गए छात्रों का डेटा (Panel 3/4 से जिन्हें Approve किया गया है)
                        # -------------------------------------------------------------------------
                        st.divider()
                        st.markdown("### ✅ Approve किए गए छात्रों का डेटा")
                        st.caption("जिन छात्रों का 'गलत विषय' Nodal / Student / Principal द्वारा Approve किया जा चुका है, वो नीचे दिखेंगे — इन्हें अब गलत नहीं गिना जाएगा।")

                        _dash_key_col = find_student_key_col(df_to_show)
                        _approved_hits = []
                        for _pos, (_idx, _row) in enumerate(df_to_show.iterrows()):
                            _skey = get_student_key(_row, _pos, _dash_key_col)
                            _appr = get_approval(_skey, "pg" if _scope == "PG" else "ug")
                            if _appr:
                                _approved_hits.append((_idx, _appr[0], _appr[1]))

                        if not _approved_hits:
                            st.info("इस डिग्री/ब्रांच में अभी तक कोई भी छात्र Approve नहीं हुआ है।")
                        else:
                            _approved_students_df = df_to_show.loc[[i for i, _, _ in _approved_hits]].copy()
                            _approved_students_df.insert(0, "✅ Approve किया (By)", [a[1] for a in _approved_hits])
                            _approved_students_df.insert(1, "🕒 Approve समय", [a[2] for a in _approved_hits])
                            st.dataframe(_approved_students_df, use_container_width=True, hide_index=True)
                            _dash_orientation = st.radio(
                                "🖨️ प्रिंट ओरिएंटेशन चुनें:",
                                options=["📄 Portrait (सीधा)", "📃 Landscape (आड़ा)"],
                                horizontal=True,
                                key=f"orientation_dash_{deg_info['display']}"
                            )
                            render_print_button(
                                _approved_students_df,
                                title=f"Approved Students List - {deg_info['display']}",
                                key_suffix=f"dash_{deg_info['display']}".replace(" ", "_").replace(".", ""),
                                orientation="landscape" if "Landscape" in _dash_orientation else "portrait"
                            )

        # =====================================================================
        # 📥 मास्टर एक्सेल डाउनलोड (P5 के सबसे नीचे) — ऊपर जितनी भी लिस्ट/डिग्री
        # बनी हैं (UG + PG दोनों), वो सब एक ही Excel फ़ाइल में 6 शीट के रूप में आ जाएँगी।
        # =====================================================================
        st.divider()
        st.markdown("### 📥 मास्टर एक्सेल डाउनलोड करें (6 शीट)")
        st.caption(
            "Sheet 1-2: UG/PG का पूरा रॉ डेटा (🔵 खाली | 🔴 गलत | 🟢 Approved)  |  "
            "Sheet 3: ब्रांच-वाइज विषय शीट (Total Admission + Minor/MDC/Voc/PW के नाम व संख्या)  |  "
            "Sheet 4: डिग्री+ब्रांच-वाइज समरी (Minor+MDC+Voc+PW सभी एक साथ)  |  "
            "Sheet 5: जिन छात्रों का कोई विषय गलत/खाली है, उनकी पूरी लिस्ट + कारण  |  "
            "Sheet 6: Sheet 5 में से जो पहले ही Approve हो चुके हैं, उनकी लिस्ट।"
        )

        _p5_ug_key_col = find_student_key_col(df_ug_all) if (df_ug_all is not None and not df_ug_all.empty) else None
        _p5_pg_key_col = find_student_key_col(df_pg_all) if (df_pg_all is not None and not df_pg_all.empty) else None
        _p5_ug_approved_keys = {a[0] for a in get_all_approvals("ug")}
        _p5_pg_approved_keys = {a[0] for a in get_all_approvals("pg")}

        _p5_master_sheets = []
        if df_ug_all is not None and not df_ug_all.empty:
            _p5_master_sheets.append({
                "df": df_ug_all, "sheet_name": "UG_Master_List",
                "deg_col": deg_col, "br_col": br_col, "minor_col": minor_col, "mdc_col": mdc_col,
                "voc_col": voc_col, "pw_col": pw_col, "master_rules": ug_master_rules,
                "approved_keys": _p5_ug_approved_keys, "key_col": _p5_ug_key_col, "prefix": "ug"
            })
        if df_pg_all is not None and not df_pg_all.empty:
            _p5_master_sheets.append({
                "df": df_pg_all, "sheet_name": "PG_Master_List",
                "deg_col": deg_col, "br_col": br_col, "minor_col": minor_col, "mdc_col": mdc_col,
                "voc_col": voc_col, "pw_col": pw_col, "master_rules": None,
                "approved_keys": _p5_pg_approved_keys, "key_col": _p5_pg_key_col, "prefix": "pg"
            })

        # ---------------------------------------------------------------------
        # 🧮 UG और PG — दोनों के लिए ब्रांच-वाइज विषय ब्लॉक्स, डिग्री+ब्रांच समरी,
        # Reason वाले छात्रों की लिस्ट और उनमें से Approve हो चुके छात्रों की लिस्ट निकालना
        # (यह वही logic है जो ऊपर 'भाग 1' में UG/PG टैब पर दिखता है, बस यहाँ पूरे स्कोप के लिए)
        # ---------------------------------------------------------------------
        def _p5_build_scope_summary(df_scope, scope_rules, prefix):
            if df_scope is None or df_scope.empty:
                return [], pd.DataFrame(), pd.DataFrame(), {}, pd.DataFrame()

            _sdb = df_scope.copy()
            _skip_kw2 = ['minor', 'mdc', 'voc', 'skill', 'pw', 'project']
            _sdc = deg_col if (deg_col and deg_col in _sdb.columns) else next(
                (c for c in _sdb.columns if any(k in str(c).lower() for k in ['deg', 'course', 'class'])), None)
            _sbc = br_col if (br_col and br_col in _sdb.columns and br_col != _sdc) else None
            if _sbc is None:
                _sbc = next((c for c in _sdb.columns if c != _sdc and any(k in str(c).lower() for k in ['branch', 'stream'])), None)
            if _sbc is None:
                _sbc = next((c for c in _sdb.columns if c != _sdc and 'subject' in str(c).lower()
                            and not any(k in str(c).lower() for k in _skip_kw2)), None)

            def _sclean_key(series):
                s_ = series.astype(str).str.strip()
                return s_.mask(series.isna() | (s_ == "") | (s_.str.lower() == "nan"), "(खाली/Blank)")

            _sdb["Degree (डिग्री)"] = _sclean_key(_sdb[_sdc]) if _sdc else "—"
            _sdb["Branch (ब्रांच)"] = _sclean_key(_sdb[_sbc]) if _sbc else "—"

            _all_cats2 = [("Minor", "minor", minor_col), ("MDC", "mdc", mdc_col), ("Voc", "voc", voc_col), ("PW", "pw", pw_col)]
            _present_cats2 = [(l_, k_, c_) for (l_, k_, c_) in _all_cats2 if c_ and c_ in _sdb.columns]

            _blocks2 = build_branch_subject_blocks(_sdb, "Degree (डिग्री)", "Branch (ब्रांच)", [(l_, c_) for l_, k_, c_ in _present_cats2])

            if not _present_cats2:
                _summary2 = _sdb.groupby(["Degree (डिग्री)", "Branch (ब्रांच)"]).size().reset_index(name="कुल छात्र (Total)")
                _summary2 = _summary2.sort_values(["Degree (डिग्री)", "कुल छात्र (Total)"], ascending=[True, False]).reset_index(drop=True)
                _summary2 = pd.concat([_summary2, pd.DataFrame([{
                    "Degree (डिग्री)": "कुल योग (GRAND TOTAL)", "Branch (ब्रांच)": "",
                    "कुल छात्र (Total)": int(_summary2["कुल छात्र (Total)"].sum())}])], ignore_index=True)
                return _blocks2, _summary2, pd.DataFrame(), {}, pd.DataFrame()

            if scope_rules is not None and len(_sdb):
                _sdb["_rd"] = _sdb.apply(_row_degree_name, axis=1)
            else:
                _sdb["_rd"] = None

            _ex_pairs2 = get_blank_exempt_pairs(prefix)
            _ex_s2 = pd.Series([(d_, b_) in _ex_pairs2 for d_, b_ in zip(_sdb["Degree (डिग्री)"], _sdb["Branch (ब्रांच)"])], index=_sdb.index)

            for _lbl, _rk, _cn in _present_cats2:
                _empty_s2 = _sdb[_cn].isna() | (_sdb[_cn].astype(str).str.strip() == "")
                _blank_s2 = _empty_s2 & ~_ex_s2
                _norm_s2 = _sdb[_cn].astype(str).map(_norm_txt)
                _wrong_s2 = pd.Series(False, index=_sdb.index)
                if scope_rules is not None:
                    for _rd_name in _sdb["_rd"].dropna().unique():
                        _allowed2 = {_norm_txt(x) for x in ((scope_rules.get(_rd_name) or {}).get(_rk, []))}
                        if _allowed2:
                            _wrong_s2 = _wrong_s2 | ((_sdb["_rd"] == _rd_name) & ~_empty_s2 & ~_norm_s2.isin(_allowed2))
                _sdb[f"_w_{_rk}"] = _wrong_s2
                _sdb[f"_b_{_rk}"] = _blank_s2

            _first_rk2 = _present_cats2[0][1]
            _agg2 = {"_total": (f"_w_{_first_rk2}", "size")}
            for _lbl, _rk, _cn in _present_cats2:
                _agg2[f"_w_{_rk}"] = (f"_w_{_rk}", "sum")
                _agg2[f"_b_{_rk}"] = (f"_b_{_rk}", "sum")
            _grp2 = _sdb.groupby(["Degree (डिग्री)", "Branch (ब्रांच)"]).agg(**_agg2).reset_index()

            def _reason_for_group2(g):
                lines = []
                for _lbl, _rk, _cn in _present_cats2:
                    parts = []
                    w = g[g[f"_w_{_rk}"]]
                    if len(w):
                        for _rd_name, _wg in w.groupby("_rd"):
                            _vc = _wg[_cn].astype(str).str.strip().value_counts()
                            _subj = ", ".join(f"{k} ({v})" for k, v in _vc.items())
                            parts.append(f"🔴 {len(_wg)} छात्रों का विषय {_rd_name or 'इस डिग्री'} के मास्टर नियम में मान्य नहीं है → {_subj}")
                    _bn = int(g[f"_b_{_rk}"].sum())
                    if _bn:
                        parts.append(f"🔵 {_bn} छात्रों का खाली है (डेटा नहीं भरा गया)")
                    if parts:
                        lines.append(f"{_lbl}: " + " ; ".join(parts))
                return "  ||  ".join(lines)

            _reasons2 = {}
            for (_d_k, _b_k), _g in _sdb.groupby(["Degree (डिग्री)", "Branch (ब्रांच)"]):
                _reasons2[(_d_k, _b_k)] = _reason_for_group2(_g)

            _summary2 = _grp2[["Degree (डिग्री)", "Branch (ब्रांच)"]].copy()
            _summary2["कुल छात्र (Total)"] = _grp2["_total"].astype(int)
            for _lbl, _rk, _cn in _present_cats2:
                _w_col2 = _grp2[f"_w_{_rk}"].astype(int)
                _b_col2 = _grp2[f"_b_{_rk}"].astype(int)
                _summary2[f"{_lbl} ✅ सही"] = _summary2["कुल छात्र (Total)"] - _w_col2 - _b_col2
                _summary2[f"{_lbl} 🔴 गलत"] = _w_col2
                _summary2[f"{_lbl} 🔵 खाली"] = _b_col2
            _summary2["📝 कारण (Reason)"] = [
                _reasons2.get((_d_k, _b_k), "")
                for _d_k, _b_k in zip(_grp2["Degree (डिग्री)"], _grp2["Branch (ब्रांच)"])
            ]
            _summary2 = _summary2.sort_values(["Degree (डिग्री)", "कुल छात्र (Total)"], ascending=[True, False]).reset_index(drop=True)
            _tot2 = {c_: "" for c_ in _summary2.columns}
            _tot2["Degree (डिग्री)"] = "कुल योग (GRAND TOTAL)"
            _tot2["Branch (ब्रांच)"] = ""
            for c_ in _summary2.columns:
                if c_ not in ("Degree (डिग्री)", "Branch (ब्रांच)", "📝 कारण (Reason)"):
                    _tot2[c_] = int(_summary2[c_].sum())
            _summary2 = pd.concat([_summary2, pd.DataFrame([_tot2])], ignore_index=True)

            _any_flag2 = pd.Series(False, index=_sdb.index)
            for _lbl, _rk, _cn in _present_cats2:
                _any_flag2 = _any_flag2 | _sdb[f"_w_{_rk}"] | _sdb[f"_b_{_rk}"]
            _stu2 = _sdb[_any_flag2].sort_values(["Degree (डिग्री)", "Branch (ब्रांच)"], kind="stable")
            _orig_cols2 = [c for c in df_scope.columns if c in _stu2.columns]

            def _stu_reason2(r):
                out_ = []
                for _lbl, _rk, _cn in _present_cats2:
                    if r[f"_w_{_rk}"]:
                        _rd_ = r["_rd"] if r["_rd"] else "इस डिग्री"
                        out_.append(f"🔴 {_lbl}: '{str(r[_cn]).strip()}' — {_rd_} के मास्टर नियम में मान्य नहीं")
                    elif r[f"_b_{_rk}"]:
                        out_.append(f"🔵 {_lbl}: खाली (डेटा नहीं भरा गया)")
                return "\n".join(out_)

            _students_df2 = _stu2[_orig_cols2].copy().reset_index(drop=True)
            _students_df2.insert(0, "क्र.सं.", range(1, len(_students_df2) + 1))
            _reason_texts2 = [_stu_reason2(r_) for _, r_ in _stu2.iterrows()]
            _students_df2["📝 कारण (Reason)"] = _reason_texts2
            _student_flags2 = {}
            for _lbl, _rk, _cn in _present_cats2:
                _student_flags2[_cn] = [
                    "w" if w_ else ("b" if b_ else "")
                    for w_, b_ in zip(_stu2[f"_w_{_rk}"].tolist(), _stu2[f"_b_{_rk}"].tolist())
                ]

            # 🟢 Sheet 6 के लिए: इनमें से जो पहले ही Approve किए जा चुके हैं
            _dash_key_col2 = find_student_key_col(_sdb)
            _approved_hits2 = []
            for _pos, (_idx, _row) in enumerate(_stu2.iterrows()):
                _skey2 = get_student_key(_row, _pos, _dash_key_col2)
                _appr2 = get_approval(_skey2, prefix)
                if _appr2:
                    _approved_hits2.append((_idx, _appr2[0], _appr2[1]))

            if _approved_hits2:
                _hit_idx2 = [i for i, _, _ in _approved_hits2]
                _approved_df2 = _stu2.loc[_hit_idx2][_orig_cols2].copy().reset_index(drop=True)
                _approved_df2.insert(0, "✅ Approve किया (By)", [a[1] for a in _approved_hits2])
                _approved_df2.insert(1, "🕒 Approve समय", [a[2] for a in _approved_hits2])
                _approved_df2["📝 कारण (जो गलत/खाली था)"] = [_stu_reason2(r_) for _, r_ in _stu2.loc[_hit_idx2].iterrows()]
            else:
                _approved_df2 = pd.DataFrame()

            return _blocks2, _summary2, _students_df2, _student_flags2, _approved_df2

        _p5_ug_blocks, _p5_ug_summary, _p5_ug_students, _p5_ug_flags, _p5_ug_approved = _p5_build_scope_summary(df_ug_all, ug_master_rules, "ug")
        _p5_pg_blocks, _p5_pg_summary, _p5_pg_students, _p5_pg_flags, _p5_pg_approved = _p5_build_scope_summary(df_pg_all, None, "pg")
        _p5_cat_labels = [l_ for l_, k_, c_ in [("Minor", "minor", minor_col), ("MDC", "mdc", mdc_col), ("Voc", "voc", voc_col), ("PW", "pw", pw_col)] if c_]

        if _p5_master_sheets or _p5_ug_blocks or _p5_pg_blocks:
            _p5_master_excel_bytes = generate_p5_master_full_excel(
                _p5_master_sheets, _p5_cat_labels,
                _p5_ug_blocks, _p5_pg_blocks,
                _p5_ug_summary, _p5_pg_summary,
                _p5_ug_students, _p5_ug_flags, _p5_ug_approved,
                _p5_pg_students, _p5_pg_flags, _p5_pg_approved
            )
            st.download_button(
                label="📥 मास्टर एक्सेल डाउनलोड करें (6 शीट: UG+PG लिस्ट, ब्रांच-वाइज विषय, समरी, Reason, Approved)",
                data=_p5_master_excel_bytes,
                file_name="P5_Master_Dashboard_Data.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key="p5_master_excel_download_btn",
                use_container_width=True
            )
        else:
            st.info("मास्टर एक्सेल के लिए फिलहाल कोई डेटा उपलब्ध नहीं है।")

# =========================================================================
# ⚙️ PANEL 6: ADMIN PANEL
# =========================================================================
elif active_panel == "⚙️ 6. Admin Panel":
    st.title("⚙️ Admin Panel - मास्टर डेटाबेस कंट्रोल")

    # =====================================================================
    # 🎨 लॉगिन स्क्रीन कस्टमाइज़ेशन (टाइटल + लोगो अपलोड)
    # =====================================================================
    if admin_section_toggle("login", "🎨 लॉगिन स्क्रीन कस्टमाइज़ करें", "यहाँ से आप लॉगिन पेज पर दिखने वाला टाइटल बदल सकते हैं, और इमोजी की जगह अपना खुद का लोगो अपलोड कर सकते हैं।"):

        logo_col, title_col = st.columns([1, 1.4])

        with logo_col:
            st.markdown("**🖼️ लोगो अपलोड करें**")
            _current_logo = get_app_setting("login_logo_b64")
            _current_mime = get_app_setting("login_logo_mime", "image/png")
            _current_w = int(get_app_setting("login_logo_width", "140"))
            _current_h = int(get_app_setting("login_logo_height", "140"))
            _current_fit = get_app_setting("login_logo_fit", "contain")
            if _current_logo:
                st.image(f"data:{_current_mime};base64,{_current_logo}", caption="अभी का ओरिजिनल लोगो", width=160)
            else:
                st.info("अभी कोई लोगो नहीं है — डिफ़ॉल्ट इमोजी (🎓🔒) दिख रहा है।")

            uploaded_logo = st.file_uploader("नया लोगो चुनें (PNG/JPG)", type=["png", "jpg", "jpeg", "webp"], key="logo_uploader")

            st.markdown("**📏 लोगो का साइज़ और फिट (पूरी पावर आपके हाथ में)**")
            wc1, wc2 = st.columns(2)
            with wc1:
                new_logo_width = st.slider("चौड़ाई / Width (px)", min_value=30, max_value=400, value=_current_w, step=5, key="logo_width_slider")
            with wc2:
                new_logo_height = st.slider("ऊंचाई / Height (px)", min_value=30, max_value=400, value=_current_h, step=5, key="logo_height_slider")

            new_logo_fit = st.radio(
                "लोगो फिट मोड:",
                options=["contain", "cover"],
                index=(0 if _current_fit == "contain" else 1),
                horizontal=True,
                key="logo_fit_radio",
                help="'contain' = पूरी image दिखेगी, कटेगी नहीं (चारों तरफ थोड़ी खाली जगह आ सकती है) | 'cover' = बॉक्स पूरा भरेगा, लेकिन extra हिस्सा क्रॉप हो सकता है"
            )
            st.caption("👉 अगर लोगो कट रहा है तो हमेशा **'contain'** मोड चुनें, और Width/Height को अपने लोगो के असली अनुपात (aspect ratio) के हिसाब से सेट करें।")

            if _current_logo:
                _preview_html = (
                    f'<div style="text-align:center; background:#f6f8fb; border:1px dashed #c7d1e0; '
                    f'border-radius:10px; padding:14px;">'
                    f'<img src="data:{_current_mime};base64,{_current_logo}" '
                    f'style="width:{new_logo_width}px; height:{new_logo_height}px; object-fit:{new_logo_fit}; '
                    f'border-radius:10px; background:#fff;" /></div>'
                )
                st.markdown("**👁️ लाइव प्रिव्यू (Login स्क्रीन पर ऐसा दिखेगा):**")
                st.markdown(_preview_html, unsafe_allow_html=True)

            if st.button("📏 साइज़ & फिट सेव करें", use_container_width=True, key="save_logo_size_btn"):
                set_app_setting("login_logo_width", str(new_logo_width))
                set_app_setting("login_logo_height", str(new_logo_height))
                set_app_setting("login_logo_fit", new_logo_fit)
                st.success(f"🎉 लोगो साइज़ ({new_logo_width}x{new_logo_height}px, {new_logo_fit}) सेव हो गया!")
                st.rerun()

            lc1, lc2 = st.columns(2)
            with lc1:
                if st.button("💾 लोगो सेव करें", use_container_width=True, disabled=(uploaded_logo is None), key="save_logo_btn"):
                    import base64 as _b64
                    logo_bytes = uploaded_logo.getvalue()
                    encoded = _b64.b64encode(logo_bytes).decode("utf-8")
                    set_app_setting("login_logo_b64", encoded)
                    set_app_setting("login_logo_mime", uploaded_logo.type or "image/png")
                    st.success("🎉 लोगो सफलतापूर्वक सेव हो गया!")
                    st.rerun()
            with lc2:
                if st.button("🗑️ लोगो हटाएं (इमोजी दिखाएं)", use_container_width=True, disabled=(not _current_logo), key="danger_delete_logo_btn"):
                    delete_app_setting("login_logo_b64")
                    delete_app_setting("login_logo_mime")
                    st.success("लोगो हटा दिया गया, अब डिफ़ॉल्ट इमोजी दिखेगा।")
                    st.rerun()

        with title_col:
            st.markdown("**✏️ टाइटल और सबटाइटल एडिट करें**")
            _cur_title = get_app_setting("login_title", "NEP Master Data System")
            _cur_subtitle = get_app_setting("login_subtitle", "अपना पैनल चुनें और आगे बढ़ने के लिए पासवर्ड डालें")
            new_title_input = st.text_input("लॉगिन पेज का टाइटल", value=_cur_title, key="login_title_input")
            new_subtitle_input = st.text_input("लॉगिन पेज का सबटाइटल", value=_cur_subtitle, key="login_subtitle_input")
            if st.button("💾 टाइटल सेव करें", key="save_login_title_btn"):
                set_app_setting("login_title", new_title_input.strip() or "NEP Master Data System")
                set_app_setting("login_subtitle", new_subtitle_input.strip() or "अपना पैनल चुनें और आगे बढ़ने के लिए पासवर्ड डालें")
                st.success("🎉 टाइटल सफलतापूर्वक अपडेट हो गया!")
                st.rerun()

    st.divider()

    # =====================================================================
    # 🔑 पैनल पासवर्ड अपडेट सिस्टम (6 पैनल्स)
    # =====================================================================
    if admin_section_toggle("passwords", "🔑 पैनल पासवर्ड अपडेट करें", "यहाँ से आप किसी भी पैनल (P1-P6) का पासवर्ड बदल सकते हैं। बदलने के बाद उस पैनल में लॉगिन के लिए नया पासवर्ड इस्तेमाल होगा।"):

        pw_cols = st.columns(3)
        new_pw_inputs = {}
        for i, (pk, pname) in enumerate(PANEL_KEY_TO_NAME.items()):
            with pw_cols[i % 3]:
                new_pw_inputs[pk] = st.text_input(f"{pname} का नया पासवर्ड", value="", type="password", key=f"pw_input_{pk}", placeholder="खाली छोड़ें तो नहीं बदलेगा")

        if st.button("🔐 पासवर्ड सेव करें", key="save_panel_passwords_btn"):
            updated_any = False
            for pk, new_pw in new_pw_inputs.items():
                if new_pw.strip():
                    set_panel_password(pk, new_pw.strip())
                    updated_any = True
            if updated_any:
                st.success("🎉 चुने गए पैनल्स के पासवर्ड सफलतापूर्वक अपडेट हो गए हैं!")
            else:
                st.info("कोई नया पासवर्ड नहीं डाला गया, कुछ भी नहीं बदला।")

    st.divider()

    # =====================================================================
    # 👁️ पैनल हाइड / अनहाइड सिस्टम (P1 से P5 तक)
    # =====================================================================
    if admin_section_toggle("panelvis", "👁️ पैनल Hide / Unhide करें (P1 से P5)", "जिस पैनल को Hide करेंगे, उसमें सही पासवर्ड डालने पर भी डेटा नहीं दिखेगा (सिर्फ पैनल का ढांचा दिखेगा)। Unhide करने पर डेटा फिर से दिखने लगेगा।"):

        hide_cols = st.columns(5)
        hide_keys = ["p1", "p2", "p3", "p4", "p5"]
        new_hidden_state = {}
        for i, pk in enumerate(hide_keys):
            with hide_cols[i]:
                current_hidden = is_panel_hidden(pk)
                new_hidden_state[pk] = st.checkbox(f"🙈 {PANEL_KEY_TO_NAME[pk]} Hide करें", value=current_hidden, key=f"hide_chk_{pk}")

        if st.button("💾 Hide/Unhide सेटिंग सेव करें", key="save_hide_settings_btn"):
            for pk, hide_flag in new_hidden_state.items():
                set_panel_hidden(pk, hide_flag)
            st.success("🎉 Hide/Unhide सेटिंग सफलतापूर्वक सेव हो गई है!")
            st.rerun()

    st.divider()

    if admin_section_toggle("backup", "📥 डेटाबेस बैकअप डाउनलोड करें", "🔵 नीला सेल = डेटा गायब है | 🔴 लाल सेल = गलत विषय (मास्टर गाइडलाइन से मिसमैच) | 🟢 हरा सेल = Approved (मान्य किया गया)"):
        df_ug_download = load_permanent_data("UG")
        df_pg_download = load_permanent_data("PG")

        # 🔒 UG मास्टर रूल्स लोड करना ताकि Admin की रंगीन डाउनलोड भी बाकी पैनल्स जैसी सही हो
        _admin_ug_rules = {}
        try:
            cursor.execute("SELECT rules_json FROM locked_rules WHERE panel_prefix = 'ug_master'")
            _locked_row = cursor.fetchone()
            if _locked_row and _locked_row[0]:
                _admin_ug_rules = json.loads(_locked_row[0])
        except Exception:
            pass

        def _detect_cols(df):
            deg_c = next((c for c in df.columns if any(k in c.lower() for k in ['deg', 'course', 'class'])), df.columns[0] if len(df.columns) else None)
            br_c = next((c for c in df.columns if any(k in c.lower() for k in ['branch', 'stream', 'subject'])), None)
            min_c = next((c for c in df.columns if 'minor' in c.lower()), None)
            mdc_c = next((c for c in df.columns if 'mdc' in c.lower()), None)
            voc_c = next((c for c in df.columns if 'voc' in c.lower() or 'skill' in c.lower()), None)
            pw_c = next((c for c in df.columns if any(k in c.lower() for k in ['pw', 'project', 'ce'])), None)
            return deg_c, br_c, min_c, mdc_c, voc_c, pw_c

        c1, c2 = st.columns(2)
        with c1:
            st.markdown("#### 🎓 UG डेटा बैकअप")
            if df_ug_download is not None and not df_ug_download.empty:
                st.download_button(
                    label="📥 UG डेटा CSV डाउनलोड करें (सादा)", 
                    data=df_ug_download.to_csv(index=False).encode('utf-8'), 
                    file_name="Approved_UG_Data_Backup.csv", 
                    mime="text/csv",
                    key="admin_ug_csv_dl"
                )
                deg_c, br_c, min_c, mdc_c, voc_c, pw_c = _detect_cols(df_ug_download)
                _ug_key_col = find_student_key_col(df_ug_download)
                _ug_approved_keys = {a[0] for a in get_all_approvals("ug")}
                ug_colored = generate_colored_excel_bytes(
                    df_ug_download, deg_c, br_c, min_c, mdc_c, voc_c, pw_c,
                    master_rules=_admin_ug_rules, sheet_name="UG_Backup",
                    approved_keys=_ug_approved_keys, key_col=_ug_key_col, prefix="ug"
                )
                st.download_button(
                    label="📥 रंगीन (🔴/🔵) UG डेटा एक्सेल डाउनलोड करें",
                    data=ug_colored,
                    file_name="Approved_UG_Colored_Backup.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key="admin_ug_colored_dl"
                )
            else: 
                st.info("UG डेटाबेस खाली है।")
            
        with c2:
            st.markdown("#### 📜 PG डेटा बैकअप")
            if df_pg_download is not None and not df_pg_download.empty:
                st.download_button(
                    label="📥 PG डेटा CSV डाउनलोड करें (सादा)", 
                    data=df_pg_download.to_csv(index=False).encode('utf-8'), 
                    file_name="Approved_PG_Data_Backup.csv", 
                    mime="text/csv",
                    key="admin_pg_csv_dl"
                )
                deg_c, br_c, min_c, mdc_c, voc_c, pw_c = _detect_cols(df_pg_download)
                _pg_key_col = find_student_key_col(df_pg_download)
                _pg_approved_keys = {a[0] for a in get_all_approvals("pg")}
                pg_colored = generate_colored_excel_bytes(
                    df_pg_download, deg_c, br_c, min_c, mdc_c, voc_c, pw_c,
                    master_rules=None, sheet_name="PG_Backup",
                    approved_keys=_pg_approved_keys, key_col=_pg_key_col, prefix="pg"
                )
                st.download_button(
                    label="📥 रंगीन (🔵) PG डेटा एक्सेल डाउनलोड करें",
                    data=pg_colored,
                    file_name="Approved_PG_Colored_Backup.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key="admin_pg_colored_dl",
                    help="PG के लिए फिलहाल कोई मास्टर सब्जेक्ट नियम सेट नहीं है, इसलिए सिर्फ खाली सेल नीले दिखेंगे।"
                )
            else: 
                st.info("PG डेटाबेस खाली है।")

    st.divider()
    if admin_section_toggle("danger", "🚨 डेंजर ज़ोन", None):
        confirm_reset = st.checkbox("मैं पूरे सिस्टम (रॉ + अप्रूव्ड दोनों डेटाबेस) को रीसेट करने की पुष्टि करता हूँ।")
        also_delete_rules = st.checkbox("⚠️ लॉक किए गए सब्जेक्ट नियम (Minor/MDC/Voc/PW रूल्स) भी डिलीट करें (सामान्यतः इसे टिक न करें)")
        if st.button("💥 ऑल डेटाबेस रीसेट करें", key="danger_reset_all_btn"):
            if confirm_reset:
                cursor.execute("DELETE FROM raw_store")
                cursor.execute("DELETE FROM perma_store")
                # 🔧 फिक्स: locked_rules अब डिफ़ॉल्ट रूप से डिलीट नहीं होगा, ताकि लॉक किए गए सब्जेक्ट नियम
                # नई फ़ाइल अपलोड करने के बाद भी सुरक्षित बने रहें
                if also_delete_rules:
                    cursor.execute("DELETE FROM locked_rules")
                conn.commit()
                st.session_state["deleted_cols"] = []
                if also_delete_rules:
                    st.success("सिस्टम पूरी तरह से रीसेट हो गया है (डेटा + लॉक किए गए नियम दोनों हट गए)!")
                else:
                    st.success("डेटा रीसेट हो गया है! लॉक किए गए सब्जेक्ट नियम सुरक्षित रखे गए हैं।")
                st.rerun()
            else: 
                st.error("कृपया पहले पुष्टि चेकबॉक्स पर टिक करें।")

# =========================================================================
# 🏁 फुटर (हर पेज के नीचे प्रोफेशनल क्रेडिट लाइन)
# =========================================================================
render_footer()
